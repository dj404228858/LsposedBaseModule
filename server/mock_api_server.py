# Python 版 mock API 服务，与 mock_api_server.js 行为对齐。运行: python mock_api_server.py
# 依赖: pip install -r requirements-mock-api.txt

from __future__ import annotations

import asyncio
import base64
import calendar
import hashlib
import json
import os
import random
import re
import io
import secrets
from urllib.parse import quote
from datetime import date, datetime, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal
from typing import Any, Callable, Optional
import math

import aiomysql
import jwt
from aiohttp import ClientSession, ClientTimeout, web

from tx_channel_map import resolve_transaction_channel_cn
from demand_interest import interest_regeneration_schedule, preview_dict_list as demand_interest_preview_dict_list

HOST = os.environ.get("MOCK_API_HOST", "127.0.0.1")
PORT = int(os.environ.get("MOCK_API_PORT", "3000"))
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
MOCK_RESPONSE_LOG_FILE = os.environ.get(
    "MOCK_RESPONSE_LOG_FILE",
    os.path.join(_SCRIPT_DIR, "logs", "mock_api_server.log"),
)
DB_HOST = os.environ.get("MOCK_DB_HOST", "127.0.0.1")
DB_PORT = int(os.environ.get("MOCK_DB_PORT", "3306"))
DB_NAME = os.environ.get("MOCK_DB_NAME", "youzeng")
DB_USER = os.environ.get("MOCK_DB_USER", "youzeng")
DB_PASSWORD = os.environ.get("MOCK_DB_PASSWORD", "WLA8J4itsc6PjBKP")
DEVICE_TABLE = os.environ.get("MOCK_DB_DEVICE_TABLE", "device_registry")
PROFILE_TABLE = os.environ.get("MOCK_DB_PROFILE_TABLE", "profile_list")
TX_TABLE = os.environ.get("MOCK_DB_TX_TABLE", "transaction_flows")
PROFILE_COL_BALANCE = "balance_amount"
TX_COL_GLOBAL = "global_busi_track_no"
TX_COL_INCM_TX_TP_CD = "incm_epn_tx_tp_cd"
TX_COL_DEVICE_COLLECTED = "device_collected"
DEVICE_COL_COLLECT_TX = "collect_transactions_enabled"
DEVICE_COL_AUTO_SYNC_TX = "auto_sync_transactions"
DEVICE_COL_AUTO_SYNC_TX_SINCE = "auto_sync_transactions_since"
PROFILE_COL_CREATOR_ADMIN = "creator_admin_id"
DEVICE_COL_CREATOR_ADMIN = "creator_admin_id"
DEVICE_COL_REMARK = "remark"
DEVICE_REMARK_MAX_LEN = 512
ADMIN_TABLE = os.environ.get("MOCK_ADMIN_TABLE", "admin_accounts")
ADMIN_ROLE_SUPER = 1
ADMIN_ROLE_SUB = 2
PROFILE_CREATE_COST_POINTS = 1
PROFILE_COL_MAIL_SEND_MODE = "mail_send_mode"
PROFILE_COL_CUST_LVL = "cust_lvl"
CUST_LVL_MIN = 1
CUST_LVL_MAX = 7
CUST_LVL_DEFAULT = 1
CUST_LVL_CN = ("", "一星", "二星", "三星", "四星", "五星", "六星", "七星")
ADMIN_COL_POINTS_MODE2 = "points_mode2"
ADMIN_COL_ALLOW_SIMULATE_MAIL = "allow_simulate_mail"
MAIL_SEND_MODE_NORMAL = 1
MAIL_SEND_MODE_SIMULATE = 2
ADMIN_JWT_SECRET = os.environ.get(
    "ADMIN_JWT_SECRET",
    "youzeng-dev-change-in-production-set-long-random-admin-jwt-secret",
)
ADMIN_JWT_EXP_DAYS = int(os.environ.get("ADMIN_JWT_EXP_DAYS", "7"))
ADMIN_BOOTSTRAP_USERNAME = os.environ.get("ADMIN_BOOTSTRAP_USERNAME", "admin")
ADMIN_BOOTSTRAP_PASSWORD = os.environ.get("ADMIN_BOOTSTRAP_PASSWORD", "admin123")

RAW_TX_COLUMN_DEFS: list[tuple[str, str]] = [
    ("accBal", "VARCHAR(64) NULL"),
    ("bkcdMask", "VARCHAR(128) NULL"),
    ("cashExgVatgCd", "VARCHAR(32) NULL"),
    ("chnlKindCode", "VARCHAR(32) NULL"),
    ("currCode", "VARCHAR(32) NULL"),
    ("dtlSeqNo", "VARCHAR(32) NULL"),
    ("dwFlagCode", "VARCHAR(32) NULL"),
    ("globalBusiTrackNo", "VARCHAR(64) NULL"),
    ("ibankFlag", "VARCHAR(32) NULL"),
    ("incmEpnTpCd", "VARCHAR(32) NULL"),
    ("incmEpnTxTpCd", "VARCHAR(32) NULL"),
    ("investProdtCdSets", "VARCHAR(255) NULL"),
    ("investProdtName", "VARCHAR(255) NULL"),
    ("mediumNo", "VARCHAR(128) NULL"),
    ("merDesc", "VARCHAR(255) NULL"),
    ("outTxSriNo", "VARCHAR(128) NULL"),
    ("persInnerAccno", "VARCHAR(128) NULL"),
    ("randomAssignNo", "VARCHAR(128) NULL"),
    ("reckinIncmEpnFlagCd", "VARCHAR(32) NULL"),
    ("saccnoSeqNo", "VARCHAR(32) NULL"),
    ("servNo", "VARCHAR(64) NULL"),
    ("subtxNo", "VARCHAR(128) NULL"),
    ("summ", "VARCHAR(255) NULL"),
    ("transInMobileNo", "VARCHAR(64) NULL"),
    ("txAmt", "VARCHAR(64) NULL"),
    ("txDate", "VARCHAR(32) NULL"),
    ("txOpsAccno", "VARCHAR(128) NULL"),
    ("txOpsName", "VARCHAR(128) NULL"),
    ("txRemark", "VARCHAR(255) NULL"),
    ("txTime", "VARCHAR(32) NULL"),
]

# 统一使用北京时间（UTC+8）生成日期/时间字符串，避免服务器时区为 UTC 时出现日期偏移。
UTC8_TZ = timezone(timedelta(hours=8))


def now_utc8() -> datetime:
    return datetime.now(UTC8_TZ)


def now_utc8_naive() -> datetime:
    return now_utc8().replace(tzinfo=None)

_pool: Optional[aiomysql.Pool] = None

HISTORY_APPLY_MAIL_TABLE = os.environ.get("MOCK_DB_HISTORY_APPLY_MAIL_TABLE", "history_apply_mail_records")
AUTO_FLOW_TEMPLATE_TABLE = os.environ.get("MOCK_DB_AUTO_FLOW_TEMPLATE_TABLE", "auto_flow_templates")
AUTO_FLOW_MAX_GENERATE = int(os.environ.get("MOCK_AUTO_FLOW_MAX_GENERATE", "8000"))


def _json_default(obj: Any) -> Any:
    if isinstance(obj, Decimal):
        return float(obj)
    if isinstance(obj, (datetime,)):
        return obj.isoformat(sep=" ")
    if isinstance(obj, bytes):
        return obj.decode("utf-8", errors="replace")
    raise TypeError(f"Object of type {type(obj)} is not JSON serializable")


def json_dumps(payload: Any) -> str:
    return json.dumps(payload, ensure_ascii=False, indent=2, default=_json_default)


async def get_pool() -> aiomysql.Pool:
    global _pool
    if _pool is None:
        _pool = await aiomysql.create_pool(
            host=DB_HOST,
            port=DB_PORT,
            user=DB_USER,
            password=DB_PASSWORD,
            db=DB_NAME,
            minsize=1,
            maxsize=8,
            charset="utf8mb4",
            autocommit=True,
        )
    return _pool


def normalize_row(row: Any) -> dict[str, Any]:
    if not row or not isinstance(row, dict):
        return {}
    return dict(row)


async def query_one(sql: str, params: Optional[tuple[Any, ...]] = None) -> Optional[dict[str, Any]]:
    pool = await get_pool()
    async with pool.acquire() as conn:
        async with conn.cursor(aiomysql.DictCursor) as cur:
            await cur.execute(sql, params or ())
            row = await cur.fetchone()
            return normalize_row(row) if row else None


async def query_all(sql: str, params: Optional[tuple[Any, ...]] = None) -> list[dict[str, Any]]:
    pool = await get_pool()
    async with pool.acquire() as conn:
        async with conn.cursor(aiomysql.DictCursor) as cur:
            await cur.execute(sql, params or ())
            rows = await cur.fetchall()
            return [normalize_row(r) for r in (rows or [])]


async def execute(sql: str, params: Optional[tuple[Any, ...]] = None) -> None:
    pool = await get_pool()
    async with pool.acquire() as conn:
        async with conn.cursor() as cur:
            await cur.execute(sql, params or ())


async def execute_rowcount(sql: str, params: Optional[tuple[Any, ...]] = None) -> int:
    pool = await get_pool()
    async with pool.acquire() as conn:
        async with conn.cursor() as cur:
            await cur.execute(sql, params or ())
            return int(cur.rowcount)


async def execute_insert(sql: str, params: Optional[tuple[Any, ...]] = None) -> int:
    pool = await get_pool()
    async with pool.acquire() as conn:
        async with conn.cursor() as cur:
            await cur.execute(sql, params or ())
            return int(cur.lastrowid)


def pick_did(request: web.Request) -> str:
    q = request.query.get("did") or ""
    h = request.headers.get("x-device-id") or request.headers.get("did") or ""
    return str(q or h or "").strip()


def to_num_or_null(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    return n


def to_id_or_null(v: Any) -> Optional[int]:
    try:
        n = int(v)
    except (TypeError, ValueError):
        return None
    if n <= 0:
        return None
    return n


def to_bool_int(v: Any) -> int:
    if v is True or v == 1:
        return 1
    s = str(v if v is not None else "").strip().lower()
    if not s:
        return 0
    if s in ("1", "true", "yes", "on"):
        return 1
    return 0


def must_non_empty_text(v: Any) -> str:
    return str(v if v is not None else "").strip()


def normalize_device_remark(v: Any) -> Optional[str]:
    """管理端设备备注：去首尾空白，空串存 NULL，超长截断。"""
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    if len(s) > DEVICE_REMARK_MAX_LEN:
        return s[:DEVICE_REMARK_MAX_LEN]
    return s


def to_nullable_id_or_invalid(v: Any) -> Any:
    if v is None:
        return None
    s = str(v).strip()
    if not s or s.lower() == "null":
        return None
    tid = to_id_or_null(s)
    if tid is None:
        return Ellipsis
    return tid


async def load_device_context(did: str) -> Optional[dict[str, Any]]:
    sql = f"""
    SELECT
      d.id AS device_id,
      d.did,
      d.profile_id,
      d.`{DEVICE_COL_COLLECT_TX}` AS collect_transactions_enabled,
      COALESCE(d.`{DEVICE_COL_AUTO_SYNC_TX}`, 1) AS auto_sync_transactions,
      d.`{DEVICE_COL_AUTO_SYNC_TX_SINCE}` AS auto_sync_transactions_since,
      p.id AS person_id,
      p.person_name,
      p.card_no,
      p.id_card,
      p.phone_no,
      p.stamp_no,
      p.title,
      p.`{PROFILE_COL_BALANCE}` AS profile_balance,
      p.`{PROFILE_COL_CUST_LVL}` AS cust_lvl
    FROM `{DEVICE_TABLE}` d
    LEFT JOIN `{PROFILE_TABLE}` p ON p.id = d.profile_id
    WHERE d.did = %s AND d.is_active = 1
    LIMIT 1
    """
    return await query_one(sql, (did,))


async def load_device_context_any(did: str) -> Optional[dict[str, Any]]:
    sql = f"""
    SELECT
      d.id AS device_id,
      d.did,
      d.profile_id,
      d.is_active,
      d.`{DEVICE_COL_COLLECT_TX}` AS collect_transactions_enabled,
      COALESCE(d.`{DEVICE_COL_AUTO_SYNC_TX}`, 1) AS auto_sync_transactions,
      d.`{DEVICE_COL_AUTO_SYNC_TX_SINCE}` AS auto_sync_transactions_since,
      p.id AS person_id,
      p.person_name,
      p.card_no,
      p.id_card,
      p.phone_no,
      p.stamp_no,
      p.title,
      p.`{PROFILE_COL_BALANCE}` AS profile_balance,
      p.`{PROFILE_COL_CUST_LVL}` AS cust_lvl
    FROM `{DEVICE_TABLE}` d
    LEFT JOIN `{PROFILE_TABLE}` p ON p.id = d.profile_id
    WHERE d.did = %s
    LIMIT 1
    """
    return await query_one(sql, (did,))


async def fetch_device_collect_transactions_enabled_by_did(did: str) -> Optional[int]:
    """现查 `device_registry.collect_transactions_enabled`（经 to_bool_int 为 0/1），无此 did 返回 None。供 mock-xhx 等避免依赖其它查询结果里可能过时的开关字段。"""
    row = await query_one(
        f"SELECT `{DEVICE_COL_COLLECT_TX}` AS c FROM `{DEVICE_TABLE}` WHERE did=%s LIMIT 1",
        (str(did or "").strip(),),
    )
    if not row:
        return None
    return to_bool_int(row.get("c"))


def is_device_profile_bound(device: Optional[dict[str, Any]]) -> bool:
    if not device or not isinstance(device, dict):
        return False
    return device.get("profile_id") is not None and device.get("person_id") is not None


async def resolve_bound_device_context_for_mock(did: str) -> dict[str, Any]:
    device = await load_device_context(did)
    if not device:
        return {"ok": False, "status": 404, "error": f"设备不存在或未启用: did={did}"}
    if not is_device_profile_bound(device):
        return {"ok": False, "status": 409, "error": f"设备未绑定资料，跳过mock替换: did={did}"}
    return {"ok": True, "device": device}


def _obfuscated_acc_no_for_psbc_app(card_no: str) -> str:
    """邮储 App 账户介质号展示风格：O + 32 位 hex，由卡号稳定映射。"""
    c = str(card_no or "").strip()
    if not c:
        return ""
    return "O" + hashlib.sha256(f"psbcAccNo|{c}".encode("utf-8")).hexdigest()[:32]


def build_qry_acc_bas_info_t080025_response(req_msg_id: str, device: Optional[dict[str, Any]]) -> dict[str, Any]:
    """sn13/api/account/qryAccBasInfo/T080025：accList[].accLevel 固定为 \"1\"。"""
    acc_list: list[dict[str, Any]] = []
    cfm_flag = "0"
    dld_flag = ""
    hhv_flag = ""
    if device and is_device_profile_bound(device):
        card = str(device.get("card_no") or "").strip()
        acc_no = _obfuscated_acc_no_for_psbc_app(card)
        bkcd = format_bkcd_mask(card) if card else ""
        cust = str(device.get("person_name") or "").strip()
        card_tp = "02" if card else ""
        acc_list.append(
            {
                "accAlias": "",
                "accLevel": "1",
                "accNo": acc_no,
                "accStaFlag": "",
                "accType": "10",
                "addChl": "0",
                "addType": "1",
                "alias": "",
                "bankCode": "403100000004",
                "bankLogo": "",
                "bankName": "中国邮政储蓄银行2",
                "bkcdMask": bkcd,
                "cardBokType": "2",
                "cardTpCode": card_tp,
                "custName": cust,
                "debitCardFlag": "",
                "defauAccFlag": "",
                "digitalCertificateAccFlag": "0",
                "eSubAccFlag": "0",
                "existDigitalCertificateAccFlag": "0",
                "hasMobiRemitAuth": True,
                "hasMobiTransAuth": True,
                "issCode": "",
                "masterCardFlag": "",
                "openPay": "",
                "randomAssignNo": acc_no,
                "topicName": "",
                "unionPayFlag": "1",
            }
        )
        cfm_flag = "1"
        dld_flag = "1"
        hhv_flag = "0"
    return {
        "code": "000000",
        "data": {
            "accList": acc_list,
            "accnumQt": "",
            "cfmFlag": cfm_flag,
            "cityList": [],
            "cozyTip": "",
            "dldFlagCd": dld_flag,
            "hhvCrcardFlag": hhv_flag,
            "mobileNo": "",
            "yhWjMigSta": "",
        },
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


# 主查询已对 `global_busi_track_no` / `incm_epn_tx_tp_cd` 使用 AS globalBusiTrackNo / incmEpnTxTpCd，
# raw 扩展列中若再选同名 camelCase 列，MySQL 会返回带表前缀的重复键，且污染 JSON。
_RAW_COLS_OMIT_IF_ALIASED_IN_TX_SELECT = frozenset({"globalBusiTrackNo", "incmEpnTxTpCd"})


def _raw_cols_sql() -> str:
    return ",\n      ".join(
        f"`{d[0]}`"
        for d in RAW_TX_COLUMN_DEFS
        if d[0] not in _RAW_COLS_OMIT_IF_ALIASED_IN_TX_SELECT
    )


async def read_raw_mock_data_by_did(did: str) -> list[dict[str, Any]]:
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        raise RuntimeError(binding["error"])
    device = binding["device"]
    tx_sql = f"""
    SELECT
      id AS tx_id,
      person_id AS person_id,
      tx_datetime AS `交易日期`,
      tx_type AS `交易类型`,
      currency AS `交易币种`,
      tx_amount AS `交易金额`,
      account_balance AS `账户余额`,
      counterparty_name AS `对手方户名`,
      counterparty_account AS `对手方账户`,
      counterparty_bank AS `对手银行`,
      remark AS `附言`,
      channel AS `交易方式`,
      `{TX_COL_GLOBAL}` AS globalBusiTrackNo,
      `{TX_COL_INCM_TX_TP_CD}` AS incmEpnTxTpCd,
      {_raw_cols_sql()}
    FROM `{TX_TABLE}`
    WHERE person_id = %s
    ORDER BY tx_datetime DESC, id DESC
    """
    rows = await query_all(tx_sql, (device["person_id"],))
    return [hydrate_tx_aliases_from_raw(dict(r)) for r in rows]


async def read_tx_by_did_and_global_track_no(did: str, global_busi_track_no: str) -> dict[str, Any]:
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        raise RuntimeError(binding["error"])
    device = binding["device"]
    tx_sql = f"""
    SELECT
      id AS tx_id,
      person_id AS person_id,
      tx_datetime AS `交易日期`,
      tx_type AS `交易类型`,
      currency AS `交易币种`,
      tx_amount AS `交易金额`,
      account_balance AS `账户余额`,
      counterparty_name AS `对手方户名`,
      counterparty_account AS `对手方账户`,
      counterparty_bank AS `对手银行`,
      remark AS `附言`,
      channel AS `交易方式`,
      `{TX_COL_GLOBAL}` AS globalBusiTrackNo,
      `{TX_COL_INCM_TX_TP_CD}` AS incmEpnTxTpCd,
      {_raw_cols_sql()}
    FROM `{TX_TABLE}`
    WHERE person_id = %s AND `{TX_COL_GLOBAL}` = %s
    ORDER BY id DESC
    LIMIT 1
    """
    row = await query_one(tx_sql, (device["person_id"], global_busi_track_no))
    return {
        "basic_info": {
            "户名": device.get("person_name") or "",
            "卡号": device.get("card_no") or "",
            "身份证号": device.get("id_card") or "",
            "电话号码": device.get("phone_no") or "",
            "印章流水": device.get("stamp_no") or "",
            "抬头": device.get("title") or "",
            "profile_balance": device.get("profile_balance"),
        },
        "transaction": hydrate_tx_aliases_from_raw(row) if row else None,
    }


def fmt_ymd_hms(input_val: Any) -> dict[str, str]:
    if input_val is None:
        return {"ymd": "", "ymdhms": ""}
    if isinstance(input_val, datetime):
        dt_obj = input_val
        y = str(dt_obj.year)
        M = str(dt_obj.month).zfill(2)
        d = str(dt_obj.day).zfill(2)
        h = str(dt_obj.hour).zfill(2)
        m = str(dt_obj.minute).zfill(2)
        s = str(dt_obj.second).zfill(2)
        ymd = f"{y}{M}{d}"
        return {"ymd": ymd, "ymdhms": f"{ymd}{h}{m}{s}"}
    t = str(input_val).strip()
    m = re.match(
        r"(\d{4})[-/]?(\d{2})[-/]?(\d{2})(?:[T\s]?(\d{2}))?:?(\d{2})?:?(\d{2})?",
        t,
        re.I,
    )
    if not m:
        return {"ymd": "", "ymdhms": ""}
    ymd = f"{m.group(1)}{m.group(2)}{m.group(3)}"
    hh = m.group(4) or "00"
    mm = m.group(5) or "00"
    ss = m.group(6) or "00"
    return {"ymd": ymd, "ymdhms": f"{ymd}{hh}{mm}{ss}"}


def to_amt_str(v: Any) -> str:
    if v is None:
        return "0.00"
    if isinstance(v, Decimal):
        return f"{float(v):.2f}"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return f"{float(v):.2f}"
    s = str(v).replace(",", "").strip()
    if not s:
        return "0.00"
    try:
        n = float(s)
    except ValueError:
        return "0.00"
    return f"{n:.2f}"


def normalize_ymd8(input_val: Any) -> str:
    if input_val is None:
        return ""
    digits = re.sub(r"[^\d]", "", str(input_val))
    if len(digits) < 8:
        return ""
    return digits[:8]


def build_hash_digits(seed: str, length: int) -> str:
    if length <= 0:
        return ""
    out = ""
    step = 0
    while len(out) < length:
        h = hashlib.sha1(f"{seed}|{step}".encode("utf-8")).hexdigest()
        n = str(int(h, 16))
        out += n
        step += 1
    return out[:length]


def build_import_raw_ids(seed: str) -> dict[str, str]:
    # 参考样本格式生成：persInnerAccno(19位数字)、randomAssignNo(40位hex)、
    # servNo(8位数字+Z+2位数字)、subtxNo(32位数字且以99开头)。
    random_assign_no = hashlib.sha1(f"{seed}|randomAssignNo".encode("utf-8")).hexdigest()
    serv_no = f"{build_hash_digits(f'{seed}|servNoA', 8)}Z{build_hash_digits(f'{seed}|servNoB', 2)}"
    subtx_no = f"99{build_hash_digits(f'{seed}|subtxNo', 30)}"
    return {
        "persInnerAccno": build_hash_digits(f"{seed}|persInnerAccno", 19),
        "randomAssignNo": random_assign_no,
        "servNo": serv_no,
        "subtxNo": subtx_no,
    }


def build_stable_tx_identity(row: Optional[dict[str, Any]], dt: dict[str, str]) -> dict[str, str]:
    tx_id = str(row.get("tx_id") if row else "") if row else ""
    person_id = str(row.get("person_id") if row else "") if row else ""
    amt = to_amt_str((row or {}).get("交易金额") or (row or {}).get("tx_amount"))
    balance = to_amt_str((row or {}).get("账户余额") or (row or {}).get("account_balance"))
    account = str((row or {}).get("对手方账户") or (row or {}).get("counterparty_account") or "")
    tx_type = str((row or {}).get("交易类型") or (row or {}).get("tx_type") or "")
    cp_name = str((row or {}).get("对手方户名") or (row or {}).get("counterparty_name") or "")
    remark = str((row or {}).get("附言") or (row or {}).get("remark") or "")
    seed = f"{tx_id}|{person_id}|{dt['ymdhms']}|{amt}|{balance}|{account}|{cp_name}|{tx_type}|{remark}"
    # 生成格式：前 8 位为交易日期(yyyymmdd)，后 23 位为“看似随机”的纯数字，总长 31。
    # 这里使用 hash 派生数字串，保证同一流水在同一输入下稳定，且满足“数字随机样式”的要求。
    ymd8 = str(dt.get("ymd") or "").strip()
    if not re.fullmatch(r"\d{8}", ymd8):
        ymd8 = ymd_today_utc8()
    tail23 = build_hash_digits(f"{seed}|globalBusiTrackNo", 23)
    global_busi_track_no = f"{ymd8}{tail23}"[:31]
    return {"globalBusiTrackNo": global_busi_track_no}


def resolve_profile_balance_str(device: Optional[dict[str, Any]]) -> str:
    if not device or not isinstance(device, dict):
        return "0.00"
    return to_amt_str(device.get("profile_balance"))


_OPENACC_DATE_RE = re.compile(r"^(\d{4})([/\-]?)(\d{1,2})([/\-]?)(\d{1,2})$")


def shift_openacc_date_back_years(raw: str, years_back: int = 3) -> Optional[str]:
    """qryAccDtl openaccDate：在原值基础上年份减 years_back（如 2026/03/24 → 2023/03/24）。"""
    s = str(raw or "").strip()
    if not s or s == "--" or years_back <= 0:
        return None
    m = _OPENACC_DATE_RE.match(s)
    if m:
        y = int(m.group(1)) - years_back
        if y < 1:
            return None
        sep = m.group(2) or ""
        mo, da = int(m.group(3)), int(m.group(5))
        if not sep:
            return f"{y:04d}{mo:02d}{da:02d}"
        sep2 = m.group(4) or sep
        return f"{y:04d}{sep}{mo:02d}{sep2}{da:02d}"
    ym = re.search(r"\d{4}", s)
    if ym:
        y = int(ym.group(0)) - years_back
        if y < 1:
            return None
        return re.sub(r"\d{4}", f"{y:04d}", s, count=1)
    return None


PROFILE_BALANCE_FIELD_KEYS: tuple[str, ...] = (
    "accBal",
    "avalBal",
    "accAvalBal",
    "curTotalAmt",
    "totalAmt",
    "totalAmount",
    "addAmt",
    "addCurTotalAmt",
    "availableAmount",
    "addAssetAmt",
    "addCurDepAmt",
    "addDepAmt",
    "addLocCurDepAmt",
    "curDepAmt",
    "depAmt",
    "locCurDepAmt",
    "totalAssetAmt",
    "masterBanknoteAccBal",
    "masterBanknoteAvalBal",
    "masterConvergeAccBal",
    "masterConvergeAvalBal",
    "masterUSDAccBal",
    "masterUSDAvalBal",
    "usableBal",
    "finBal",
    "fixBal",
)


def apply_profile_balance_to_node(node: Any, balance_str: str) -> None:
    """与 Lsposed YouzengPsbcHelpers.applyProfileBalanceToPayload 一致：递归替换余额类字段。"""
    if isinstance(node, dict):
        for k in PROFILE_BALANCE_FIELD_KEYS:
            if k in node and not isinstance(node[k], (dict, list)):
                node[k] = balance_str
        for v in node.values():
            apply_profile_balance_to_node(v, balance_str)
    elif isinstance(node, list):
        for item in node:
            apply_profile_balance_to_node(item, balance_str)


def apply_openacc_date_shift_to_qry_acc_dtl_payload(payload: dict[str, Any], years_back: int = 3) -> bool:
    """就地修改网关明文 JSON 的 data.openaccDate / openAccDate。"""
    if not isinstance(payload, dict):
        return False
    data = payload.get("data")
    if not isinstance(data, dict):
        return False
    changed = False
    for key in ("openaccDate", "openAccDate"):
        if key not in data:
            continue
        raw = str(data.get(key) or "").strip()
        shifted = shift_openacc_date_back_years(raw, years_back)
        if shifted and shifted != raw:
            data[key] = shifted
            changed = True
    return changed


def apply_qry_acc_dtl_acc_level(payload: dict[str, Any], acc_level: str = "1") -> bool:
    """qryAccDtl/T080002：data.accLevel 固定为指定值（默认 \"1\"）。"""
    if not isinstance(payload, dict):
        return False
    data = payload.get("data")
    if not isinstance(data, dict):
        return False
    if str(data.get("accLevel") or "") == str(acc_level):
        return False
    data["accLevel"] = str(acc_level)
    return True


async def build_qry_acc_dtl_patched_response(
    req_msg_id: str, did: str, host_rsp_plain: str
) -> dict[str, Any]:
    """
    qryAccDtl/T080002：在邮储网关明文上覆盖资料余额、accLevel=1、openaccDate 前移 3 年。
    Hook 或 sn13 路由传入 hostRspPlain（SM4 解密后的 JSON 字符串）。
    """
    hp = str(host_rsp_plain or "").strip()
    if not hp:
        return {
            "code": "000016",
            "msg": "missing hostRspPlain",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        return {
            "code": "000016",
            "msg": str(binding.get("error") or "device not bound"),
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    try:
        obj = json.loads(hp)
    except json.JSONDecodeError as e:
        return {
            "code": "000016",
            "msg": f"invalid hostRspPlain json: {e}",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    if not isinstance(obj, dict):
        return {
            "code": "000016",
            "msg": "hostRspPlain root must be object",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    device = binding["device"]
    bal = resolve_profile_balance_str(device)
    apply_profile_balance_to_node(obj, bal)
    apply_qry_acc_dtl_acc_level(obj, "1")
    apply_openacc_date_shift_to_qry_acc_dtl_payload(obj, 3)
    obj["reqMsgId"] = req_msg_id or str(obj.get("reqMsgId") or "")
    _data = (obj.get("data") or {}) if isinstance(obj.get("data"), dict) else {}
    print(
        "[mock-api] qryAccDtl patched did=%s balance=%s accLevel=%s openaccDate=%s"
        % (did, bal, _data.get("accLevel"), _data.get("openaccDate"))
    )
    return obj


async def build_qry_trans_acc_bal_patched_response(
    req_msg_id: str, did: str, host_rsp_plain: str
) -> dict[str, Any]:
    """
    qryTransAccBal/T080770：在网关明文上把 avalBal（及同额 lenderNowSelfOwnFunds）同步为资料余额。
    """
    hp = str(host_rsp_plain or "").strip()
    if not hp:
        return {
            "code": "000016",
            "msg": "missing hostRspPlain",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        return {
            "code": "000016",
            "msg": str(binding.get("error") or "device not bound"),
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    try:
        obj = json.loads(hp)
    except json.JSONDecodeError as e:
        return {
            "code": "000016",
            "msg": f"invalid hostRspPlain json: {e}",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    if not isinstance(obj, dict):
        return {
            "code": "000016",
            "msg": "hostRspPlain root must be object",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    if str(obj.get("code") or "") != "000000":
        return {
            "code": "000016",
            "msg": f"host qryTransAccBal not success: {obj.get('code')}",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    device = binding["device"]
    bal = resolve_profile_balance_str(device)
    apply_profile_balance_to_node(obj, bal)
    data = obj.get("data") if isinstance(obj.get("data"), dict) else None
    if isinstance(data, dict) and "lenderNowSelfOwnFunds" in data:
        data["lenderNowSelfOwnFunds"] = bal
    obj["reqMsgId"] = req_msg_id or str(obj.get("reqMsgId") or "")
    _data = data if isinstance(data, dict) else {}
    print(
        "[mock-api] qryTransAccBal T080770 patched did=%s avalBal=%s lenderNowSelfOwnFunds=%s"
        % (did, _data.get("avalBal"), _data.get("lenderNowSelfOwnFunds"))
    )
    return obj


_STAR_LVL_ICON_RE = re.compile(r"(星级标签_)(\d+)(星)")
_STAR_LVL_ICON_DEFAULT = (
    "https://static.mobile-bank.psbc.com//static/my/starLvl/星级标签_{n}星@3x.png"
)


def cust_lvl_star_icon_url(lvl: int, current: str = "") -> str:
    """「我的」页星级角标：按 N 星改写 starLvlIcon 文件名。"""
    n = normalize_cust_lvl(lvl)
    cur = str(current or "").strip()
    if cur and _STAR_LVL_ICON_RE.search(cur):
        return _STAR_LVL_ICON_RE.sub(rf"\g<1>{n}\g<3>", cur, count=1)
    if cur:
        replaced = re.sub(r"_(\d+)星", f"_{n}星", cur, count=1)
        if replaced != cur:
            return replaced
    return _STAR_LVL_ICON_DEFAULT.format(n=n)


def apply_t020104_cust_lvl(payload: dict[str, Any], cust_lvl: int) -> bool:
    """pageDataQuery/T020104：覆盖 custLvl / custLvlCode，并同步 starLvlIcon。"""
    if not isinstance(payload, dict):
        return False
    data = payload.get("data")
    if not isinstance(data, dict):
        return False
    lvl = normalize_cust_lvl(cust_lvl)
    code = cust_lvl_to_code(lvl)
    label = cust_lvl_to_label(lvl)
    icon = cust_lvl_star_icon_url(lvl, str(data.get("starLvlIcon") or ""))
    changed = False
    if str(data.get("custLvlCode") or "") != code:
        data["custLvlCode"] = code
        changed = True
    if label and str(data.get("custLvl") or "") != label:
        data["custLvl"] = label
        changed = True
    if str(data.get("starLvlIcon") or "") != icon:
        data["starLvlIcon"] = icon
        changed = True
    return changed


async def build_t020104_cust_lvl_patched_response(
    req_msg_id: str, did: str, host_rsp_plain: str
) -> dict[str, Any]:
    """
    pageDataQuery/T020104：在邮储网关明文上覆盖资料账户星级。
    后台 1 星 → custLvlCode=0，2 星 → 1，以此类推。
    """
    hp = str(host_rsp_plain or "").strip()
    if not hp:
        return {
            "code": "000016",
            "msg": "missing hostRspPlain",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        return {
            "code": "000016",
            "msg": str(binding.get("error") or "device not bound"),
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    try:
        obj = json.loads(hp)
    except json.JSONDecodeError as e:
        return {
            "code": "000016",
            "msg": f"invalid hostRspPlain json: {e}",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    if not isinstance(obj, dict):
        return {
            "code": "000016",
            "msg": "hostRspPlain root must be object",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    device = binding["device"]
    lvl = normalize_cust_lvl(device.get("cust_lvl"))
    apply_t020104_cust_lvl(obj, lvl)
    obj["reqMsgId"] = req_msg_id or str(obj.get("reqMsgId") or "")
    _data = (obj.get("data") or {}) if isinstance(obj.get("data"), dict) else {}
    print(
        "[mock-api] T020104 patched did=%s custLvl=%s custLvlCode=%s starLvlIcon=%s"
        % (did, _data.get("custLvl"), _data.get("custLvlCode"), _data.get("starLvlIcon"))
    )
    return obj


def apply_t070708_cust_lvl(payload: dict[str, Any], cust_lvl: int) -> bool:
    """initQyzq/T070708：只覆盖 custCurStarLvl（1 起算，与「我的」展示星级相同）。

    本接口的 custLvl / custLable 不是 T020104 的 custLvlCode。
    原包 1 星是 custCurStarLvl=1、custLvl=0；若再把 custLvl 写成 0 起算码，
    前端会叠成 当前星+码（4+3=7）。
    """
    if not isinstance(payload, dict):
        return False
    data = payload.get("data")
    if not isinstance(data, dict):
        return False
    star = str(normalize_cust_lvl(cust_lvl))
    if str(data.get("custCurStarLvl") or "") == star:
        return False
    data["custCurStarLvl"] = star
    return True


async def build_t070708_cust_lvl_patched_response(
    req_msg_id: str, did: str, host_rsp_plain: str
) -> dict[str, Any]:
    """
    initQyzq/T070708：在邮储网关明文上覆盖资料账户星级。
    后台 N 星 → custCurStarLvl=N（与「我的」一致，不改 custLvl）。
    """
    hp = str(host_rsp_plain or "").strip()
    if not hp:
        return {
            "code": "000016",
            "msg": "missing hostRspPlain",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        return {
            "code": "000016",
            "msg": str(binding.get("error") or "device not bound"),
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    try:
        obj = json.loads(hp)
    except json.JSONDecodeError as e:
        return {
            "code": "000016",
            "msg": f"invalid hostRspPlain json: {e}",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    if not isinstance(obj, dict):
        return {
            "code": "000016",
            "msg": "hostRspPlain root must be object",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    device = binding["device"]
    lvl = normalize_cust_lvl(device.get("cust_lvl"))
    apply_t070708_cust_lvl(obj, lvl)
    obj["reqMsgId"] = req_msg_id or str(obj.get("reqMsgId") or "")
    _data = (obj.get("data") or {}) if isinstance(obj.get("data"), dict) else {}
    print(
        "[mock-api] T070708 patched did=%s custCurStarLvl=%s (custLvl/growthLevel untouched)"
        % (did, _data.get("custCurStarLvl"))
    )
    return obj


def hydrate_tx_aliases_from_raw(row: Optional[dict[str, Any]]) -> dict[str, Any]:
    r = dict(row or {})
    tx_type = str(r.get("交易类型") or r.get("tx_type") or r.get("summ") or "").strip()
    tx_amt = r.get("交易金额") if r.get("交易金额") is not None else (
        r.get("tx_amount") if r.get("tx_amount") is not None else r.get("txAmt")
    )
    acc_bal = r.get("账户余额") if r.get("账户余额") is not None else (
        r.get("account_balance") if r.get("account_balance") is not None else r.get("accBal")
    )
    tx_ops_name = r.get("对手方户名") if r.get("对手方户名") is not None else r.get("txOpsName")
    tx_ops_accno = r.get("对手方账户") if r.get("对手方账户") is not None else r.get("txOpsAccno")
    tx_ops_bank = r.get("对手银行") if r.get("对手银行") is not None else (
        r.get("counterparty_bank") if r.get("counterparty_bank") is not None else ""
    )
    tx_remark = r.get("附言") if r.get("附言") is not None else r.get("txRemark")
    tx_channel = r.get("交易方式") if r.get("交易方式") is not None else r.get("chnlKindCode")
    if tx_type and r.get("交易类型") is None:
        r["交易类型"] = tx_type
    if tx_amt is not None and r.get("交易金额") is None:
        r["交易金额"] = tx_amt
    if acc_bal is not None and r.get("账户余额") is None:
        r["账户余额"] = acc_bal
    if tx_ops_name is not None and r.get("对手方户名") is None:
        r["对手方户名"] = tx_ops_name
    if tx_ops_accno is not None and r.get("对手方账户") is None:
        r["对手方账户"] = tx_ops_accno
    if tx_ops_bank is not None and r.get("对手银行") is None:
        r["对手银行"] = tx_ops_bank
    if tx_remark is not None and r.get("附言") is None:
        r["附言"] = tx_remark
    if tx_channel is not None and r.get("交易方式") is None:
        r["交易方式"] = tx_channel
    if not r.get("txDate") and r.get("交易日期"):
        r["txDate"] = normalize_ymd8(r["交易日期"])
    tx_type = str(r.get("交易类型") or r.get("tx_type") or r.get("summ") or "").strip()
    if tx_type and not str(r.get("summ") or "").strip():
        r["summ"] = tx_type
    if r.get("txAmt") is None:
        amt_src = r.get("交易金额") if r.get("交易金额") is not None else r.get("tx_amount")
        if amt_src is not None:
            r["txAmt"] = amt_src
    if not str(r.get("incmEpnTxTpCd") or "").strip():
        tp_cd = incm_epn_tp_cd_from_record(r)
        r["incmEpnTxTpCd"] = resolve_incm_epn_tx_tp_cd_from_row(r, tx_type, tp_cd)
    if str(r.get("dwFlagCode") or "").strip() not in ("1", "2"):
        r["dwFlagCode"] = incm_epn_tp_cd_from_record(r)
    gb_snake = r.get(TX_COL_GLOBAL)
    if gb_snake is not None and str(gb_snake).strip() and not str(r.get("globalBusiTrackNo") or "").strip():
        r["globalBusiTrackNo"] = str(gb_snake).strip()
    return r


def resolve_tx_identity_from_row(row: Optional[dict[str, Any]], dt: dict[str, str]) -> dict[str, str]:
    gb = str(row.get("globalBusiTrackNo") or "").strip() if row else ""
    if gb:
        return {"globalBusiTrackNo": gb}
    return build_stable_tx_identity(row, dt)


def build_tx_row_from_body(body: dict[str, Any], tx_id: int, person_id: int) -> dict[str, Any]:
    return {
        "tx_id": tx_id,
        "person_id": person_id,
        "tx_datetime": body.get("tx_datetime") or "1970-01-01 00:00:00",
        "tx_type": body.get("tx_type"),
        "tx_amount": to_num_or_null(body.get("tx_amount")) or 0,
        "account_balance": to_num_or_null(body.get("account_balance")),
        "counterparty_name": body.get("counterparty_name"),
        "counterparty_account": body.get("counterparty_account"),
        "remark": body.get("remark"),
        "incmEpnTxTpCd": body.get("incmEpnTxTpCd"),
    }


def normalize_raw_tx_value(v: Any) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def build_raw_tx_model_from_body(body: dict[str, Any], global_track_no: str, tx_tp_cd: str) -> dict[str, Optional[str]]:
    model: dict[str, Optional[str]] = {}
    for name, _ in RAW_TX_COLUMN_DEFS:
        if name == "summ" and body.get("_preserve_summ_raw"):
            v = body.get(name)
            model[name] = None if v is None else str(v)
        else:
            model[name] = normalize_raw_tx_value(body.get(name))
    if not model.get("globalBusiTrackNo"):
        model["globalBusiTrackNo"] = normalize_raw_tx_value(global_track_no)
    if not model.get("incmEpnTxTpCd"):
        model["incmEpnTxTpCd"] = normalize_raw_tx_value(tx_tp_cd)
    return model


async def ensure_tx_identity_columns_and_backfill() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, TX_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    if TX_COL_GLOBAL.lower() not in existing:
        await execute(f"ALTER TABLE `{TX_TABLE}` ADD COLUMN `{TX_COL_GLOBAL}` VARCHAR(64) NULL")
    if TX_COL_INCM_TX_TP_CD.lower() not in existing:
        await execute(f"ALTER TABLE `{TX_TABLE}` ADD COLUMN `{TX_COL_INCM_TX_TP_CD}` VARCHAR(16) NULL")

    missing_rows = await query_all(
        f"""SELECT id AS tx_id, person_id, tx_datetime, tx_type, tx_amount, account_balance,
            counterparty_name, counterparty_account, remark,
            `{TX_COL_GLOBAL}` AS globalBusiTrackNo,
            `{TX_COL_INCM_TX_TP_CD}` AS incmEpnTxTpCd
     FROM `{TX_TABLE}`
     WHERE COALESCE(`{TX_COL_GLOBAL}`, '') = ''
        OR COALESCE(`{TX_COL_INCM_TX_TP_CD}`, '') = ''"""
    )
    for row in missing_rows or []:
        dt = fmt_ymd_hms(row.get("tx_datetime"))
        identity = resolve_tx_identity_from_row(row, dt)
        tx_type_text = "" if row.get("tx_type") is None else str(row.get("tx_type"))
        amt_num_raw = float(to_amt_str(row.get("tx_amount")))
        is_out = amt_num_raw < 0
        is_in_text = "汇入" in tx_type_text or "转入" in tx_type_text
        is_out_text = "汇出" in tx_type_text or "转出" in tx_type_text or "支出" in tx_type_text
        tp_cd = "2" if is_out else "1"
        if is_in_text and not is_out_text:
            tp_cd = "1"
        if is_out_text and not is_in_text:
            tp_cd = "2"
        tx_tp_cd = str(row.get("incmEpnTxTpCd") or "") or resolve_incm_epn_tx_tp_cd(tx_type_text, tp_cd)
        await execute(
            f"UPDATE `{TX_TABLE}` SET `{TX_COL_GLOBAL}`=%s, `{TX_COL_INCM_TX_TP_CD}`=%s WHERE id=%s",
            (identity["globalBusiTrackNo"], tx_tp_cd, int(row["tx_id"])),
        )


async def ensure_raw_tx_model_columns() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, TX_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    for name, ddl in RAW_TX_COLUMN_DEFS:
        if str(name).lower() not in existing:
            await execute(f"ALTER TABLE `{TX_TABLE}` ADD COLUMN `{name}` {ddl}")


async def ensure_tx_device_collected_column() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, TX_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    if TX_COL_DEVICE_COLLECTED.lower() not in existing:
        await execute(
            f"ALTER TABLE `{TX_TABLE}` ADD COLUMN `{TX_COL_DEVICE_COLLECTED}` "
            "TINYINT(1) NOT NULL DEFAULT 0"
        )


async def drop_legacy_tx_identity_columns() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, TX_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    if "subtx_no" in existing:
        await execute(f"ALTER TABLE `{TX_TABLE}` DROP COLUMN `subtx_no`")
    if "random_assign_no" in existing:
        await execute(f"ALTER TABLE `{TX_TABLE}` DROP COLUMN `random_assign_no`")


async def ensure_history_apply_mail_table() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, HISTORY_APPLY_MAIL_TABLE),
    )
    if not rows:
        await execute(
            f"""CREATE TABLE `{HISTORY_APPLY_MAIL_TABLE}` (
  `id` BIGINT NOT NULL AUTO_INCREMENT,
  `person_id` BIGINT NOT NULL,
  `apply_date` VARCHAR(8) NOT NULL,
  `apply_time` VARCHAR(6) NOT NULL,
  `begin_date` VARCHAR(8) NOT NULL,
  `dline_date` VARCHAR(8) NOT NULL,
  `deliver_status` VARCHAR(8) NOT NULL DEFAULT '1',
  `draw_no` VARCHAR(64) NOT NULL DEFAULT '',
  `email` VARCHAR(255) NOT NULL DEFAULT '',
  `file_id` VARCHAR(255) NOT NULL DEFAULT '',
  `flag` VARCHAR(8) NOT NULL DEFAULT '0',
  `mail_no` VARCHAR(64) NOT NULL DEFAULT '',
  `medium_no` VARCHAR(128) NOT NULL DEFAULT '',
  `send_batch` VARCHAR(255) NOT NULL DEFAULT '',
  `tp_flag` VARCHAR(8) NOT NULL DEFAULT '0',
  `created_at` DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
  PRIMARY KEY (`id`),
  KEY `idx_person_created` (`person_id`, `created_at`),
  KEY `idx_person_apply` (`person_id`, `apply_date`, `apply_time`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4"""
        )
        return


async def ensure_auto_flow_templates_table() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, AUTO_FLOW_TEMPLATE_TABLE),
    )
    if rows:
        return
    await execute(
        f"""CREATE TABLE `{AUTO_FLOW_TEMPLATE_TABLE}` (
  `id` BIGINT NOT NULL AUTO_INCREMENT,
  `ref_name` VARCHAR(255) NOT NULL DEFAULT '',
  `gen_settings` JSON NOT NULL,
  `lines` JSON NOT NULL,
  `creator_admin_id` BIGINT NOT NULL,
  `created_at` TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
  `updated_at` TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
  PRIMARY KEY (`id`),
  KEY `idx_auto_flow_creator` (`creator_admin_id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4"""
    )


def _history_apply_send_batch(mail_no: str, draw_no: str, seq_no: str, apply_at: datetime) -> str:
    ts = int(apply_at.timestamp())
    md5 = hashlib.md5(f"{mail_no}|{draw_no}|{ts}".encode("utf-8")).hexdigest()
    prefix = str(seq_no or "").strip() or "0"
    return f"{prefix}_fast.hefei.{ts}.{md5}"


def _format_history_mail_seq_no(n: int) -> str:
    """邮件主题序列号：当天第几封邮件（真机常见 001/002/...）。"""
    n = int(n)
    if n <= 0:
        return "001"
    # 超过 999 时不截断，直接用自然数（仍保持兼容）
    return f"{n:03d}" if n <= 999 else str(n)


async def _next_history_mail_seq_no_for_date(apply_date: str) -> str:
    """按自然日统计发送次数，返回下一次的序列号字符串（001 起）。"""
    row = await query_one(
        f"SELECT COUNT(1) AS cnt FROM `{HISTORY_APPLY_MAIL_TABLE}` WHERE apply_date=%s",
        (str(apply_date or "").strip(),),
    )
    cnt = int((row or {}).get("cnt") or 0)
    return _format_history_mail_seq_no(cnt + 1)


async def insert_history_apply_mail_record(
    *,
    person_id: int,
    medium_no: str | None,
    to_email: str,
    query: dict[str, Any],
    apply_at: datetime,
    mail_seq_no: str,
) -> None:
    begin_date = normalize_ymd8(query.get("beginDate")) or apply_at.strftime("%Y%m%d")
    dline_date = normalize_ymd8(query.get("dlineDate")) or apply_at.strftime("%Y%m%d")
    draw_no = str(query.get("drawNo") or "").strip()
    seq_no = str(mail_seq_no or "").strip() or "001"
    # mailNo 参考真机格式：YYYYMMDD-001（当天第几封）
    mail_no = f"{apply_at.strftime('%Y%m%d')}-{seq_no}"
    send_batch = _history_apply_send_batch(mail_no, draw_no, seq_no, apply_at)
    masked_medium = format_bkcd_mask(str(medium_no or "").strip()) if medium_no else ""
    await execute(
        f"""INSERT INTO `{HISTORY_APPLY_MAIL_TABLE}`
  (person_id, apply_date, apply_time, begin_date, dline_date, deliver_status, draw_no, email, file_id, flag, mail_no, medium_no, send_batch, tp_flag)
  VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
        (
            int(person_id),
            apply_at.strftime("%Y%m%d"),
            apply_at.strftime("%H%M%S"),
            begin_date,
            dline_date,
            "1",
            draw_no,
            str(to_email or "").strip(),
            "",
            "0",
            mail_no,
            masked_medium,
            send_batch,
            "0",
        ),
    )


async def list_history_apply_mail_records(person_id: int, limit: int = 50) -> list[dict[str, Any]]:
    lim = max(1, min(int(limit), 200))
    rows = await query_all(
        f"""SELECT apply_date, apply_time, begin_date, dline_date, deliver_status, draw_no,
            email, file_id, flag, mail_no, medium_no, send_batch, tp_flag
     FROM `{HISTORY_APPLY_MAIL_TABLE}`
     WHERE person_id=%s
     ORDER BY apply_date DESC, apply_time DESC, id DESC
     LIMIT {lim}""",
        (int(person_id),),
    )
    out: list[dict[str, Any]] = []
    for r in rows or []:
        out.append(
            {
                "applyDate": str(r.get("apply_date") or ""),
                "applyTime": str(r.get("apply_time") or ""),
                "beginDate": str(r.get("begin_date") or ""),
                "deliverStatus": str(r.get("deliver_status") or "1"),
                "dlineDate": str(r.get("dline_date") or ""),
                "drawNo": str(r.get("draw_no") or ""),
                "email": str(r.get("email") or ""),
                "fileId": str(r.get("file_id") or ""),
                "flag": str(r.get("flag") or "0"),
                "mailNo": str(r.get("mail_no") or ""),
                "mediumNo": str(r.get("medium_no") or ""),
                "sendBatch": str(r.get("send_batch") or ""),
                "tpFlag": str(r.get("tp_flag") or "0"),
            }
        )
    return out


async def ensure_profile_balance_column() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, PROFILE_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    if PROFILE_COL_BALANCE.lower() not in existing:
        await execute(
            f"ALTER TABLE `{PROFILE_TABLE}` ADD COLUMN `{PROFILE_COL_BALANCE}` DECIMAL(18,2) NULL"
        )


async def ensure_device_auto_sync_tx_column() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, DEVICE_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    if DEVICE_COL_AUTO_SYNC_TX.lower() not in existing:
        await execute(
            f"ALTER TABLE `{DEVICE_TABLE}` ADD COLUMN `{DEVICE_COL_AUTO_SYNC_TX}` "
            "TINYINT(1) NOT NULL DEFAULT 1"
        )
    if DEVICE_COL_AUTO_SYNC_TX_SINCE.lower() not in existing:
        await execute(
            f"ALTER TABLE `{DEVICE_TABLE}` ADD COLUMN `{DEVICE_COL_AUTO_SYNC_TX_SINCE}` "
            "DATETIME NULL COMMENT '自动同步开启时刻，仅同步该时间之后的流水'"
        )
        await execute(
            f"UPDATE `{DEVICE_TABLE}` SET `{DEVICE_COL_AUTO_SYNC_TX_SINCE}`=%s "
            f"WHERE `{DEVICE_COL_AUTO_SYNC_TX}`=1 AND `{DEVICE_COL_AUTO_SYNC_TX_SINCE}` IS NULL",
            (now_utc8_naive(),),
        )


async def ensure_profile_creator_admin_column() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, PROFILE_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    if PROFILE_COL_CREATOR_ADMIN.lower() not in existing:
        await execute(f"ALTER TABLE `{PROFILE_TABLE}` ADD COLUMN `{PROFILE_COL_CREATOR_ADMIN}` BIGINT NULL")
        await execute(
            f"ALTER TABLE `{PROFILE_TABLE}` ADD INDEX `idx_profile_creator_admin` (`{PROFILE_COL_CREATOR_ADMIN}`)"
        )


async def ensure_profile_mail_send_mode_column() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, PROFILE_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    if PROFILE_COL_MAIL_SEND_MODE.lower() not in existing:
        await execute(
            f"ALTER TABLE `{PROFILE_TABLE}` ADD COLUMN `{PROFILE_COL_MAIL_SEND_MODE}` "
            f"TINYINT NOT NULL DEFAULT {MAIL_SEND_MODE_NORMAL} "
            "COMMENT '1=普通发件 2=模拟真实地址'"
        )


async def ensure_profile_cust_lvl_column() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, PROFILE_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    if PROFILE_COL_CUST_LVL.lower() not in existing:
        await execute(
            f"ALTER TABLE `{PROFILE_TABLE}` ADD COLUMN `{PROFILE_COL_CUST_LVL}` "
            f"TINYINT NOT NULL DEFAULT {CUST_LVL_DEFAULT} "
            "COMMENT '账户星级 1-7，接口 custLvlCode=星级-1'"
        )


async def ensure_admin_points_mode2_column() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, ADMIN_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    if ADMIN_COL_POINTS_MODE2.lower() not in existing:
        await execute(
            f"ALTER TABLE `{ADMIN_TABLE}` ADD COLUMN `{ADMIN_COL_POINTS_MODE2}` "
            "INT NOT NULL DEFAULT 0 COMMENT '模拟真实地址资料积分'"
        )


async def ensure_admin_allow_simulate_mail_column() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, ADMIN_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    if ADMIN_COL_ALLOW_SIMULATE_MAIL.lower() not in existing:
        await execute(
            f"ALTER TABLE `{ADMIN_TABLE}` ADD COLUMN `{ADMIN_COL_ALLOW_SIMULATE_MAIL}` "
            "TINYINT(1) NOT NULL DEFAULT 0 COMMENT '1=允许二级使用模拟真实发件'"
        )


def normalize_mail_send_mode(value: Any) -> int:
    try:
        n = int(value)
    except (TypeError, ValueError):
        return MAIL_SEND_MODE_NORMAL
    return MAIL_SEND_MODE_SIMULATE if n == MAIL_SEND_MODE_SIMULATE else MAIL_SEND_MODE_NORMAL


def normalize_cust_lvl(value: Any) -> int:
    """资料账户星级：1–7，非法值回落到 1 星。"""
    try:
        n = int(str(value).strip())
    except (TypeError, ValueError):
        return CUST_LVL_DEFAULT
    if n < CUST_LVL_MIN or n > CUST_LVL_MAX:
        return CUST_LVL_DEFAULT
    return n


def cust_lvl_to_code(lvl: int) -> str:
    """1 星 → \"0\"，2 星 → \"1\"，以此类推。"""
    return str(normalize_cust_lvl(lvl) - 1)


def cust_lvl_to_label(lvl: int) -> str:
    n = normalize_cust_lvl(lvl)
    if 0 < n < len(CUST_LVL_CN):
        return CUST_LVL_CN[n]
    return ""


def admin_allows_simulate_mail(admin: dict[str, Any]) -> bool:
    """一级始终可用；二级需 allow_simulate_mail=1。"""
    if int(admin.get("role") or 0) == ADMIN_ROLE_SUPER:
        return True
    return to_bool_int(admin.get("allow_simulate_mail")) == 1


def _admin_select_points_and_perms() -> str:
    return (
        f"points_balance, `{ADMIN_COL_POINTS_MODE2}` AS points_mode2, "
        f"`{ADMIN_COL_ALLOW_SIMULATE_MAIL}` AS allow_simulate_mail"
    )


async def ensure_device_creator_admin_column() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, DEVICE_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    if DEVICE_COL_CREATOR_ADMIN.lower() not in existing:
        await execute(f"ALTER TABLE `{DEVICE_TABLE}` ADD COLUMN `{DEVICE_COL_CREATOR_ADMIN}` BIGINT NULL")
        await execute(
            f"ALTER TABLE `{DEVICE_TABLE}` ADD INDEX `idx_device_creator_admin` (`{DEVICE_COL_CREATOR_ADMIN}`)"
        )


async def ensure_device_remark_column() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s""",
        (DB_NAME, DEVICE_TABLE),
    )
    existing = {str(r.get("COLUMN_NAME") or "").lower() for r in rows}
    if DEVICE_COL_REMARK.lower() not in existing:
        await execute(
            f"ALTER TABLE `{DEVICE_TABLE}` ADD COLUMN `{DEVICE_COL_REMARK}` VARCHAR({DEVICE_REMARK_MAX_LEN}) DEFAULT NULL"
        )


async def ensure_admin_accounts_table() -> None:
    await execute(
        f"""
        CREATE TABLE IF NOT EXISTS `{ADMIN_TABLE}` (
          `id` BIGINT NOT NULL AUTO_INCREMENT,
          `username` VARCHAR(64) NOT NULL,
          `password_hash` VARCHAR(255) NOT NULL,
          `role` TINYINT NOT NULL COMMENT '1=一级 2=二级',
          `parent_admin_id` BIGINT DEFAULT NULL,
          `points_balance` INT NOT NULL DEFAULT 0 COMMENT '普通发件资料积分',
          `points_mode2` INT NOT NULL DEFAULT 0 COMMENT '模拟真实地址资料积分',
          `allow_simulate_mail` TINYINT(1) NOT NULL DEFAULT 0 COMMENT '1=允许二级使用模拟真实发件',
          `created_at` TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
          `updated_at` TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
          PRIMARY KEY (`id`),
          UNIQUE KEY `uk_admin_username` (`username`),
          KEY `idx_admin_parent` (`parent_admin_id`)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
        """
    )


async def bootstrap_default_super_admin_if_empty() -> None:
    cnt = await query_one(f"SELECT COUNT(1) AS c FROM `{ADMIN_TABLE}`")
    if cnt and int(cnt.get("c") or 0) > 0:
        return
    h = hash_admin_password(ADMIN_BOOTSTRAP_PASSWORD)
    await execute_insert(
        f"""INSERT INTO `{ADMIN_TABLE}` (username, password_hash, role, parent_admin_id, points_balance)
        VALUES (%s, %s, %s, NULL, 0)""",
        (ADMIN_BOOTSTRAP_USERNAME, h, ADMIN_ROLE_SUPER),
    )
    print(
        f"[mock-api] 已创建默认一级管理员: username={ADMIN_BOOTSTRAP_USERNAME} "
        f"（请登录后修改或设置环境变量 ADMIN_BOOTSTRAP_PASSWORD）",
    )


def device_auto_sync_transactions_enabled(device: dict[str, Any]) -> bool:
    """设备是否允许 mock-xhx 按开启时刻之后的流水自动采集入库（默认开启以兼容旧数据）。"""
    v = device.get("auto_sync_transactions")
    if v is None:
        return True
    return to_bool_int(v) == 1


def device_auto_sync_since_dt(device: Optional[dict[str, Any]]) -> Optional[datetime]:
    if not device:
        return None
    return _parse_naive_datetime_for_compare(device.get("auto_sync_transactions_since"))


async def stamp_device_auto_sync_since(device_id: int, since_dt: Optional[datetime]) -> None:
    await execute(
        f"UPDATE `{DEVICE_TABLE}` SET `{DEVICE_COL_AUTO_SYNC_TX_SINCE}`=%s WHERE id=%s",
        (since_dt, int(device_id)),
    )


async def resolve_auto_sync_since_for_collect(device: dict[str, Any]) -> Optional[datetime]:
    """返回自动同步起点；已开启但尚未记时则写入当前北京时间，本批不回溯历史流水。"""
    if not device_auto_sync_transactions_enabled(device):
        return None
    since_dt = device_auto_sync_since_dt(device)
    if since_dt is not None:
        return since_dt
    device_id = to_id_or_null(device.get("device_id") or device.get("id"))
    if not device_id:
        return None
    since_dt = now_utc8_naive()
    await stamp_device_auto_sync_since(device_id, since_dt)
    device["auto_sync_transactions_since"] = since_dt
    return since_dt


async def build_profile_balance_response(req_msg_id: str, did: str) -> dict[str, Any]:
    device = await load_device_context(did)
    if not device:
        return {
            "code": "000016",
            "msg": "暂时无法处理您的请求，请返回或退出后重试",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    return {
        "code": "000000",
        "data": {"balance": resolve_profile_balance_str(device)},
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


async def ensure_device_profile_nullable() -> None:
    row = await query_one(
        """SELECT IS_NULLABLE
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s AND COLUMN_NAME = 'profile_id'
     LIMIT 1""",
        (DB_NAME, DEVICE_TABLE),
    )
    nullable = str(row.get("IS_NULLABLE") or "").upper() == "YES" if row else False
    if not nullable:
        await execute(f"ALTER TABLE `{DEVICE_TABLE}` MODIFY COLUMN profile_id BIGINT NULL")


async def ensure_device_collect_tx_column() -> None:
    rows = await query_all(
        """SELECT COLUMN_NAME
     FROM INFORMATION_SCHEMA.COLUMNS
     WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s AND COLUMN_NAME = %s""",
        (DB_NAME, DEVICE_TABLE, DEVICE_COL_COLLECT_TX),
    )
    if not rows:
        await execute(
            f"ALTER TABLE `{DEVICE_TABLE}` ADD COLUMN `{DEVICE_COL_COLLECT_TX}` TINYINT(1) NOT NULL DEFAULT 0"
        )


def to_tx_datetime_string(tx_date_raw: Any, tx_time_raw: Any) -> str:
    digits = re.sub(r"[^\d]", "", str(tx_time_raw if tx_time_raw is not None else ""))
    if len(digits) >= 14:
        s = digits[:14]
        return f"{s[:4]}-{s[4:6]}-{s[6:8]} {s[8:10]}:{s[10:12]}:{s[12:14]}"
    ymd = normalize_ymd8(tx_date_raw) or "19700101"
    return f"{ymd[:4]}-{ymd[4:6]}-{ymd[6:8]} 00:00:00"


def infer_counterparty_bank_name(tx_ops_accno: Any, _tx_ops_bank_no: Any, _tx_ops_bank_nm: Any) -> Optional[str]:
    acc = re.sub(r"[^\d]", "", str(tx_ops_accno if tx_ops_accno is not None else ""))
    if not acc:
        return None
    prefix_map = [
        ("621225", "中国工商银行"),
        ("622123", "中国工商银行"),
        ("621226", "中国工商银行"),
        ("622848", "中国农业银行"),
        ("622846", "中国农业银行"),
        ("621798", "中国农业银行"),
        ("621788", "中国银行"),
        ("621700", "中国建设银行"),
        ("622700", "中国建设银行"),
        ("436742", "中国建设银行"),
        ("622188", "中国邮政储蓄银行"),
        ("622180", "中国邮政储蓄银行"),
        ("621799", "中国邮政储蓄银行"),
    ]
    for prefix, bank in prefix_map:
        if acc.startswith(prefix):
            return bank
    return None


def _interest_needle_from_parts(*parts: Any) -> str:
    return "".join(str(p or "").strip() for p in parts)


def is_interest_like_transaction(*, summ: Any = "", tx_type: Any = "", remark: Any = "", tp_cd: Any = "") -> bool:
    """判断是否利息类流水（与子类编码 1004 对齐）：用于采集与存量合并。"""
    n = _interest_needle_from_parts(summ, tx_type, remark)
    if any(k in n for k in ("利息", "结息", "活期利息", "存款利息", "贷款利息")):
        return True
    return str(tp_cd or "").strip() == "1004"


def extract_collected_tx_rows_from_payload(payload: Any) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    if not payload or not isinstance(payload, dict):
        return out
    blocks = payload.get("data", {}).get("fieldItemInfo") if isinstance(payload.get("data"), dict) else None
    if not isinstance(blocks, list):
        return out
    for block in blocks:
        lst = block.get("detailList") if isinstance(block, dict) else None
        if not isinstance(lst, list):
            continue
        for item in lst:
            if not item or not isinstance(item, dict):
                continue
            global_busi_track_no = str(item.get("globalBusiTrackNo") or "").strip()
            if not global_busi_track_no:
                continue
            tx_amt_raw = to_num_or_null(item.get("txAmt")) or 0
            is_out = str(item.get("dwFlagCode") or item.get("incmEpnTpCd") or "1").strip() == "2"
            tx_amount = -abs(tx_amt_raw) if is_out else abs(tx_amt_raw)
            raw = {
                "accBal": None if item.get("accBal") is None else str(item.get("accBal")),
                "bkcdMask": None if item.get("bkcdMask") is None else str(item.get("bkcdMask")),
                "cashExgVatgCd": None if item.get("cashExgVatgCd") is None else str(item.get("cashExgVatgCd")),
                "chnlKindCode": None if item.get("chnlKindCode") is None else str(item.get("chnlKindCode")),
                "currCode": None if item.get("currCode") is None else str(item.get("currCode")),
                "dtlSeqNo": None if item.get("dtlSeqNo") is None else str(item.get("dtlSeqNo")),
                "dwFlagCode": None if item.get("dwFlagCode") is None else str(item.get("dwFlagCode")),
                "globalBusiTrackNo": None if item.get("globalBusiTrackNo") is None else str(item.get("globalBusiTrackNo")),
                "ibankFlag": None if item.get("ibankFlag") is None else str(item.get("ibankFlag")),
                "incmEpnTpCd": None if item.get("incmEpnTpCd") is None else str(item.get("incmEpnTpCd")),
                "incmEpnTxTpCd": None if item.get("incmEpnTxTpCd") is None else str(item.get("incmEpnTxTpCd")),
                "investProdtCdSets": None if item.get("investProdtCdSets") is None else str(item.get("investProdtCdSets")),
                "investProdtName": None if item.get("investProdtName") is None else str(item.get("investProdtName")),
                "mediumNo": None if item.get("mediumNo") is None else str(item.get("mediumNo")),
                "merDesc": None if item.get("merDesc") is None else str(item.get("merDesc")),
                "outTxSriNo": None if item.get("outTxSriNo") is None else str(item.get("outTxSriNo")),
                "persInnerAccno": None if item.get("persInnerAccno") is None else str(item.get("persInnerAccno")),
                "randomAssignNo": None if item.get("randomAssignNo") is None else str(item.get("randomAssignNo")),
                "reckinIncmEpnFlagCd": None if item.get("reckinIncmEpnFlagCd") is None else str(item.get("reckinIncmEpnFlagCd")),
                "saccnoSeqNo": None if item.get("saccnoSeqNo") is None else str(item.get("saccnoSeqNo")),
                "servNo": None if item.get("servNo") is None else str(item.get("servNo")),
                "subtxNo": None if item.get("subtxNo") is None else str(item.get("subtxNo")),
                "summ": None if item.get("summ") is None else str(item.get("summ")),
                "transInMobileNo": None if item.get("transInMobileNo") is None else str(item.get("transInMobileNo")),
                "txAmt": None if item.get("txAmt") is None else str(item.get("txAmt")),
                "txDate": None if item.get("txDate") is None else str(item.get("txDate")),
                "txOpsAccno": None if item.get("txOpsAccno") is None else str(item.get("txOpsAccno")),
                "txOpsName": None if item.get("txOpsName") is None else str(item.get("txOpsName")),
                "txRemark": None if item.get("txRemark") is None else str(item.get("txRemark")),
                "txTime": None if item.get("txTime") is None else str(item.get("txTime")),
            }
            summ_v = item.get("summ") or ""
            chnl_gate = str(item.get("chnlKindCode") or "").strip()
            channel_cn = resolve_transaction_channel_cn(
                summ=str(summ_v or ""),
                tx_type=str(summ_v or ""),
                remark=str(item.get("txRemark") or ""),
                channel=chnl_gate,
                chnl_kind_code=chnl_gate,
            )
            apply_out_tx_sri_no_for_import_or_collect(
                raw,
                summ_for_check=str(summ_v or "").strip(),
                ymd8=normalize_ymd8(item.get("txDate")),
                global_track_no=global_busi_track_no,
            )
            out.append(
                {
                    "tx_datetime": to_tx_datetime_string(item.get("txDate"), item.get("txTime")),
                    "tx_type": item.get("summ"),
                    "currency": str(item.get("currCode")) if item.get("currCode") else "人民币",
                    "tx_amount": tx_amount,
                    "account_balance": to_num_or_null(item.get("accBal")),
                    "counterparty_name": item.get("txOpsName"),
                    "counterparty_account": item.get("txOpsAccno"),
                    "counterparty_bank": item.get("txOpsBankNm"),
                    "remark": item.get("txRemark"),
                    "channel": channel_cn or chnl_gate,
                    "incm_epn_tx_tp_cd": item.get("incmEpnTxTpCd"),
                    "global_busi_track_no": global_busi_track_no,
                    "raw": raw,
                }
            )
    return out


def _parse_naive_datetime_for_compare(val: Any) -> Optional[datetime]:
    """将 tx_datetime / DB 返回值规范为 naive datetime，供比较大小。"""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.replace(tzinfo=None) if val.tzinfo else val
    if isinstance(val, date) and not isinstance(val, datetime):
        return datetime(val.year, val.month, val.day)
    s = str(val).strip()
    if not s:
        return None
    if len(s) >= 19:
        try:
            return datetime.strptime(s[:19], "%Y-%m-%d %H:%M:%S")
        except ValueError:
            pass
    if len(s) >= 16:
        try:
            return datetime.strptime(s[:16], "%Y-%m-%d %H:%M")
        except ValueError:
            pass
    if len(s) >= 10:
        try:
            return datetime.strptime(s[:10], "%Y-%m-%d")
        except ValueError:
            pass
    return None


def _tx_row_leq(a_dt: Any, a_id: int, b_dt: Any, b_id: int) -> bool:
    """比较两条流水在时间轴上的先后：(tx_datetime, id) 升序。"""
    da = _parse_naive_datetime_for_compare(a_dt)
    db = _parse_naive_datetime_for_compare(b_dt)
    if da is not None and db is not None:
        if da != db:
            return da < db
    else:
        sa = str(a_dt or "")[:32]
        sb = str(b_dt or "")[:32]
        if sa != sb:
            return sa < sb
    return int(a_id) <= int(b_id)


def max_tx_datetime_from_xhx_host_payload(payload: dict[str, Any]) -> Optional[datetime]:
    """从 Hook 下发的 hostRspPlain（网关 JSON）中解析明细，取最晚一笔交易时间。"""
    rows = extract_collected_tx_rows_from_payload(payload)
    best: Optional[datetime] = None
    for row in rows:
        dt = _parse_naive_datetime_for_compare(row.get("tx_datetime"))
        if dt is None:
            continue
        if best is None or dt > best:
            best = dt
    return best


def _calendar_date_yyyy_mm_dd_from_tx_datetime(val: Any) -> str:
    """与自然日匹配的日期串 YYYY-MM-DD（结息常为固定某日）。"""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    s = str(val or "").strip()
    if len(s) >= 10:
        return s[:10]
    return ""


async def fetch_latest_tx_datetime_for_person(person_id: int) -> Optional[datetime]:
    """库中该自然人下交易流水的最大 tx_datetime。"""
    row = await query_one(
        f"SELECT MAX(tx_datetime) AS m FROM `{TX_TABLE}` WHERE person_id=%s",
        (int(person_id),),
    )
    if not row:
        return None
    return _parse_naive_datetime_for_compare(row.get("m"))


async def collect_transactions_by_did(
    did: str, payload_obj: Any, *, min_tx_datetime: Optional[datetime] = None
) -> dict[str, Any]:
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        return {"ok": False, "status": binding["status"], "error": binding["error"]}
    rows = extract_collected_tx_rows_from_payload(payload_obj)
    if not rows:
        return {"ok": True, "inserted": 0, "skipped": 0, "total": 0}
    person_id_pre = int(binding["device"]["person_id"])
    min_dt = _parse_naive_datetime_for_compare(min_tx_datetime) if min_tx_datetime is not None else None
    # 入库前列表「时间上最后一笔」的赛后余额，供新入账行在其基础上滚余额
    balance_snapshot = await fetch_latest_tx_balance_snapshot(person_id_pre)
    inserted = 0
    skipped = 0
    inserted_ids: list[int] = []
    for row in rows:
        if min_dt is not None:
            row_dt = _parse_naive_datetime_for_compare(row.get("tx_datetime"))
            if row_dt is None or row_dt <= min_dt:
                skipped += 1
                continue
        exists = await query_one(
            f"""SELECT id, counterparty_bank FROM `{TX_TABLE}` WHERE person_id=%s AND `{TX_COL_GLOBAL}`=%s LIMIT 1""",
            (person_id_pre, row["global_busi_track_no"]),
        )
        if exists:
            existing_bank = str(exists.get("counterparty_bank") or "").strip()
            incoming_bank = str(row.get("counterparty_bank") or "").strip()
            if not existing_bank and incoming_bank:
                await execute(
                    f"UPDATE `{TX_TABLE}` SET counterparty_bank=%s WHERE id=%s AND person_id=%s",
                    (incoming_bank, int(exists["id"]), person_id_pre),
                )
            skipped += 1
            continue
        base_cols = [
            "person_id",
            "tx_datetime",
            "tx_type",
            "currency",
            "tx_amount",
            "account_balance",
            "counterparty_name",
            "counterparty_account",
            "counterparty_bank",
            "remark",
            "channel",
            TX_COL_INCM_TX_TP_CD,
            TX_COL_GLOBAL,
            TX_COL_DEVICE_COLLECTED,
        ]
        raw_cols = [d[0] for d in RAW_TX_COLUMN_DEFS]
        insert_cols = base_cols + raw_cols
        raw_dict = row.get("raw") or {}
        insert_vals: list[Any] = [
            int(binding["device"]["person_id"]),
            row["tx_datetime"],
            row["tx_type"],
            row["currency"],
            row["tx_amount"],
            row["account_balance"],
            row["counterparty_name"],
            row["counterparty_account"],
            row["counterparty_bank"],
            row["remark"],
            row["channel"],
            row["incm_epn_tx_tp_cd"],
            row["global_busi_track_no"],
            1,
        ] + [raw_dict.get(col) for col in raw_cols]
        placeholders = ", ".join(["%s"] * len(insert_cols))
        quoted = ", ".join(f"`{c}`" for c in insert_cols)
        new_tx_id = await execute_insert(
            f"INSERT INTO `{TX_TABLE}` ({quoted}) VALUES ({placeholders})", tuple(insert_vals)
        )
        inserted_ids.append(int(new_tx_id))
        inserted += 1
    out: dict[str, Any] = {"ok": True, "inserted": inserted, "skipped": skipped, "total": len(rows)}
    # 利息同日去重等在管理端「重算余额」中处理（admin_recalculate_transaction_balances）。
    if inserted > 0:
        out["balance_recalc"] = await adjust_collected_transaction_balances_only(
            person_id_pre, inserted_ids, balance_snapshot
        )
    return out


TX_TYPE_KEYWORDS: dict[str, list[tuple[str, str]]] = {
    "1": [
        ("本行汇入", "1008"),
        ("他行汇入", "1008"),
        ("薪酬", "1001"),
        ("借款", "1002"),
        ("退款", "1003"),
        ("利息", "1004"),
        ("结息", "1004"),
        ("投资收益", "1005"),
        ("报销", "1006"),
        ("补贴", "1006"),
        ("社保公积金", "1007"),
        ("社保", "1007"),
        ("公积金", "1007"),
        ("他人转入", "1008"),
        ("转入", "1008"),
        ("现金收入", "1009"),
        ("现金", "1009"),
        ("其他收入", "1010"),
        ("贵金属", "1011"),
        ("公益", "1012"),
        ("手续费", "1013"),
        ("住房租房", "1014"),
        ("租房", "1014"),
    ],
    "2": [
        ("本行汇出", "2010"),
        ("跨行汇出", "2010"),
        ("购物", "2001"),
        ("消费", "2001"),
        ("交通", "2002"),
        ("出行", "2002"),
        ("娱乐", "2003"),
        ("教育", "2004"),
        ("生活", "2005"),
        ("医疗", "2006"),
        ("健康", "2006"),
        ("缴费", "2007"),
        ("充值", "2007"),
        ("转账给他人", "2010"),
        ("转出", "2010"),
        ("汇出", "2010"),
        ("还款", "2011"),
        ("其他支出", "2012"),
        ("餐饮", "2013"),
        ("美食", "2013"),
        ("服饰", "2014"),
        ("美容", "2014"),
        ("数码", "2015"),
        ("电器", "2015"),
        ("家居", "2016"),
        ("家装", "2016"),
        ("旅游", "2017"),
        ("酒店", "2017"),
        ("住房租房", "2018"),
        ("养车", "2019"),
        ("爱车", "2019"),
        ("现金支出", "2020"),
        ("ATM取款", "2020"),
        ("卡现金取", "2020"),
        ("取款", "2020"),
        ("红包", "2021"),
        ("手续费", "2022"),
        ("贵金属", "2023"),
        ("社保公积金", "2024"),
        ("党团工会", "2025"),
    ],
}


def resolve_incm_epn_tx_tp_cd(tx_type_text: str, tp_cd: str) -> str:
    normalized_tp_cd = "2" if str(tp_cd) == "2" else "1"
    text = "" if tx_type_text is None else str(tx_type_text).strip()
    if not text:
        return "2010" if normalized_tp_cd == "2" else "1008"
    candidates = TX_TYPE_KEYWORDS.get(normalized_tp_cd, [])
    for kw, code in candidates:
        if kw in text:
            return code
    return "2010" if normalized_tp_cd == "2" else "1008"


def resolve_incm_epn_tx_tp_cd_from_row(row: Optional[dict[str, Any]], tx_type_text: str, tp_cd: str) -> str:
    from_row = str(row.get("incmEpnTxTpCd") or "").strip() if row else ""
    if from_row:
        return from_row
    return resolve_incm_epn_tx_tp_cd(tx_type_text, tp_cd)


def incm_epn_tp_cd_from_record(r: dict[str, Any]) -> str:
    """从流水推断收支方向：1=收入，2=支出。仅当 dwFlagCode/incmEpnTpCd 为 1/2 时才采用。"""
    dw = str(first_non_empty_value(r.get("dwFlagCode"), r.get("dw_flag_code")) or "").strip()
    if dw == "2":
        return "2"
    if dw == "1":
        return "1"
    iep = str(first_non_empty_value(r.get("incmEpnTpCd"), r.get("incm_epn_tp_cd")) or "").strip()
    if iep == "2":
        return "2"
    if iep == "1":
        return "1"
    amt_raw = first_non_empty_value(r.get("txAmt"), r.get("tx_amt"), r.get("交易金额"), r.get("amount"), "0")
    try:
        amt = float(str(amt_raw).replace(",", "").strip() or "0")
    except Exception:
        amt = 0.0
    if amt < 0:
        return "2"
    return "1"


def record_amt_abs(r: dict[str, Any]) -> float:
    amt_raw = first_non_empty_value(r.get("txAmt"), r.get("tx_amt"), r.get("交易金额"), r.get("amount"), "0")
    try:
        return abs(float(str(amt_raw).replace(",", "").strip() or "0"))
    except Exception:
        return 0.0


def record_tx_date_ymd8(r: dict[str, Any]) -> str:
    return normalize_ymd8(r.get("txDate") or r.get("交易日期") or r.get("tx_datetime") or "")


def record_matches_incm_epn_month(r: dict[str, Any], incm_month: str) -> bool:
    if not incm_month or len(incm_month) != 6:
        return True
    tx_date = record_tx_date_ymd8(r)
    return bool(tx_date and tx_date[:6] == incm_month)


BANK_CODE_MAP = [
    ("工商", "102100099996", "中国工商银行"),
    ("农业", "103100000026", "中国农业银行"),
    ("中国银行", "104100000004", "中国银行"),
    ("建设", "105100000017", "中国建设银行"),
    ("交通", "301290000007", "交通银行"),
    ("邮储", "403100000004", "中国邮政储蓄银行"),
    ("邮政", "403100000004", "中国邮政储蓄银行"),
    ("招商", "308584000013", "招商银行"),
    ("中信", "302100011000", "中信银行"),
    ("平安", "307584007998", "平安银行"),
]


def resolve_bank_meta(bank_text: Any) -> dict[str, str]:
    text = "" if bank_text is None else str(bank_text).strip()
    if not text:
        return {"txOpsBankNm": "", "txOpsBankNo": ""}
    for kw, code, name in BANK_CODE_MAP:
        if kw in text:
            return {"txOpsBankNm": name, "txOpsBankNo": code}
    return {"txOpsBankNm": text, "txOpsBankNo": ""}


def first_non_empty_value(*values: Any) -> str:
    for value in values:
        if value is None:
            continue
        t = str(value).strip()
        if t:
            return t
    return ""


# 导入 Excel / 采集入库：摘要为这些类型时写入 outTxSriNo（31 位：交易日期 YYYYMMDD + 23 位数字）
_OUT_TX_SRI_NO_ELIGIBLE_SUMM_TYPES = frozenset({"快捷支付", "微信转账", "银联入账", "网联入账"})


def _random_digit_string(n: int) -> str:
    if n <= 0:
        return ""
    return "".join(str(secrets.randbelow(10)) for _ in range(n))


def _generate_out_tx_sri_no_31(*, ymd8: str) -> str:
    """对外交易流水号：前 8 位为交易日 YYYYMMDD，后 23 位为随机数字。"""
    y = normalize_ymd8(ymd8)
    if len(y) < 8:
        y = ymd_today_utc8()
    prefix = y[:8]
    return prefix + _random_digit_string(23)


def apply_out_tx_sri_no_for_import_or_collect(
    target: dict[str, Any],
    *,
    summ_for_check: str,
    ymd8: str,
    global_track_no: str,
) -> None:
    """
    与历史 mock-xhx 规则一致：四类摘要生成 31 位 outTxSriNo；否则仅在空时填入 global_track_no。
    """
    summ = str(summ_for_check or "").strip()
    if summ in _OUT_TX_SRI_NO_ELIGIBLE_SUMM_TYPES:
        tx_ymd = normalize_ymd8(ymd8)
        prefix_ok = tx_ymd[:8] if len(tx_ymd) >= 8 else ymd_today_utc8()[:8]
        existing = str(
            first_non_empty_value(target.get("outTxSriNo"), target.get("out_tx_sri_no")) or ""
        ).strip()
        keep_existing = (
            len(existing) == 31
            and existing.isdigit()
            and existing.startswith(prefix_ok)
        )
        target["outTxSriNo"] = (
            existing if keep_existing else _generate_out_tx_sri_no_31(ymd8=tx_ymd or prefix_ok)
        )
        return
    cur = target.get("outTxSriNo")
    if cur is None or (isinstance(cur, str) and not str(cur).strip()):
        g = str(global_track_no or "").strip()
        if g:
            target["outTxSriNo"] = g


def format_bkcd_mask(value: Any) -> str:
    """银行卡掩码：与真机常见格式一致，如 19 位卡号 -> 622180*********1744（前 6 + * + 后 4）。"""
    text = str(value if value is not None else "").strip()
    if not text:
        return ""
    n = len(text)
    if n <= 8:
        return text
    if n >= 10:
        return f"{text[:6]}{'*' * (n - 10)}{text[-4:]}"
    # n == 9：不宜 6+4（重叠），沿用 4 + * + 4
    return f"{text[:4]}{'*' * (n - 8)}{text[-4:]}"


def to_curr_code(value: Any) -> str:
    text = str(value if value is not None else "").strip()
    if not text:
        return "156"
    if re.fullmatch(r"\d+", text):
        return text
    if "人民币" in text or text.upper() == "CNY":
        return "156"
    return text


def build_tx_detail_data_from_record(
    record: dict[str, Any], basic_info: dict[str, Any], *, force_profile_medium: bool = False
) -> dict[str, Any]:
    r = record or {}
    tx_date_time = first_non_empty_value(
        r.get("tx_datetime"), r.get("交易日期"), r.get("txTime"), r.get("txDate")
    )
    dt = fmt_ymd_hms(tx_date_time)
    tx_date = first_non_empty_value(r.get("txDate"), dt["ymd"])
    tx_time = first_non_empty_value(r.get("txTime"), dt["ymdhms"])
    raw_amt = first_non_empty_value(r.get("交易金额"), r.get("tx_amount"), r.get("txAmt"))
    try:
        tx_amount_num = float(to_amt_str(raw_amt)) if str(raw_amt or "").strip() else 0.0
    except ValueError:
        tx_amount_num = 0.0
    inferred_dw = "2" if tx_amount_num < 0 else "1"
    dw_flag_code = first_non_empty_value(
        r.get("dwFlagCode"),
        r.get("dw_flag_code"),
        r.get("incm_epn_tp_cd"),
        r.get("incmEpnTpCd"),
        inferred_dw,
        "1",
    )
    incm_epn_tp_cd = first_non_empty_value(
        r.get("incm_epn_tp_cd"), r.get("incmEpnTpCd"), "2" if dw_flag_code == "2" else "1"
    )
    summ = first_non_empty_value(r.get("summ"), r.get("交易类型"), r.get("tx_type"))
    incm_epn_tx_tp_cd = first_non_empty_value(
        r.get(TX_COL_INCM_TX_TP_CD),
        r.get("incmEpnTxTpCd"),
        resolve_incm_epn_tx_tp_cd_from_row(r, summ, incm_epn_tp_cd),
    )
    profile_card = str(basic_info.get("卡号") or "").strip()
    if force_profile_medium and profile_card:
        medium_no = profile_card
    else:
        medium_no = first_non_empty_value(r.get("mediumNo"), basic_info.get("卡号"))
    profile_bkcd_mask = format_bkcd_mask(profile_card) if profile_card else ""
    tx_ops_accno = first_non_empty_value(r.get("txOpsAccno"), r.get("对手方账户"), r.get("counterparty_account"))
    tx_ops_name = first_non_empty_value(r.get("txOpsName"), r.get("对手方户名"), r.get("counterparty_name"))
    bank_from_text = first_non_empty_value(r.get("txOpsBankNm"), r.get("对手银行"), r.get("counterparty_bank"))
    bank_meta = resolve_bank_meta(bank_from_text)
    inferred_bank_name = infer_counterparty_bank_name(
        tx_ops_accno, r.get("txOpsBankNo"), bank_meta.get("txOpsBankNm") or bank_from_text
    )
    tx_ops_bank_nm = first_non_empty_value(
        r.get("txOpsBankNm"), bank_meta.get("txOpsBankNm"), inferred_bank_name, bank_from_text
    )
    tx_ops_bank_no = first_non_empty_value(r.get("txOpsBankNo"), bank_meta.get("txOpsBankNo"), "403100000004")
    acc_bal = first_non_empty_value(
        r.get("accBal"),
        to_amt_str(first_non_empty_value(r.get("账户余额"), r.get("account_balance"))),
    )
    chnl_kind_code_raw = first_non_empty_value(r.get("chnlKindCode"), r.get("交易方式"), r.get("channel"))
    chnl_kind_code = chnl_kind_code_raw if re.fullmatch(r"\d+", chnl_kind_code_raw) else "20"
    tx_remark = first_non_empty_value(r.get("txRemark"), r.get("附言"), r.get("remark"), "无附言")
    default_data: dict[str, Any] = {
        "accBal": "0.00",
        "accountBokList": [],
        "bkcdMask": "",
        "cardBokType": "2",
        "cashExgVatgCd": "2",
        "cashTranFlagCode": "1",
        "chnlKindCode": "20",
        "currCode": "156",
        "dtlSeqNo": "1",
        "dwFlagCode": "1",
        "flag": "1",
        "globalBusiTrackNo": "",
        "ibankFlag": "1",
        "incmEpnTpCd": "1",
        "incmEpnTxTpCd": "1008",
        "investProdtCdSets": "",
        "investProdtName": "",
        "mediumNo": "",
        "merDesc": "",
        "outTxSriNo": "",
        "reckinIncmEpnFlagCd": "1",
        "saccnoSeqNo": "1",
        "servNo": "",
        "summ": "",
        "tacctPtscrt": "",
        "transInMobileNo": "",
        "txAmt": "0.00",
        "txDate": "",
        "txHappAddr": "",
        "txOpsAccno": "",
        "txOpsBankNm": "",
        "txOpsBankNo": "",
        "txOpsName": "",
        "txRemark": "",
        "txTime": "",
    }
    out = {**default_data}
    if force_profile_medium:
        bkcd_mask_out = profile_bkcd_mask
    else:
        bkcd_mask_out = first_non_empty_value(
            r.get("bkcdMask"),
            format_bkcd_mask(medium_no),
            profile_bkcd_mask,
        )
    out.update(
        {
            "accBal": first_non_empty_value(r.get("accBal"), acc_bal),
            "bkcdMask": bkcd_mask_out,
            "cardBokType": first_non_empty_value(r.get("cardBokType"), default_data["cardBokType"]),
            "cashExgVatgCd": first_non_empty_value(r.get("cashExgVatgCd"), default_data["cashExgVatgCd"]),
            "cashTranFlagCode": first_non_empty_value(r.get("cashTranFlagCode"), default_data["cashTranFlagCode"]),
            "chnlKindCode": first_non_empty_value(r.get("chnlKindCode"), chnl_kind_code),
            "currCode": to_curr_code(first_non_empty_value(r.get("currCode"), r.get("交易币种"), r.get("currency"))),
            "dtlSeqNo": first_non_empty_value(
                r.get("dtlSeqNo"),
                r.get("tx_id"),
                "" if r.get("id") is None else str(r.get("id")),
                default_data["dtlSeqNo"],
            ),
            "dwFlagCode": first_non_empty_value(r.get("dwFlagCode"), r.get("dw_flag_code"), dw_flag_code),
            "flag": first_non_empty_value(r.get("flag"), default_data["flag"]),
            "globalBusiTrackNo": first_non_empty_value(
                r.get("globalBusiTrackNo"),
                r.get("global_busi_track_no"),
                r.get(TX_COL_GLOBAL),
            ),
            "ibankFlag": first_non_empty_value(r.get("ibankFlag"), default_data["ibankFlag"]),
            "incmEpnTpCd": first_non_empty_value(r.get("incm_epn_tp_cd"), r.get("incmEpnTpCd"), incm_epn_tp_cd),
            "incmEpnTxTpCd": first_non_empty_value(
                r.get(TX_COL_INCM_TX_TP_CD), r.get("incmEpnTxTpCd"), incm_epn_tx_tp_cd
            ),
            "investProdtCdSets": first_non_empty_value(r.get("investProdtCdSets"), default_data["investProdtCdSets"]),
            "investProdtName": first_non_empty_value(r.get("investProdtName"), default_data["investProdtName"]),
            "mediumNo": "O3e2410589f9c30c7a4cc9f8bc415035d44",
            "merDesc": first_non_empty_value(r.get("merDesc"), default_data["merDesc"]),
            "outTxSriNo": first_non_empty_value(r.get("outTxSriNo"), r.get("out_tx_sri_no"), default_data["outTxSriNo"]),
            "reckinIncmEpnFlagCd": first_non_empty_value(r.get("reckinIncmEpnFlagCd"), default_data["reckinIncmEpnFlagCd"]),
            "saccnoSeqNo": first_non_empty_value(r.get("saccnoSeqNo"), default_data["saccnoSeqNo"]),
            "servNo": first_non_empty_value(r.get("servNo"), default_data["servNo"]),
            "summ": first_non_empty_value(r.get("summ"), r.get("交易类型"), r.get("tx_type"), summ),
            "tacctPtscrt": first_non_empty_value(r.get("tacctPtscrt"), default_data["tacctPtscrt"]),
            "transInMobileNo": first_non_empty_value(r.get("transInMobileNo"), default_data["transInMobileNo"]),
            "txAmt": first_non_empty_value(r.get("txAmt"), f"{abs(tx_amount_num):.2f}"),
            "txDate": first_non_empty_value(r.get("txDate"), tx_date),
            "txHappAddr": first_non_empty_value(r.get("txHappAddr"), default_data["txHappAddr"]),
            "txOpsAccno": first_non_empty_value(
                r.get("txOpsAccno"), r.get("对手方账户"), r.get("counterparty_account"), tx_ops_accno
            ),
            "txOpsBankNm": first_non_empty_value(
                r.get("txOpsBankNm"), r.get("对手银行"), r.get("counterparty_bank"), tx_ops_bank_nm
            ),
            "txOpsBankNo": first_non_empty_value(r.get("txOpsBankNo"), tx_ops_bank_no),
            "txOpsName": first_non_empty_value(
                r.get("txOpsName"), r.get("对手方户名"), r.get("counterparty_name"), tx_ops_name
            ),
            "txRemark": first_non_empty_value(r.get("txRemark"), r.get("附言"), r.get("remark"), tx_remark),
            "txTime": first_non_empty_value(r.get("txTime"), tx_time),
        }
    )
    return out


async def build_tx_detail_response(req_msg_id: str, did: str, query: dict[str, Any]) -> dict[str, Any]:
    target_track_no = str(query.get("globalBusiTrackNo") or "").strip()
    if not target_track_no:
        return {
            "code": "000016",
            "msg": "暂时无法处理您的请求，请返回或退出后重试",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    direct = await read_tx_by_did_and_global_track_no(did, target_track_no)
    rr = direct.get("transaction")
    if not rr:
        return {
            "code": "000016",
            "msg": "暂时无法处理您的请求，请返回或退出后重试",
            "showType": "2",
            "reqMsgId": req_msg_id or "",
        }
    detail_data = build_tx_detail_data_from_record(rr, direct.get("basic_info") or {})
    return {
        "code": "000000",
        "data": detail_data,
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


def ymd_today_utc8() -> str:
    d = now_utc8()
    return f"{d.year}{str(d.month).zfill(2)}{str(d.day).zfill(2)}"


async def build_history_transaction_detail_apply_response(
    req_msg_id: str, did: str, query: dict[str, Any]
) -> dict[str, Any]:
    cust_no = str(query.get("custNo") or "").strip() or "843942355"
    dline = normalize_ymd8(query.get("dlineDate")) or ymd_today_utc8()
    gen_file_name = f"14001990000_SJYHLSJYMXSX_0000_{dline}_A_0001_5152.xml{cust_no}"
    return {
        "code": "000000",
        "data": {"fileRecordHeadNum": "6", "genFileName": gen_file_name, "ovrnFlag": "0"},
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


def _history_apply_email_param_enabled() -> bool:
    return os.environ.get("MOCK_HISTORY_TX_APPLY_EMAIL", "1").strip().lower() not in (
        "0",
        "false",
        "no",
        "off",
    )


def _looks_like_email(addr: str) -> bool:
    s = (addr or "").strip()
    if len(s) < 5 or "@" not in s:
        return False
    local, _, domain = s.partition("@")
    return bool(local) and bool(domain) and "." in domain


def _range_compact_from_history_query(query: dict[str, Any]) -> str:
    bd = normalize_ymd8(query.get("beginDate"))
    dd = normalize_ymd8(query.get("dlineDate")) or ymd_today_utc8()
    if not bd:
        bd = dd
    if bd > dd:
        bd, dd = dd, bd
    return f"{bd}-{dd}"


async def dispatch_history_apply_email_if_requested(
    *,
    to_email: str,
    person_id: int,
    customer_name: str,
    query: dict[str, Any],
) -> None:
    """在历史明细申请 mock 成功后发送 ZIP 邮件（阻塞 IO 放到线程池）。"""
    if not _history_apply_email_param_enabled():
        print("[mock-api] history apply mail skipped (MOCK_HISTORY_TX_APPLY_EMAIL disabled)")
        return
    try:
        import history_tx_mail  # 惰性导入，避免未安装 PDF 依赖时无法启动 mock 服务
    except ImportError as e:
        print(f"[mock-api] history apply mail skipped (import history_tx_mail): {e}")
        return

    apply_at = now_utc8()
    apply_date = apply_at.strftime("%Y%m%d")
    range_compact = _range_compact_from_history_query(query)
    draw_no = str(query.get("drawNo") or "").strip()
    # 约定：history_apply_mail_records.draw_no 就是 ZIP 解压密码；若未传则自动生成 6 位数字
    if not draw_no:
        import secrets

        draw_no = "".join(str(secrets.randbelow(10)) for _ in range(6))
        query = dict(query)
        query["drawNo"] = draw_no

    # 邮件主题序列号：当天发送的次数（001 起），与 drawNo（ZIP 密码）解耦
    mail_seq_no = await _next_history_mail_seq_no_for_date(apply_date)

    # 先写入发送记录，再发送邮件
    prow = await query_one(
        f"SELECT card_no, `{PROFILE_COL_MAIL_SEND_MODE}` AS mail_send_mode FROM `{PROFILE_TABLE}` WHERE id=%s LIMIT 1",
        (int(person_id),),
    )
    card_no = str((prow or {}).get("card_no") or "").strip()
    mail_send_mode = normalize_mail_send_mode((prow or {}).get("mail_send_mode"))
    await insert_history_apply_mail_record(
        person_id=person_id,
        medium_no=card_no or None,
        to_email=to_email,
        query=query,
        apply_at=apply_at,
        mail_seq_no=mail_seq_no,
    )

    def _run() -> None:
        history_tx_mail.send_for_person_psbc(
            to_email,
            person_id,
            customer_name=(customer_name or None),
            apply_at=apply_at,
            range_compact=range_compact,
            mail_seq_date=apply_at.strftime("%Y%m%d"),
            mail_seq_no=mail_seq_no,
            zip_password=draw_no,
            mail_send_mode=mail_send_mode,
        )

    await asyncio.to_thread(_run)


# 真机「收支分析」列表接口 detailList 单条仅含以下 camelCase 字段（与网银 JSON 一致）。
XHX_DETAIL_LIST_KEYS: tuple[str, ...] = (
    "accBal",
    "bkcdMask",
    "cashExgVatgCd",
    "chnlKindCode",
    "currCode",
    "dtlSeqNo",
    "dwFlagCode",
    "globalBusiTrackNo",
    "ibankFlag",
    "incmEpnTpCd",
    "incmEpnTxTpCd",
    "investProdtCdSets",
    "investProdtName",
    "mediumNo",
    "merDesc",
    "persInnerAccno",
    "randomAssignNo",
    "reckinIncmEpnFlagCd",
    "saccnoSeqNo",
    "servNo",
    "subtxNo",
    "summ",
    "transInMobileNo",
    "txAmt",
    "txDate",
    "txOpsAccno",
    "txOpsName",
    "txRemark",
    "txTime",
)


def build_xhx_detail_list_item(record: dict[str, Any], basic_info: dict[str, Any], dtl_seq_no: int) -> dict[str, str]:
    # mock-xhx 返回须与绑定资料卡号一致，避免库内流水仍带旧卡号/掩码
    full = build_tx_detail_data_from_record(record, basic_info, force_profile_medium=True)
    full["dtlSeqNo"] = str(dtl_seq_no)
    out: dict[str, str] = {}
    for k in XHX_DETAIL_LIST_KEYS:
        v = full.get(k)
        if v is None:
            out[k] = ""
        else:
            out[k] = str(v)
    return out


def _parse_tx_tp_cd_set(query: dict[str, Any]) -> set[str]:
    out: set[str] = set()
    raw = query.get("txTpCdList")
    if raw is None:
        return out
    if isinstance(raw, list):
        for item in raw:
            code = str(item or "").strip()
            if code:
                out.add(code)
        return out
    text = str(raw).strip()
    if not text:
        return out
    for item in text.split(","):
        code = str(item or "").strip()
        if code:
            out.add(code)
    return out


async def build_xhx_response(req_msg_id: str, did: str, query: dict[str, Any]) -> dict[str, Any]:
    """
    query 可由 POST /api/mock-xhx 的 JSON body 拼装并传入（仅 URL ?did= 为短参数），其中：
      - hostRspPlain（可选）：Hook 传来的原网关 SM4 解密后明文（JSON 字符串）。
    """
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        raise RuntimeError(binding["error"])
    device = binding["device"]
    basic_info = {
        "户名": device.get("person_name") or "",
        "卡号": device.get("card_no") or "",
        "身份证号": device.get("id_card") or "",
        "电话号码": device.get("phone_no") or "",
        "印章流水": device.get("stamp_no") or "",
        "抬头": device.get("title") or "",
        "profile_balance": device.get("profile_balance"),
    }
    recs = await read_raw_mock_data_by_did(did)
    begin_date = normalize_ymd8(query.get("beginDate"))
    dline_date = normalize_ymd8(query.get("dlineDate"))
    tx_tp_cd_set = _parse_tx_tp_cd_set(query)
    ebank_qry_type_flag_cd = str(query.get("ebankQryTypeFlagCd") or "").strip()
    if begin_date and dline_date and begin_date > dline_date:
        begin_date, dline_date = dline_date, begin_date

    # 自定义查询：不按交易月份拆分，直接汇总到单一月份节点
    custom_month: str | None = None
    if ebank_qry_type_flag_cd == "1":
        anchor = dline_date or begin_date or ymd_today_utc8()
        custom_month = anchor[:6] if anchor else None
    month_map: dict[str, dict[str, Any]] = {}
    for r0 in recs:
        r = r0 or {}
        tx_date = normalize_ymd8(r.get("txDate") or r.get("交易日期") or r.get("tx_datetime") or "")
        if not tx_date:
            continue
        if begin_date and tx_date < begin_date:
            continue
        if dline_date and tx_date > dline_date:
            continue
        tx_tp_cd = str(r.get("incmEpnTxTpCd") or r.get("incm_epn_tx_tp_cd") or "").strip()
        if tx_tp_cd_set and (not tx_tp_cd or tx_tp_cd not in tx_tp_cd_set):
            continue
        month = custom_month or tx_date[:6]
        if month not in month_map:
            month_map[month] = {"_rows": [], "expnTotAmt": 0.0, "incomeTotalAmt": 0.0}
        month_map[month]["_rows"].append(r)
    months = sorted(month_map.keys(), reverse=True)
    total = 0
    field_item_info = []
    for mm in months:
        node = month_map[mm]
        rows_m = node["_rows"]
        n = len(rows_m)
        details: list[dict[str, str]] = []
        for i, r in enumerate(rows_m):
            seq = n - i
            d = build_xhx_detail_list_item(r, basic_info, seq)
            details.append(d)
            # 汇总金额以原始记录为准，避免 detailList 中 txAmt 格式/缺失导致统计偏差
            amt_raw = first_non_empty_value(r.get("txAmt"), r.get("tx_amt"), r.get("交易金额"), r.get("amount"), "0")
            try:
                amount_abs = abs(float(str(amt_raw).replace(",", "").strip() or "0"))
            except Exception:
                amount_abs = float("nan")
            if not math.isnan(amount_abs) and amount_abs > 0:
                dw = str(first_non_empty_value(r.get("dwFlagCode"), r.get("incmEpnTpCd"), d.get("dwFlagCode"), "1") or "1").strip()
                if dw == "2":
                    node["expnTotAmt"] += amount_abs
                else:
                    node["incomeTotalAmt"] += amount_abs
        total += len(details)
        field_item_info.append(
            {
                "detailList": details,
                "expnTotAmt": f"{node['expnTotAmt']:.2f}",
                "incmEpnMonth": mm,
                "incomeTotalAmt": f"{node['incomeTotalAmt']:.2f}",
            }
        )
    now = datetime.now()
    qry_time = (
        f"{now.year}{str(now.month).zfill(2)}{str(now.day).zfill(2)}"
        f"{str(now.hour).zfill(2)}{str(now.minute).zfill(2)}{str(now.second).zfill(2)}"
    )
    return {
        "code": "000000",
        "data": {
            "curQryReturnNum": str(total),
            "fieldItemInfo": field_item_info,
            "haveNextDataFlag": "0",
            "qryResultTnum": str(total),
            "qryTime": qry_time,
        },
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


# txTrsfQry/T080233：来往记录查询时，网关用内部户名/账号标识平台（如微信 1000050201）
_WALLET_TRSF_QUERY_ACCNO: dict[str, frozenset[str]] = {
    "1000050201": frozenset({"微信转账"}),
}


def _normalize_account_digits(s: str) -> str:
    return re.sub(r"\D", "", str(s or ""))


def _accounts_match_for_trsf(want: str, have: str) -> bool:
    w = _normalize_account_digits(want)
    h = _normalize_account_digits(have)
    if not w or not h:
        return False
    if w == h:
        return True
    shorter, longer = (w, h) if len(w) <= len(h) else (h, w)
    return len(shorter) >= 8 and longer.endswith(shorter)


def _names_match_for_trsf(want: str, have: str) -> bool:
    w = str(want or "").strip()
    h = str(have or "").strip()
    if not w or not h:
        return False
    return w == h or w in h or h in w


def _is_wallet_platform_tx_record(r: dict[str, Any]) -> bool:
    summ = str(first_non_empty_value(r.get("summ"), r.get("交易类型"), r.get("tx_type")) or "").strip()
    if summ in _OUT_TX_SRI_NO_ELIGIBLE_SUMM_TYPES:
        return True
    cp = str(first_non_empty_value(r.get("counterparty_name"), r.get("txOpsName"), r.get("merDesc")) or "")
    if "财付通" in cp or "支付宝" in cp:
        return True
    chnl = str(first_non_empty_value(r.get("chnlKindCode"), r.get("交易方式"), r.get("channel")) or "")
    return chnl == "30"


def _record_matches_tx_trsf_query(
    r: dict[str, Any], *, query_accno: str, query_name: str
) -> bool:
    """来往记录筛选：有 txOpsAccno 时只按对手方卡号/账号匹配，txOpsName 可选。"""
    query_accno = str(query_accno or "").strip()
    query_name = str(query_name or "").strip()
    if not query_accno and not query_name:
        return False

    rec_acc = first_non_empty_value(r.get("txOpsAccno"), r.get("对手方账户"), r.get("counterparty_account"))
    rec_name = first_non_empty_value(r.get("txOpsName"), r.get("对手方户名"), r.get("counterparty_name"))
    summ = str(first_non_empty_value(r.get("summ"), r.get("交易类型"), r.get("tx_type")) or "").strip()

    wallet_names = _WALLET_TRSF_QUERY_ACCNO.get(query_accno)
    if wallet_names is not None:
        name_ok = (not query_name) or query_name in wallet_names or _names_match_for_trsf(query_name, "微信转账")
        if not name_ok:
            return False
        return _is_wallet_platform_tx_record(r) or summ in _OUT_TX_SRI_NO_ELIGIBLE_SUMM_TYPES

    if query_accno:
        return _accounts_match_for_trsf(query_accno, rec_acc)

    return _names_match_for_trsf(query_name, rec_name) or _names_match_for_trsf(query_name, summ)


def build_empty_tx_trsf_qry_response(req_msg_id: str) -> dict[str, Any]:
    return {
        "code": "000000",
        "data": {
            "curQryReturnNum": "0",
            "fieldItemInfo": [],
            "haveNextDataFlag": "0",
        },
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


def pick_req_msg_id(request: web.Request, body: Optional[dict[str, Any]] = None) -> str:
    """从 header / query / POST 明文（含 tokenInfo.reqMsgId）解析 reqMsgId。"""
    for key in ("x-req-msg-id", "reqMsgId", "ReqMsgId"):
        h = str(request.headers.get(key) or "").strip()
        if h:
            return h
    q = str(request.query.get("reqMsgId") or "").strip()
    if q:
        return q
    b = body or {}
    top = str(b.get("reqMsgId") or "").strip()
    if top:
        return top
    ti = b.get("tokenInfo")
    if isinstance(ti, dict):
        nested = str(ti.get("reqMsgId") or "").strip()
        if nested:
            return nested
    return ""


_TX_TRSF_QRY_FIELD_KEYS: tuple[str, ...] = (
    "beginDate",
    "dlineDate",
    "txOpsAccno",
    "txOpsName",
    "bgnIndexNo",
    "curQryReqNum",
)


def _try_json_dict_from_text(text: Any) -> Optional[dict[str, Any]]:
    s = str(text or "").strip()
    if not s or s[0] not in ("{", "["):
        return None
    try:
        obj = json.loads(s)
    except json.JSONDecodeError:
        return None
    return obj if isinstance(obj, dict) else None


def _decode_b64_json_maybe(raw: Any) -> Optional[dict[str, Any]]:
    s = str(raw or "").strip()
    if not s:
        return None
    try:
        text = base64.b64decode(s).decode("utf-8", errors="replace")
    except Exception:
        return None
    return _try_json_dict_from_text(text)


def _pick_tx_trsf_qry_fields_from_obj(obj: Optional[dict[str, Any]]) -> dict[str, str]:
    if not obj or not isinstance(obj, dict):
        return {}
    out: dict[str, str] = {}
    alias_map: dict[str, tuple[str, ...]] = {
        "txOpsAccno": ("txOpsAccno", "counterparty_account", "对手方账户", "txOpsAccNo"),
        "txOpsName": ("txOpsName", "counterparty_name", "对手方户名", "txOpsNm"),
        "beginDate": ("beginDate",),
        "dlineDate": ("dlineDate",),
        "bgnIndexNo": ("bgnIndexNo",),
        "curQryReqNum": ("curQryReqNum",),
    }
    for key in _TX_TRSF_QRY_FIELD_KEYS:
        for alias in alias_map.get(key, (key,)):
            v = obj.get(alias)
            if v is None:
                continue
            s = str(v).strip()
            if s:
                out[key] = s
                break
    return out


def _merge_tx_trsf_qry_fields(*parts: Optional[dict[str, str]]) -> dict[str, str]:
    merged: dict[str, str] = {}
    for part in parts:
        if not part:
            continue
        for k, v in part.items():
            if v:
                merged[k] = v
    return merged


def tx_trsf_qry_payload_from_request(request: web.Request, body: dict[str, Any]) -> dict[str, str]:
    """
    解析 txTrsfQry/T080233 查询参数。与 mock-tx-detail / mock-tx-incm-epn-analy-sum 一致：
    优先读 URL query，再读 POST body 顶层，并从 hostReqPlain / reqPlain / hostReqPayloadB64 等补充。
    """
    b = body if isinstance(body, dict) else {}

    from_query = _pick_tx_trsf_qry_fields_from_obj({k: request.query.get(k) for k in _TX_TRSF_QRY_FIELD_KEYS})
    from_body = _pick_tx_trsf_qry_fields_from_obj(b)

    plain_sources: list[Optional[dict[str, Any]]] = []
    for key in ("hostReqPlain", "reqPlain", "plainRequest", "hostReqPlainText"):
        plain_sources.append(_try_json_dict_from_text(b.get(key)))
        plain_sources.append(_try_json_dict_from_text(request.query.get(key)))
    for b64_key in ("hostReqPayloadB64", "reqPayloadB64", "hostPayloadB64"):
        for src in (b.get(b64_key), request.query.get(b64_key)):
            decoded = _decode_b64_json_maybe(src)
            if decoded and ("txOpsAccno" in decoded or "txOpsName" in decoded or "beginDate" in decoded):
                plain_sources.append(decoded)

    from_plain: dict[str, str] = {}
    for obj in plain_sources:
        from_plain = _merge_tx_trsf_qry_fields(from_plain, _pick_tx_trsf_qry_fields_from_obj(obj))

    merged = _merge_tx_trsf_qry_fields(from_query, from_body, from_plain)
    return {
        "beginDate": merged.get("beginDate") or "",
        "dlineDate": merged.get("dlineDate") or "",
        "txOpsAccno": merged.get("txOpsAccno") or "",
        "txOpsName": merged.get("txOpsName") or "",
        "bgnIndexNo": merged.get("bgnIndexNo") or "1",
        "curQryReqNum": merged.get("curQryReqNum") or "30",
    }


async def serve_tx_trsf_qry_request(request: web.Request, body_any: dict[str, Any]) -> web.Response:
    """来往记录 mock：供 /sn13/.../T080233 与 /api/mock-tx-trsf* 共用。"""
    req_msg_id = pick_req_msg_id(request, body_any)
    did = pick_did(request)
    if not did:
        return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
    device_any = await load_device_context_any(did)
    if device_any and to_bool_int(device_any.get("is_active")) != 1:
        empty = build_empty_tx_trsf_qry_response(str(req_msg_id))
        print("服务端响应(txTrsfQry 设备未开启)===============>", empty)
        return json_response(empty)
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        return json_response({"ok": False, "error": binding["error"]}, int(binding["status"]))
    q_payload = tx_trsf_qry_payload_from_request(request, body_any)
    if not q_payload.get("txOpsAccno"):
        print(
            "[mock-api] txTrsfQry 缺少 txOpsAccno；"
            "请转发对手方卡号/账号（txOpsName 可不传，仅按卡号检索）"
        )
    data = await build_tx_trsf_qry_response(str(req_msg_id), did, q_payload)
    print(
        "服务端响应(txTrsfQry)===============>",
        {
            "curQryReturnNum": (data.get("data") or {}).get("curQryReturnNum"),
            "txOpsAccno": q_payload.get("txOpsAccno"),
            "txOpsName": q_payload.get("txOpsName"),
            "beginDate": q_payload.get("beginDate"),
            "dlineDate": q_payload.get("dlineDate"),
            "body_keys": sorted(b for b in (body_any or {}).keys()),
            "query_keys": sorted(request.query.keys()),
        },
    )
    return json_response(data)


async def build_tx_trsf_qry_response(req_msg_id: str, did: str, query: dict[str, Any]) -> dict[str, Any]:
    """对齐 app 端 txTrsfQry/T080233：按对手方卡号/账号 + 日期区间查询来往明细（户名可选）。"""
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        raise RuntimeError(binding["error"])
    device = binding["device"]
    basic_info = {
        "户名": device.get("person_name") or "",
        "卡号": device.get("card_no") or "",
        "身份证号": device.get("id_card") or "",
        "电话号码": device.get("phone_no") or "",
        "印章流水": device.get("stamp_no") or "",
        "抬头": device.get("title") or "",
        "profile_balance": device.get("profile_balance"),
    }
    recs = await read_raw_mock_data_by_did(did)
    begin_date = normalize_ymd8(query.get("beginDate"))
    dline_date = normalize_ymd8(query.get("dlineDate"))
    query_accno = str(query.get("txOpsAccno") or "").strip()
    query_name = str(query.get("txOpsName") or "").strip()
    if begin_date and dline_date and begin_date > dline_date:
        begin_date, dline_date = dline_date, begin_date

    try:
        bgn_index = max(1, int(str(query.get("bgnIndexNo") or "1").strip() or "1"))
    except ValueError:
        bgn_index = 1
    try:
        page_size = max(1, min(500, int(str(query.get("curQryReqNum") or "30").strip() or "30")))
    except ValueError:
        page_size = 30

    matched: list[dict[str, Any]] = []
    for r0 in recs:
        r = r0 or {}
        if not _record_matches_tx_trsf_query(r, query_accno=query_accno, query_name=query_name):
            continue
        tx_date = normalize_ymd8(r.get("txDate") or r.get("交易日期") or r.get("tx_datetime") or "")
        if not tx_date:
            continue
        if begin_date and tx_date < begin_date:
            continue
        if dline_date and tx_date > dline_date:
            continue
        matched.append(r)

    total_matched = len(matched)
    start = bgn_index - 1
    page_rows = matched[start : start + page_size]
    have_next = "1" if start + page_size < total_matched else "0"

    month_map: dict[str, dict[str, Any]] = {}
    for r in page_rows:
        tx_date = normalize_ymd8(r.get("txDate") or r.get("交易日期") or r.get("tx_datetime") or "")
        mm = tx_date[:6] if tx_date else ""
        if not mm:
            continue
        if mm not in month_map:
            month_map[mm] = {"_rows": [], "expnTotAmt": 0.0, "incomeTotalAmt": 0.0}
        month_map[mm]["_rows"].append(r)

    months = sorted(month_map.keys(), reverse=True)
    field_item_info: list[dict[str, Any]] = []
    page_count = 0
    for mm in months:
        node = month_map[mm]
        rows_m = node["_rows"]
        n = len(rows_m)
        details: list[dict[str, str]] = []
        for i, r in enumerate(rows_m):
            seq = n - i
            d = build_xhx_detail_list_item(r, basic_info, seq)
            if query_accno in _WALLET_TRSF_QUERY_ACCNO:
                d["txOpsAccno"] = query_accno
                if query_name:
                    d["txOpsName"] = query_name
            details.append(d)
            amt_raw = first_non_empty_value(r.get("txAmt"), r.get("tx_amt"), r.get("交易金额"), r.get("amount"), "0")
            try:
                amount_abs = abs(float(str(amt_raw).replace(",", "").strip() or "0"))
            except Exception:
                amount_abs = float("nan")
            if not math.isnan(amount_abs) and amount_abs > 0:
                dw = str(
                    first_non_empty_value(r.get("dwFlagCode"), r.get("incmEpnTpCd"), d.get("dwFlagCode"), "1") or "1"
                ).strip()
                if dw == "2":
                    node["expnTotAmt"] += amount_abs
                else:
                    node["incomeTotalAmt"] += amount_abs
        page_count += len(details)
        field_item_info.append(
            {
                "detailList": details,
                "expnTotAmt": f"{node['expnTotAmt']:.2f}",
                "incmEpnMonth": mm,
                "incomeTotalAmt": f"{node['incomeTotalAmt']:.2f}",
                "tnum": str(len(details)),
            }
        )

    return {
        "code": "000000",
        "data": {
            "curQryReturnNum": str(page_count),
            "fieldItemInfo": field_item_info,
            "haveNextDataFlag": have_next,
        },
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


def _parse_host_incm_epn_field_item_axis(fi: Any) -> tuple[list[str], str]:
    """
    从网关 T080239 data.fieldItemInfo 判断统计维度。
    返回 (有序键列表, 'month'|'year')：month 键为 YYYYMM，year 键为 YYYY。
    若同时存在有效月份项，优先按月；否则若存在年份项则按年。
    """
    if not isinstance(fi, list):
        return [], "month"
    month_keys: list[str] = []
    year_keys: list[str] = []
    for it in fi:
        if not isinstance(it, dict):
            continue
        m = str(it.get("incmEpnMonth") or "").strip()
        y = str(it.get("incmEpnYear") or "").strip()
        if len(m) == 6 and m.isdigit():
            if m not in month_keys:
                month_keys.append(m)
        elif len(y) == 4 and y.isdigit():
            if y not in year_keys:
                year_keys.append(y)
    if month_keys:
        return month_keys, "month"
    if year_keys:
        return year_keys, "year"
    return [], "month"


def _infer_selected_incm_epn_year_from_field_items(fi: list[Any]) -> str:
    """与年度 fieldItemInfo 对齐：自后向前取首条 tnum>0 的 incmEpnYear；否则取列表中最后一条有效年份。"""
    if not isinstance(fi, list):
        return ""
    for it in reversed(fi):
        if not isinstance(it, dict):
            continue
        yy = str(it.get("incmEpnYear") or "").strip()
        if len(yy) != 4 or not yy.isdigit():
            continue
        try:
            if int(str(it.get("tnum") or "0").strip() or "0") > 0:
                return yy
        except ValueError:
            continue
    for it in reversed(fi):
        if not isinstance(it, dict):
            continue
        yy = str(it.get("incmEpnYear") or "").strip()
        if len(yy) == 4 and yy.isdigit():
            return yy
    return ""


async def build_incm_epn_analy_sum_response(req_msg_id: str, did: str, query: dict[str, Any]) -> dict[str, Any]:
    """
    对齐 app 端 `qryIncmEpnAnalySum/T080239` 的常见返回结构：
    - data: { fieldItemInfo: [ {incmEpnMonth, incmEpnYear, tnum, totalAmt}, ... ] }
    按 raw mock 聚合：host 模板为「按年」（仅有 incmEpnYear）时按自然年与模板年份序列对齐；
    否则按月份（YYYYMM）；无模板时默认按月（无月数据再按年）。
    """
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        raise RuntimeError(binding["error"])
    recs = await read_raw_mock_data_by_did(did)
    begin_date = normalize_ymd8(query.get("beginDate"))
    dline_date = normalize_ymd8(query.get("dlineDate"))
    want_tp = str(query.get("incmEpnTpCd") or "").strip()  # 1/2
    host_rsp_plain = str(query.get("hostRspPlain") or "")
    if begin_date and dline_date and begin_date > dline_date:
        begin_date, dline_date = dline_date, begin_date

    month_map: dict[str, dict[str, Any]] = {}
    year_map: dict[str, dict[str, Any]] = {}
    for r0 in recs:
        r = r0 or {}
        tx_date = normalize_ymd8(r.get("txDate") or r.get("交易日期") or r.get("tx_datetime") or "")
        if not tx_date:
            continue
        if begin_date and tx_date < begin_date:
            continue
        if dline_date and tx_date > dline_date:
            continue

        # 进/出：优先使用 dwFlagCode，其次 incmEpnTpCd；2 表示支出，其余按收入处理
        dw = str(first_non_empty_value(r.get("dwFlagCode"), r.get("incmEpnTpCd"), "1") or "1").strip()
        incm_epn_tp_cd = "2" if dw == "2" else "1"
        if want_tp in ("1", "2") and incm_epn_tp_cd != want_tp:
            continue

        mm = tx_date[:6]
        yy = tx_date[:4]
        if mm not in month_map:
            month_map[mm] = {"tnum": 0, "total": 0.0}
        if yy not in year_map:
            year_map[yy] = {"tnum": 0, "total": 0.0}
        amt_raw = first_non_empty_value(r.get("txAmt"), r.get("tx_amt"), r.get("交易金额"), r.get("amount"), "0")
        try:
            amt = float(str(amt_raw).replace(",", "").strip() or "0")
        except Exception:
            amt = 0.0
        amt_abs = abs(amt)
        if amt_abs <= 0:
            continue
        month_map[mm]["tnum"] += 1
        month_map[mm]["total"] += amt_abs
        year_map[yy]["tnum"] += 1
        year_map[yy]["total"] += amt_abs

    host_fi: list[Any] = []
    if host_rsp_plain:
        try:
            obj = json.loads(host_rsp_plain)
            host_fi = ((obj or {}).get("data") or {}).get("fieldItemInfo") or []
        except Exception:
            host_fi = []
    template_keys, axis = _parse_host_incm_epn_field_item_axis(host_fi)

    if template_keys:
        keys = template_keys
    elif axis == "year" and year_map:
        keys = sorted(year_map.keys())
    else:
        keys = sorted(month_map.keys()) if month_map else sorted(year_map.keys())
        axis = "month" if month_map else "year"

    def _fmt_amt(v: float) -> str:
        return "0" if abs(v) < 0.0005 else f"{v:.2f}"

    field_item_info: list[dict[str, Any]] = []
    if axis == "year":
        for yy in keys:
            node = year_map.get(yy) or {"tnum": 0, "total": 0.0}
            field_item_info.append(
                {
                    "incmEpnMonth": "",
                    "incmEpnYear": yy,
                    "tnum": str(int(node["tnum"])),
                    "totalAmt": _fmt_amt(float(node["total"])),
                }
            )
    else:
        for mm in keys:
            node = month_map.get(mm) or {"tnum": 0, "total": 0.0}
            field_item_info.append(
                {
                    "incmEpnMonth": mm,
                    "incmEpnYear": "",
                    "tnum": str(int(node["tnum"])),
                    "totalAmt": _fmt_amt(float(node["total"])),
                }
            )
    return {
        "code": "000000",
        "data": {"fieldItemInfo": field_item_info},
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


def _prev_incm_epn_month(yyyymm: str) -> str:
    """YYYYMM → 上一自然月 YYYYMM；非法则返回空串。"""
    s = str(yyyymm or "").strip()
    if len(s) != 6 or not s.isdigit():
        return ""
    y = int(s[:4])
    m = int(s[4:6])
    if m < 1 or m > 12:
        return ""
    if m == 1:
        return f"{y - 1}12"
    return f"{y}{m - 1:02d}"


def _iter_yyyymm_between(begin_date: str, dline_date: str) -> list[str]:
    """YYYYMMDD 闭区间 → 覆盖的自然月 YYYYMM 列表（升序）。"""
    bd = normalize_ymd8(begin_date)
    dd = normalize_ymd8(dline_date)
    if len(bd) < 6 or len(dd) < 6:
        return []
    y, m = int(bd[:4]), int(bd[4:6])
    ey, em = int(dd[:4]), int(dd[4:6])
    if m < 1 or m > 12 or em < 1 or em > 12:
        return []
    out: list[str] = []
    cy, cm = y, m
    while (cy, cm) <= (ey, em):
        out.append(f"{cy}{cm:02d}")
        if cm == 12:
            cy += 1
            cm = 1
        else:
            cm += 1
    return out


def _iter_yyyy_between(begin_date: str, dline_date: str) -> list[str]:
    """YYYYMMDD 闭区间 → 覆盖的自然年 YYYY 列表（升序）。"""
    bd = normalize_ymd8(begin_date)
    dd = normalize_ymd8(dline_date)
    if len(bd) < 4 or len(dd) < 4:
        return []
    sy, ey = int(bd[:4]), int(dd[:4])
    if sy > ey:
        sy, ey = ey, sy
    return [str(y) for y in range(sy, ey + 1)]


def _sum_raw_incm_epn_abs_amt(
    recs: list[Any],
    *,
    want_tp: str,
    year_scope: bool,
    incm_epn_year: str,
    incm_month: str,
    begin_date: str,
    dline_date: str,
) -> tuple[float, dict[tuple[str, str], dict[str, Any]]]:
    """
    按范围聚合 raw 流水：返回 (总额绝对值, agg[(tp,txTpCd)]->{sum,num})。
    总额统计不要求 incmEpnTxTpCd；分类聚合时缺失类型则按 summ/交易类型推断。
    """
    agg: dict[tuple[str, str], dict[str, Any]] = {}
    total_amt = 0.0
    for r0 in recs:
        r = r0 or {}
        tx_date = normalize_ymd8(r.get("txDate") or r.get("交易日期") or r.get("tx_datetime") or "")
        scope_need_date = bool(begin_date or dline_date or incm_month or year_scope)
        if scope_need_date and not tx_date:
            continue
        if year_scope:
            if len(tx_date) < 4 or tx_date[:4] != incm_epn_year:
                continue
        elif incm_month and tx_date and len(incm_month) == 6 and tx_date[:6] != incm_month:
            continue
        if begin_date and tx_date and tx_date < begin_date:
            continue
        if dline_date and tx_date and tx_date > dline_date:
            continue

        incm_epn_tp_cd = incm_epn_tp_cd_from_record(r)
        if want_tp in ("1", "2") and incm_epn_tp_cd != want_tp:
            continue

        amt_abs = record_amt_abs(r)
        if amt_abs <= 0:
            continue

        total_amt += amt_abs

        tx_tp_cd = str(r.get("incmEpnTxTpCd") or r.get("incm_epn_tx_tp_cd") or "").strip()
        if not tx_tp_cd:
            summ = str(first_non_empty_value(r.get("summ"), r.get("交易类型"), r.get("tx_type")) or "")
            tx_tp_cd = resolve_incm_epn_tx_tp_cd(summ, incm_epn_tp_cd)

        key = (incm_epn_tp_cd, tx_tp_cd)
        if key not in agg:
            agg[key] = {"tp": incm_epn_tp_cd, "cd": tx_tp_cd, "sum": 0.0, "num": 0}
        agg[key]["sum"] += amt_abs
        agg[key]["num"] += 1
    return total_amt, agg


async def build_tx_incm_epn_analy_sum_response(req_msg_id: str, did: str, query: dict[str, Any]) -> dict[str, Any]:
    """
    对齐 app 端 `qryTxIncmEpnAnalySum/T080240` 的常见返回结构：
    - data: { incmEpnChgAmount, lastmonYearIncmEpnAmt, resultList, tnum, totalAmt }
    本实现：
    - 依据 raw mock 交易按 (incmEpnTpCd, incmEpnTxTpCd) 聚合 txAmt/txNum
    - measureUnitCode=0310：按自然年汇总（整年）；=0307：按 incmEpnMonth 单月
      （App 常同时带 incmEpnYear+incmEpnMonth，必须以 measureUnitCode 区分，不能见年就按年）
    - 未传 measureUnitCode 时：有合法四位年则按年，否则按月
    - 未传 incmEpnYear 且非强制按月时，若 hostRspPlain 的 data.fieldItemInfo 为「按年」模板，
      则从模板推断选中年份（与 T080239 一致）
    - lastmonYearIncmEpnAmt：按月取上一月总额，按年取上一年总额；incmEpnChgAmount=本期-上期
    - 若传入 hostRspPlain，则优先复用其 resultList 结构/顺序作为模板（避免缺项/顺序差异）
    """
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        raise RuntimeError(binding["error"])
    recs = await read_raw_mock_data_by_did(did)

    begin_date = normalize_ymd8(query.get("beginDate"))
    dline_date = normalize_ymd8(query.get("dlineDate"))
    incm_month = str(query.get("incmEpnMonth") or "").strip()  # YYYYMM
    incm_epn_year = str(query.get("incmEpnYear") or "").strip()  # YYYY，年度查询
    measure_unit = str(query.get("measureUnitCode") or "").strip()
    want_tp = str(query.get("incmEpnTpCd") or "").strip()  # 1/2
    host_rsp_plain = str(query.get("hostRspPlain") or "")

    # 0307=月度，0310=年度。二者常同时出现 year/month 参数，以 measureUnitCode 为准。
    if measure_unit == "0307":
        year_scope = False
    elif measure_unit == "0310":
        year_scope = bool(len(incm_epn_year) == 4 and incm_epn_year.isdigit())
        if not year_scope and len(incm_month) == 6 and incm_month.isdigit():
            incm_epn_year = incm_month[:4]
            year_scope = True
    else:
        year_scope = bool(len(incm_epn_year) == 4 and incm_epn_year.isdigit())

    if not year_scope and measure_unit != "0307" and host_rsp_plain:
        try:
            obj = json.loads(host_rsp_plain)
            fi = ((obj or {}).get("data") or {}).get("fieldItemInfo") or []
            _tk, ax_host = _parse_host_incm_epn_field_item_axis(fi)
            if ax_host == "year" and isinstance(fi, list):
                iy = _infer_selected_incm_epn_year_from_field_items(fi)
                if len(iy) == 4 and iy.isdigit():
                    incm_epn_year = iy
                    year_scope = True
        except Exception:
            pass
    if begin_date and dline_date and begin_date > dline_date:
        begin_date, dline_date = dline_date, begin_date

    total_amt, agg = _sum_raw_incm_epn_abs_amt(
        recs,
        want_tp=want_tp,
        year_scope=year_scope,
        incm_epn_year=incm_epn_year,
        incm_month=incm_month,
        begin_date=begin_date,
        dline_date=dline_date,
    )

    # 上期总额：月度→上一自然月；年度→上一年
    last_amt = 0.0
    if year_scope and len(incm_epn_year) == 4 and incm_epn_year.isdigit():
        prev_year = str(int(incm_epn_year) - 1)
        last_amt, _ = _sum_raw_incm_epn_abs_amt(
            recs,
            want_tp=want_tp,
            year_scope=True,
            incm_epn_year=prev_year,
            incm_month="",
            begin_date="",
            dline_date="",
        )
    elif (not year_scope) and len(incm_month) == 6 and incm_month.isdigit():
        prev_month = _prev_incm_epn_month(incm_month)
        if prev_month:
            last_amt, _ = _sum_raw_incm_epn_abs_amt(
                recs,
                want_tp=want_tp,
                year_scope=False,
                incm_epn_year="",
                incm_month=prev_month,
                begin_date="",
                dline_date="",
            )
    chg_amt = total_amt - last_amt

    # 先从 hostRspPlain 取模板 resultList（若有）
    template: list[dict[str, Any]] = []
    if host_rsp_plain:
        try:
            obj = json.loads(host_rsp_plain)
            rl = ((obj or {}).get("data") or {}).get("resultList") or []
            if isinstance(rl, list):
                for it in rl:
                    if isinstance(it, dict):
                        template.append(dict(it))
        except Exception:
            template = []

    def _ratio_text(v: float) -> str:
        if v <= 0:
            return "0"
        # 真机常见返回为 "1" 而非 "1.0000"
        s = f"{v:.4f}".rstrip("0").rstrip(".")
        return s or "0"

    result_list: list[dict[str, Any]] = []
    used: set[tuple[str, str]] = set()

    if template:
        for it in template:
            tp = str(it.get("incmEpnTpCd") or "").strip()
            cd = str(it.get("incmEpnTxTpCd") or "").strip()
            if not tp or not cd:
                continue
            key = (tp, cd)
            node = agg.get(key)
            tx_amt = float(node["sum"]) if node else 0.0
            tx_num = int(node["num"]) if node else 0
            it["incmEpnTpCd"] = tp
            it["incmEpnTxTpCd"] = cd
            it["txAmt"] = f"{tx_amt:.2f}"
            it["txNum"] = str(tx_num)
            it["txTypeAmtDataChgVal"] = f"{tx_amt:.2f}"
            it["txTypeAmtRatio"] = (_ratio_text(tx_amt / total_amt) if total_amt > 0 else "0")
            # lastmonYearTxamt 等字段保留模板原值（或缺省）
            if "lastmonYearTxamt" not in it:
                it["lastmonYearTxamt"] = "0.00"
            result_list.append(it)
            used.add(key)

    # 追加模板中没有的聚合项（按 tp/cd 排序）
    for (tp, cd), node in sorted(agg.items(), key=lambda x: (x[0][0], x[0][1])):
        if (tp, cd) in used:
            continue
        tx_amt = float(node["sum"])
        tx_num = int(node["num"])
        result_list.append(
            {
                "incmEpnTpCd": tp,
                "incmEpnTxTpCd": cd,
                "lastmonYearTxamt": "0.00",
                "txAmt": f"{tx_amt:.2f}",
                "txNum": str(tx_num),
                "txTypeAmtDataChgVal": f"{tx_amt:.2f}",
                "txTypeAmtRatio": (_ratio_text(tx_amt / total_amt) if total_amt > 0 else "0"),
            }
        )

    total_tx_num = sum(int(v["num"]) for v in agg.values())
    data = {
        "incmEpnChgAmount": f"{chg_amt:.2f}",
        "lastmonYearIncmEpnAmt": f"{last_amt:.2f}",
        "resultList": result_list,
        "tnum": str(total_tx_num if total_tx_num > 0 else len(result_list)),
        "totalAmt": f"{total_amt:.2f}",
    }
    print(
        "[mock-api] txIncmEpnAnalySum aggregate",
        {
            "incmEpnMonth": incm_month,
            "incmEpnTpCd": want_tp,
            "measureUnitCode": measure_unit,
            "totalAmt": data["totalAmt"],
            "resultSize": len(result_list),
            "aggSize": len(agg),
        },
    )
    return {
        "code": "000000",
        "data": data,
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


async def build_my_incm_epn_response(req_msg_id: str, did: str, query: dict[str, Any]) -> dict[str, Any]:
    """
    对齐 app 端 `qryMyIncmEpn/T080505` 的常见返回结构：
    - data: { expnTotAmt, incomeTotalAmt, jumpInfo, mmlyAccblSta, mmlyAccblTips }
    这里根据 raw mock 交易汇总：支出/收入分别求绝对值总和。
    """
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        raise RuntimeError(binding["error"])
    recs = await read_raw_mock_data_by_did(did)
    begin_date = normalize_ymd8(query.get("beginDate"))
    dline_date = normalize_ymd8(query.get("dlineDate"))
    if begin_date and dline_date and begin_date > dline_date:
        begin_date, dline_date = dline_date, begin_date

    # 该接口语义为“本月收支”：以 dlineDate（未传则今日）所在月份为准
    if dline_date:
        anchor = dline_date
    else:
        anchor = ymd_today_utc8()
        dline_date = anchor
    month_begin = f"{anchor[:6]}01"
    begin_date = month_begin

    expn = 0.0
    incm = 0.0
    for r0 in recs:
        r = r0 or {}
        tx_date = normalize_ymd8(r.get("txDate") or r.get("交易日期") or r.get("tx_datetime") or "")
        if (begin_date or dline_date) and not tx_date:
            continue
        if begin_date and tx_date and tx_date < begin_date:
            continue
        if dline_date and tx_date and tx_date > dline_date:
            continue
        amt_raw = first_non_empty_value(r.get("txAmt"), r.get("tx_amt"), r.get("交易金额"), r.get("amount"), "0")
        try:
            amt = float(str(amt_raw).replace(",", "").strip() or "0")
        except Exception:
            amt = 0.0
        amt_abs = abs(amt)
        if amt_abs <= 0:
            continue
        dw = str(first_non_empty_value(r.get("dwFlagCode"), r.get("incmEpnTpCd"), "1") or "1").strip()
        if dw == "2":
            expn += amt_abs
        else:
            incm += amt_abs

    jump_info = json_dumps(
        {
            "myIncm": {"pageId": "1548", "h5AppData": {"type": "1"}},
            "myExpn": {"pageId": "1548", "h5AppData": {"type": "0"}},
            "mmlyAccbl": {"pageId": "1784", "tips": "查看上月账单"},
        }
    )
    return {
        "code": "000000",
        "data": {
            "expnTotAmt": f"{expn:.2f}",
            "incomeTotalAmt": f"{incm:.2f}",
            "jumpInfo": jump_info,
            "mmlyAccblSta": "0",
            "mmlyAccblTips": "查看上月账单",
        },
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


# 真机 qryTopDtlList/T080492 detailList 常见字段（收支分析「大额/Top」明细）
TOP_DTL_LIST_KEYS: tuple[str, ...] = (
    "txAmt",
    "txDate",
    "txTime",
    "summ",
    "merDesc",
    "incmEpnTpCd",
    "incmEpnTxTpCd",
    "dwFlagCode",
    "mediumNo",
    "bkcdMask",
    "txOpsName",
    "txOpsAccno",
    "txRemark",
    "dtlSeqNo",
    "globalBusiTrackNo",
    "chnlKindCode",
    "currCode",
)


async def build_top_dtl_list_response(req_msg_id: str, did: str, query: dict[str, Any]) -> dict[str, Any]:
    """
    对齐 app 端 `qryTopDtlList/T080492`：
    - data: { detailList, incmEpnMonth, qryTime, tnum, totalAmt }
    - 按 incmEpnMonth（YYYYMM）+ incmEpnTpCd 过滤，按金额绝对值降序取前 datasize 条
    """
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        raise RuntimeError(binding["error"])
    device = binding["device"]
    basic_info = {
        "户名": device.get("person_name") or "",
        "卡号": device.get("card_no") or "",
        "身份证号": device.get("id_card") or "",
        "电话号码": device.get("phone_no") or "",
        "印章流水": device.get("stamp_no") or "",
        "抬头": device.get("title") or "",
        "profile_balance": device.get("profile_balance"),
    }
    recs = await read_raw_mock_data_by_did(did)
    incm_month = str(query.get("incmEpnMonth") or "").strip()
    want_tp = str(query.get("incmEpnTpCd") or "").strip()
    try:
        data_size = int(str(query.get("datasize") or query.get("dataSize") or "5").strip() or "5")
    except Exception:
        data_size = 5
    if data_size <= 0:
        data_size = 5
    if data_size > 50:
        data_size = 50

    scored: list[tuple[float, dict[str, Any]]] = []
    for r0 in recs:
        r = r0 or {}
        if not record_matches_incm_epn_month(r, incm_month):
            continue
        incm_epn_tp_cd = incm_epn_tp_cd_from_record(r)
        if want_tp in ("1", "2") and incm_epn_tp_cd != want_tp:
            continue
        tx_tp_filter = str(query.get("incmEpnTxTpCd") or "").strip()
        if tx_tp_filter:
            rec_tp = str(r.get("incmEpnTxTpCd") or r.get("incm_epn_tx_tp_cd") or "").strip()
            if not rec_tp:
                summ = str(first_non_empty_value(r.get("summ"), r.get("交易类型"), r.get("tx_type")) or "")
                rec_tp = resolve_incm_epn_tx_tp_cd(summ, incm_epn_tp_cd)
            if rec_tp != tx_tp_filter:
                continue
        amt_abs = record_amt_abs(r)
        if amt_abs <= 0:
            continue
        scored.append((amt_abs, r))

    scored.sort(key=lambda x: (-x[0], str(x[1].get("txDate") or ""), str(x[1].get("txTime") or "")))
    top_rows = [r for _, r in scored[:data_size]]
    total_amt = sum(a for a, _ in scored)

    details: list[dict[str, str]] = []
    for i, r in enumerate(top_rows):
        full = build_xhx_detail_list_item(r, basic_info, i + 1)
        item: dict[str, str] = {}
        for k in TOP_DTL_LIST_KEYS:
            item[k] = str(full.get(k) or "")
        details.append(item)

    now = datetime.now()
    qry_time = (
        f"{now.year}{str(now.month).zfill(2)}{str(now.day).zfill(2)}"
        f"{str(now.hour).zfill(2)}{str(now.minute).zfill(2)}{str(now.second).zfill(2)}"
    )
    return {
        "code": "000000",
        "data": {
            "detailList": details,
            "incmEpnMonth": incm_month if len(incm_month) == 6 else "",
            "qryTime": qry_time,
            "tnum": str(len(details)),
            "totalAmt": f"{total_amt:.2f}",
        },
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


def _pick_top_dtl_list_fields_from_obj(obj: Optional[dict[str, Any]]) -> dict[str, str]:
    out: dict[str, str] = {}
    if not isinstance(obj, dict):
        return out
    for k in ("incmEpnMonth", "incmEpnTpCd", "incmEpnTxTpCd"):
        v = obj.get(k)
        if v is not None and str(v).strip():
            out[k] = str(v).strip()
    ds = obj.get("datasize", obj.get("dataSize"))
    if ds is not None and str(ds).strip():
        out["datasize"] = str(ds).strip()
    return out


def top_dtl_list_payload_from_request(request: web.Request, body: dict[str, Any]) -> dict[str, str]:
    """解析 qryTopDtlList/T080492：query + body + hostReqPlain。"""
    b = body if isinstance(body, dict) else {}
    from_query = _pick_top_dtl_list_fields_from_obj({k: request.query.get(k) for k in (
        "incmEpnMonth", "incmEpnTpCd", "incmEpnTxTpCd", "datasize", "dataSize"
    )})
    from_body = _pick_top_dtl_list_fields_from_obj(b)
    from_plain: dict[str, str] = {}
    for key in ("hostReqPlain", "reqPlain", "plainRequest", "hostReqPlainText"):
        obj = _try_json_dict_from_text(b.get(key)) or _try_json_dict_from_text(request.query.get(key))
        if obj:
            merged = _pick_top_dtl_list_fields_from_obj(obj)
            for k, v in merged.items():
                if v:
                    from_plain[k] = v
    merged: dict[str, str] = {}
    for part in (from_query, from_body, from_plain):
        for k, v in part.items():
            if v:
                merged[k] = v
    return {
        "incmEpnMonth": merged.get("incmEpnMonth") or "",
        "incmEpnTpCd": merged.get("incmEpnTpCd") or "",
        "incmEpnTxTpCd": merged.get("incmEpnTxTpCd") or "",
        "datasize": merged.get("datasize") or "5",
    }


async def serve_top_dtl_list_request(request: web.Request, body_any: dict[str, Any]) -> web.Response:
    req_msg_id = pick_req_msg_id(request, body_any)
    did = pick_did(request)
    if not did:
        return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
    device_any = await load_device_context_any(did)
    if device_any and to_bool_int(device_any.get("is_active")) != 1:
        empty = {
            "code": "000000",
            "data": {"detailList": [], "incmEpnMonth": "", "qryTime": "", "tnum": "0", "totalAmt": "0.00"},
            "msg": "交易成功",
            "showType": "0",
            "reqMsgId": str(req_msg_id),
        }
        print("服务端响应(qryTopDtlList 设备未开启)===============>", empty)
        return json_response(empty)
    q_payload = top_dtl_list_payload_from_request(request, body_any)
    data = await build_top_dtl_list_response(str(req_msg_id), did, q_payload)
    print(
        "服务端响应(qryTopDtlList)===============>",
        {
            "tnum": (data.get("data") or {}).get("tnum"),
            "totalAmt": (data.get("data") or {}).get("totalAmt"),
            "incmEpnMonth": q_payload.get("incmEpnMonth"),
            "incmEpnTpCd": q_payload.get("incmEpnTpCd"),
            "datasize": q_payload.get("datasize"),
        },
    )
    return json_response(data)


# qryTransDetail/T080245：账单页搜索，fieldItemInfo 为扁平流水列表（非按月分组）
QRY_TRANS_DETAIL_KEYS: tuple[str, ...] = (
    "accBal",
    "bkcdMask",
    "cashExgVatgCd",
    "currCode",
    "dtlSeqNo",
    "dwFlagCode",
    "globalBusiTrackNo",
    "ibankFlag",
    "incmEpnTpCd",
    "incmEpnTxTpCd",
    "investProdtCdSets",
    "investProdtName",
    "mediumNo",
    "persInnerAccno",
    "randomAssignNo",
    "reckinIncmEpnFlagCd",
    "saccnoSeqNo",
    "subtxNo",
    "summ",
    "transInMobileNo",
    "txAmt",
    "txDate",
    "txOpsAccno",
    "txOpsName",
    "txTime",
)


def build_empty_qry_trans_detail_response(req_msg_id: str) -> dict[str, Any]:
    return {
        "code": "000000",
        "data": {
            "curQryReturnNum": "0",
            "expnTotAmt": "0.00",
            "fieldItemInfo": [],
            "haveNextDataFlag": "0",
            "incomeTotalAmt": "0.00",
            "qryResultTnum": "0",
            "qryTimeTotalExpnNum": "0",
            "qryTimeTotalIncomeNum": "0",
        },
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


def _record_matches_qry_cond(r: dict[str, Any], qry_cond: str) -> bool:
    q = str(qry_cond or "").strip()
    if not q:
        return True
    hay = " ".join(
        [
            str(first_non_empty_value(r.get("summ"), r.get("交易类型"), r.get("tx_type")) or ""),
            str(first_non_empty_value(r.get("txOpsName"), r.get("对手方户名"), r.get("counterparty_name")) or ""),
            str(first_non_empty_value(r.get("txOpsAccno"), r.get("对手方账户"), r.get("counterparty_account")) or ""),
            str(first_non_empty_value(r.get("merDesc")) or ""),
            str(first_non_empty_value(r.get("txRemark"), r.get("附言"), r.get("remark")) or ""),
            str(first_non_empty_value(r.get("txAmt"), r.get("交易金额"), r.get("amount")) or ""),
        ]
    )
    return q in hay


def _pick_medium_overlay_from_obj(obj: Optional[dict[str, Any]]) -> dict[str, str]:
    out: dict[str, str] = {}
    if not isinstance(obj, dict):
        return out
    for k in ("mediumNo", "persInnerAccno", "saccnoSeqNo"):
        v = obj.get(k)
        if v is not None and str(v).strip():
            out[k] = str(v).strip()
    ml = obj.get("mediumNoList")
    if isinstance(ml, list) and ml:
        first = ml[0] if isinstance(ml[0], dict) else {}
        for k in ("mediumNo", "persInnerAccno", "saccnoSeqNo"):
            if out.get(k):
                continue
            v = first.get(k) if isinstance(first, dict) else None
            if v is not None and str(v).strip():
                out[k] = str(v).strip()
    return out


def _pick_qry_trans_detail_fields_from_obj(obj: Optional[dict[str, Any]]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    if not isinstance(obj, dict):
        return out
    for k in (
        "beginDate",
        "dlineDate",
        "incmEpnTpCd",
        "incmEpnTxTpCd",
        "qryCond",
        "cfmFlag",
        "bgnIndexNo",
        "curQryReqNum",
        "mediumNo",
        "persInnerAccno",
        "saccnoSeqNo",
    ):
        v = obj.get(k)
        if v is None:
            continue
        s = str(v).strip()
        if s:
            out[k] = s
    tx_tp = obj.get("txTpCdList")
    if isinstance(tx_tp, list) and tx_tp:
        codes: list[str] = []
        for item in tx_tp:
            if isinstance(item, dict):
                code = str(item.get("incmEpnTxTpCd") or "").strip()
            else:
                code = str(item or "").strip()
            if code:
                codes.append(code)
        if codes:
            out["txTpCdList"] = codes
    overlay = _pick_medium_overlay_from_obj(obj)
    for k, v in overlay.items():
        if v and not out.get(k):
            out[k] = v
    return out


def qry_trans_detail_payload_from_request(request: web.Request, body: dict[str, Any]) -> dict[str, Any]:
    """解析 qryTransDetail/T080245：query + body + hostReqPlain。"""
    b = body if isinstance(body, dict) else {}
    from_query = _pick_qry_trans_detail_fields_from_obj({k: request.query.get(k) for k in (
        "beginDate",
        "dlineDate",
        "incmEpnTpCd",
        "incmEpnTxTpCd",
        "qryCond",
        "cfmFlag",
        "bgnIndexNo",
        "curQryReqNum",
        "mediumNo",
        "persInnerAccno",
        "saccnoSeqNo",
    )})
    from_body = _pick_qry_trans_detail_fields_from_obj(b)
    from_plain: dict[str, Any] = {}
    for key in ("hostReqPlain", "reqPlain", "plainRequest", "hostReqPlainText"):
        obj = _try_json_dict_from_text(b.get(key)) or _try_json_dict_from_text(request.query.get(key))
        if obj:
            merged = _pick_qry_trans_detail_fields_from_obj(obj)
            for k, v in merged.items():
                if v and not from_plain.get(k):
                    from_plain[k] = v
    for b64_key in ("hostReqPayloadB64", "reqPayloadB64", "hostPayloadB64"):
        for src in (b.get(b64_key), request.query.get(b64_key)):
            decoded = _decode_b64_json_maybe(src)
            if decoded:
                merged = _pick_qry_trans_detail_fields_from_obj(decoded)
                for k, v in merged.items():
                    if v and not from_plain.get(k):
                        from_plain[k] = v
    merged: dict[str, Any] = {}
    for part in (from_query, from_body, from_plain):
        for k, v in part.items():
            if v and not merged.get(k):
                merged[k] = v
    return {
        "beginDate": merged.get("beginDate") or "",
        "dlineDate": merged.get("dlineDate") or "",
        "incmEpnTpCd": merged.get("incmEpnTpCd") or "",
        "incmEpnTxTpCd": merged.get("incmEpnTxTpCd") or "",
        "qryCond": merged.get("qryCond") or "",
        "cfmFlag": merged.get("cfmFlag") or "",
        "bgnIndexNo": merged.get("bgnIndexNo") or "1",
        "curQryReqNum": merged.get("curQryReqNum") or "30",
        "mediumNo": merged.get("mediumNo") or "",
        "persInnerAccno": merged.get("persInnerAccno") or "",
        "saccnoSeqNo": merged.get("saccnoSeqNo") or "",
        "txTpCdList": merged.get("txTpCdList") or [],
    }


def build_qry_trans_detail_item(
    record: dict[str, Any], basic_info: dict[str, Any], overlay: dict[str, str]
) -> dict[str, str]:
    d = build_xhx_detail_list_item(record, basic_info, 1)
    d["dtlSeqNo"] = str(first_non_empty_value(record.get("dtlSeqNo"), d.get("dtlSeqNo"), "1") or "1")
    for k in ("persInnerAccno", "randomAssignNo", "subtxNo"):
        v = first_non_empty_value(record.get(k), d.get(k))
        if v:
            d[k] = v
    for k in ("mediumNo", "persInnerAccno", "saccnoSeqNo"):
        v = str((overlay or {}).get(k) or "").strip()
        if v:
            d[k] = v
    out: dict[str, str] = {}
    for k in QRY_TRANS_DETAIL_KEYS:
        out[k] = str(d.get(k) or "")
    return out


async def build_qry_trans_detail_response(req_msg_id: str, did: str, query: dict[str, Any]) -> dict[str, Any]:
    """
    对齐账单页搜索 qryTransDetail/T080245：
    - incmEpnTpCd=1 仅收入、=2 仅支出、空则全部
    - qryCond 关键词匹配摘要/对手方/账号/金额
    - beginDate/dlineDate 闭区间，bgnIndexNo/curQryReqNum 分页
    - data.fieldItemInfo 为扁平流水列表
    """
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        raise RuntimeError(binding["error"])
    device = binding["device"]
    basic_info = {
        "户名": device.get("person_name") or "",
        "卡号": device.get("card_no") or "",
        "身份证号": device.get("id_card") or "",
        "电话号码": device.get("phone_no") or "",
        "印章流水": device.get("stamp_no") or "",
        "抬头": device.get("title") or "",
        "profile_balance": device.get("profile_balance"),
    }
    recs = await read_raw_mock_data_by_did(did)
    begin_date = normalize_ymd8(query.get("beginDate"))
    dline_date = normalize_ymd8(query.get("dlineDate"))
    want_tp = str(query.get("incmEpnTpCd") or "").strip()
    qry_cond = str(query.get("qryCond") or "").strip()
    tx_tp_filter = str(query.get("incmEpnTxTpCd") or "").strip()
    tx_tp_cd_set = _parse_tx_tp_cd_set(query)
    if tx_tp_filter:
        tx_tp_cd_set.add(tx_tp_filter)
    if begin_date and dline_date and begin_date > dline_date:
        begin_date, dline_date = dline_date, begin_date
    try:
        bgn_index = max(1, int(str(query.get("bgnIndexNo") or "1").strip() or "1"))
    except ValueError:
        bgn_index = 1
    try:
        page_size = max(1, min(500, int(str(query.get("curQryReqNum") or "30").strip() or "30")))
    except ValueError:
        page_size = 30
    overlay = {
        "mediumNo": str(query.get("mediumNo") or "").strip(),
        "persInnerAccno": str(query.get("persInnerAccno") or "").strip(),
        "saccnoSeqNo": str(query.get("saccnoSeqNo") or "").strip() or "1",
    }

    matched: list[dict[str, Any]] = []
    expn_amt = 0.0
    incm_amt = 0.0
    expn_num = 0
    incm_num = 0
    for r0 in recs:
        r = r0 or {}
        tx_date = record_tx_date_ymd8(r)
        if not tx_date:
            continue
        if begin_date and tx_date < begin_date:
            continue
        if dline_date and tx_date > dline_date:
            continue
        incm_epn_tp_cd = incm_epn_tp_cd_from_record(r)
        if want_tp in ("1", "2") and incm_epn_tp_cd != want_tp:
            continue
        if tx_tp_cd_set:
            rec_tp = str(r.get("incmEpnTxTpCd") or r.get("incm_epn_tx_tp_cd") or "").strip()
            if not rec_tp:
                summ = str(first_non_empty_value(r.get("summ"), r.get("交易类型"), r.get("tx_type")) or "")
                rec_tp = resolve_incm_epn_tx_tp_cd(summ, incm_epn_tp_cd)
            if rec_tp not in tx_tp_cd_set:
                continue
        if not _record_matches_qry_cond(r, qry_cond):
            continue
        amt_abs = record_amt_abs(r)
        if incm_epn_tp_cd == "2":
            expn_amt += amt_abs
            expn_num += 1
        else:
            incm_amt += amt_abs
            incm_num += 1
        matched.append(r)

    total_matched = len(matched)
    start = bgn_index - 1
    page_rows = matched[start : start + page_size]
    have_next = "1" if start + page_size < total_matched else "0"
    field_item_info = [build_qry_trans_detail_item(r, basic_info, overlay) for r in page_rows]
    return {
        "code": "000000",
        "data": {
            "curQryReturnNum": str(len(field_item_info)),
            "expnTotAmt": f"{expn_amt:.2f}",
            "fieldItemInfo": field_item_info,
            "haveNextDataFlag": have_next,
            "incomeTotalAmt": f"{incm_amt:.2f}",
            "qryResultTnum": str(total_matched),
            "qryTimeTotalExpnNum": str(expn_num),
            "qryTimeTotalIncomeNum": str(incm_num),
        },
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


async def serve_qry_trans_detail_request(request: web.Request, body_any: dict[str, Any]) -> web.Response:
    req_msg_id = pick_req_msg_id(request, body_any)
    did = pick_did(request)
    if not did:
        return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
    device_any = await load_device_context_any(did)
    if device_any and to_bool_int(device_any.get("is_active")) != 1:
        empty = build_empty_qry_trans_detail_response(str(req_msg_id))
        print("服务端响应(qryTransDetail 设备未开启)===============>", empty)
        return json_response(empty)
    q_payload = qry_trans_detail_payload_from_request(request, body_any)
    data = await build_qry_trans_detail_response(str(req_msg_id), did, q_payload)
    print(
        "服务端响应(qryTransDetail)===============>",
        {
            "curQryReturnNum": (data.get("data") or {}).get("curQryReturnNum"),
            "qryResultTnum": (data.get("data") or {}).get("qryResultTnum"),
            "expnTotAmt": (data.get("data") or {}).get("expnTotAmt"),
            "incomeTotalAmt": (data.get("data") or {}).get("incomeTotalAmt"),
            "incmEpnTpCd": q_payload.get("incmEpnTpCd"),
            "qryCond": q_payload.get("qryCond"),
            "beginDate": q_payload.get("beginDate"),
            "dlineDate": q_payload.get("dlineDate"),
        },
    )
    return json_response(data)


def _pick_tx_incm_epn_analy_fields(obj: Optional[dict[str, Any]]) -> dict[str, str]:
    out: dict[str, str] = {}
    if not isinstance(obj, dict):
        return out
    for k in (
        "beginDate", "dlineDate", "incmEpnMonth", "incmEpnYear", "incmEpnTpCd",
        "incmEpnTxTpCd", "measureUnitCode",
    ):
        v = obj.get(k)
        if v is None:
            continue
        sv = str(v).strip()
        if sv:
            out[k] = sv
    return out


def tx_incm_epn_analy_sum_payload_from_request(request: web.Request, body: dict[str, Any]) -> dict[str, str]:
    b = body if isinstance(body, dict) else {}
    from_query = _pick_tx_incm_epn_analy_fields({k: request.query.get(k) for k in (
        "beginDate", "dlineDate", "incmEpnMonth", "incmEpnYear", "incmEpnTpCd",
        "incmEpnTxTpCd", "measureUnitCode",
    )})
    from_body = _pick_tx_incm_epn_analy_fields(b)
    host_rsp_plain = str(b.get("hostRspPlain") or "").strip()
    if not host_rsp_plain:
        b64 = str(b.get("hostPayloadB64") or request.query.get("hostPayloadB64") or "").strip()
        if b64:
            try:
                host_rsp_plain = base64.b64decode(b64).decode("utf-8", errors="replace")
            except Exception:
                host_rsp_plain = ""
    from_plain: dict[str, str] = {}
    for key in ("hostReqPlain", "reqPlain", "plainRequest", "hostReqPlainText"):
        obj = _try_json_dict_from_text(b.get(key)) or _try_json_dict_from_text(request.query.get(key))
        if obj:
            merged = _pick_tx_incm_epn_analy_fields(obj)
            for k, v in merged.items():
                if v:
                    from_plain[k] = v
    merged: dict[str, str] = {}
    for part in (from_query, from_body, from_plain):
        for k, v in part.items():
            if v:
                merged[k] = v
    return {
        "beginDate": merged.get("beginDate") or "",
        "dlineDate": merged.get("dlineDate") or "",
        "incmEpnMonth": merged.get("incmEpnMonth") or "",
        "incmEpnYear": merged.get("incmEpnYear") or "",
        "incmEpnTpCd": merged.get("incmEpnTpCd") or "",
        "incmEpnTxTpCd": merged.get("incmEpnTxTpCd") or "",
        "measureUnitCode": merged.get("measureUnitCode") or "",
        "hostRspPlain": host_rsp_plain,
    }


async def serve_tx_incm_epn_analy_sum_request(request: web.Request, body_any: dict[str, Any]) -> web.Response:
    req_msg_id = pick_req_msg_id(request, body_any)
    did = pick_did(request)
    if not did:
        return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
    q_payload = tx_incm_epn_analy_sum_payload_from_request(request, body_any)
    data = await build_tx_incm_epn_analy_sum_response(str(req_msg_id), did, q_payload)
    print("服务端响应(txIncmEpnAnalySum)===============>", data)
    return json_response(data)


def _pick_imex_sum_fields(obj: Optional[dict[str, Any]]) -> dict[str, str]:
    out: dict[str, str] = {}
    if not isinstance(obj, dict):
        return out
    for k in ("beginDate", "dlineDate", "incmEpnMonth", "incmEpnYear", "measureUnitCode"):
        v = obj.get(k)
        if v is None:
            continue
        sv = str(v).strip()
        if sv:
            out[k] = sv
    return out


def imex_sum_data_payload_from_request(request: web.Request, body: dict[str, Any]) -> dict[str, str]:
    b = body if isinstance(body, dict) else {}
    from_query = _pick_imex_sum_fields({k: request.query.get(k) for k in (
        "beginDate", "dlineDate", "incmEpnMonth", "incmEpnYear", "measureUnitCode",
    )})
    from_body = _pick_imex_sum_fields(b)
    host_rsp_plain = str(b.get("hostRspPlain") or "").strip()
    if not host_rsp_plain:
        b64 = str(b.get("hostPayloadB64") or request.query.get("hostPayloadB64") or "").strip()
        if b64:
            try:
                host_rsp_plain = base64.b64decode(b64).decode("utf-8", errors="replace")
            except Exception:
                host_rsp_plain = ""
    from_plain: dict[str, str] = {}
    for key in ("hostReqPlain", "reqPlain", "plainRequest", "hostReqPlainText"):
        obj = _try_json_dict_from_text(b.get(key)) or _try_json_dict_from_text(request.query.get(key))
        if obj:
            merged = _pick_imex_sum_fields(obj)
            for k, v in merged.items():
                if v:
                    from_plain[k] = v
    merged: dict[str, str] = {}
    for part in (from_query, from_body, from_plain):
        for k, v in part.items():
            if v:
                merged[k] = v
    return {
        "beginDate": merged.get("beginDate") or "",
        "dlineDate": merged.get("dlineDate") or "",
        "incmEpnMonth": merged.get("incmEpnMonth") or "",
        "incmEpnYear": merged.get("incmEpnYear") or "",
        "measureUnitCode": merged.get("measureUnitCode") or "",
        "hostRspPlain": host_rsp_plain,
    }


async def build_imex_sum_data_response(req_msg_id: str, did: str, query: dict[str, Any]) -> dict[str, Any]:
    """
    对齐 app 端 `qryImexSumData/T080764`：
    - data.fieldItemInfo: [{ incmEpnMonth, incmEpnYear, totIcmAmt, totIcmNum, totExpnAmt, totExpnNum }, ...]
    - measureUnitCode=0307：按月（incmEpnMonth=YYYYMM）
    - measureUnitCode=0310：按年（incmEpnYear=YYYY）；App 年度视图读此维度，按月返回会导致年度金额/图表错误
    """
    binding = await resolve_bound_device_context_for_mock(did)
    if not binding["ok"]:
        raise RuntimeError(binding["error"])
    recs = await read_raw_mock_data_by_did(did)

    begin_date = normalize_ymd8(query.get("beginDate"))
    dline_date = normalize_ymd8(query.get("dlineDate"))
    measure_unit = str(query.get("measureUnitCode") or "").strip()
    host_rsp_plain = str(query.get("hostRspPlain") or "")
    if begin_date and dline_date and begin_date > dline_date:
        begin_date, dline_date = dline_date, begin_date

    month_map: dict[str, dict[str, Any]] = {}
    year_map: dict[str, dict[str, Any]] = {}
    for r0 in recs:
        r = r0 or {}
        tx_date = record_tx_date_ymd8(r)
        if not tx_date:
            continue
        if begin_date and tx_date < begin_date:
            continue
        if dline_date and tx_date > dline_date:
            continue
        mm = tx_date[:6]
        yy = tx_date[:4]
        if mm not in month_map:
            month_map[mm] = {"icm_amt": 0.0, "icm_num": 0, "expn_amt": 0.0, "expn_num": 0}
        if yy not in year_map:
            year_map[yy] = {"icm_amt": 0.0, "icm_num": 0, "expn_amt": 0.0, "expn_num": 0}
        amt_abs = record_amt_abs(r)
        if amt_abs <= 0:
            continue
        tp = incm_epn_tp_cd_from_record(r)
        if tp == "2":
            month_map[mm]["expn_amt"] += amt_abs
            month_map[mm]["expn_num"] += 1
            year_map[yy]["expn_amt"] += amt_abs
            year_map[yy]["expn_num"] += 1
        else:
            month_map[mm]["icm_amt"] += amt_abs
            month_map[mm]["icm_num"] += 1
            year_map[yy]["icm_amt"] += amt_abs
            year_map[yy]["icm_num"] += 1

    host_fi: list[Any] = []
    if host_rsp_plain:
        try:
            obj = json.loads(host_rsp_plain)
            host_fi = ((obj or {}).get("data") or {}).get("fieldItemInfo") or []
        except Exception:
            host_fi = []

    template_keys, axis_host = _parse_host_incm_epn_field_item_axis(host_fi)
    if measure_unit == "0310":
        axis = "year"
    elif measure_unit == "0307":
        axis = "month"
    elif template_keys:
        axis = axis_host
    else:
        axis = "month"

    def _fmt_amt(v: float) -> str:
        return "0" if abs(v) < 0.0005 else f"{v:.2f}"

    field_item_info: list[dict[str, Any]] = []
    if axis == "year":
        keys: list[str] = []
        if template_keys and axis_host == "year":
            keys = template_keys
        if not keys:
            keys = _iter_yyyy_between(begin_date, dline_date)
        if not keys:
            keys = sorted(year_map.keys())
        for yy in keys:
            node = year_map.get(yy) or {"icm_amt": 0.0, "icm_num": 0, "expn_amt": 0.0, "expn_num": 0}
            field_item_info.append(
                {
                    "incmEpnMonth": "",
                    "incmEpnYear": yy,
                    "totIcmAmt": _fmt_amt(float(node["icm_amt"])),
                    "totIcmNum": str(int(node["icm_num"])),
                    "totExpnAmt": _fmt_amt(float(node["expn_amt"])),
                    "totExpnNum": str(int(node["expn_num"])),
                }
            )
    else:
        keys = []
        if template_keys and axis_host == "month":
            keys = template_keys
        if not keys:
            keys = _iter_yyyymm_between(begin_date, dline_date)
        if not keys:
            keys = sorted(month_map.keys())
        for mm in keys:
            node = month_map.get(mm) or {"icm_amt": 0.0, "icm_num": 0, "expn_amt": 0.0, "expn_num": 0}
            field_item_info.append(
                {
                    "incmEpnMonth": mm,
                    "incmEpnYear": "",
                    "totIcmAmt": _fmt_amt(float(node["icm_amt"])),
                    "totIcmNum": str(int(node["icm_num"])),
                    "totExpnAmt": _fmt_amt(float(node["expn_amt"])),
                    "totExpnNum": str(int(node["expn_num"])),
                }
            )

    print(
        "[mock-api] imexSumData aggregate",
        {
            "beginDate": begin_date,
            "dlineDate": dline_date,
            "measureUnitCode": measure_unit,
            "axis": axis,
            "items": len(field_item_info),
            "sample": field_item_info[-1] if field_item_info else None,
        },
    )
    return {
        "code": "000000",
        "data": {"fieldItemInfo": field_item_info},
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


async def serve_imex_sum_data_request(request: web.Request, body_any: dict[str, Any]) -> web.Response:
    req_msg_id = pick_req_msg_id(request, body_any)
    did = pick_did(request)
    if not did:
        return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
    q_payload = imex_sum_data_payload_from_request(request, body_any)
    data = await build_imex_sum_data_response(str(req_msg_id), did, q_payload)
    print("服务端响应(imexSumData)===============>", data)
    return json_response(data)


def build_empty_xhx_response(req_msg_id: str) -> dict[str, Any]:
    now = datetime.now()
    qry_time = (
        f"{now.year}{str(now.month).zfill(2)}{str(now.day).zfill(2)}"
        f"{str(now.hour).zfill(2)}{str(now.minute).zfill(2)}{str(now.second).zfill(2)}"
    )
    return {
        "code": "000000",
        "data": {
            "curQryReturnNum": "0",
            "fieldItemInfo": [],
            "haveNextDataFlag": "0",
            "qryResultTnum": "0",
            "qryTime": qry_time,
        },
        "msg": "交易成功",
        "showType": "0",
        "reqMsgId": req_msg_id or "",
    }


_DEFAULT_CORS_HEADERS: dict[str, str] = {
    "Access-Control-Allow-Origin": "*",
    # 含 Authorization；X-Admin-Token 为部分反向代理剥离 Authorization 时的备用传参
    "Access-Control-Allow-Headers": (
        "Content-Type, Authorization, X-Admin-Token, x-device-id, x-req-msg-id, Accept, "
        "X-Requested-With, did"
    ),
    "Access-Control-Allow-Methods": "GET,POST,PUT,DELETE,OPTIONS",
}


def _parse_cors_allowed_origins() -> set[str]:
    """
    返回允许的 Origin 白名单（完整匹配，含 scheme/host/port）。
    可用环境变量 MOCK_CORS_ORIGINS 覆盖（逗号分隔）。
    """
    raw = str(os.environ.get("MOCK_CORS_ORIGINS") or "").strip()
    if raw:
        return {s.strip() for s in raw.split(",") if s and s.strip()}
    # 默认放行常见管理端/本地开发域名（含 https 与 http）
    return {
        "https://a.psbclqd.com",
        "http://a.psbclqd.com",
        "https://psbclqd.com",
        "http://psbclqd.com",
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "http://localhost:3000",
        "http://127.0.0.1:3000",
    }


_CORS_ALLOWED_ORIGINS = _parse_cors_allowed_origins()


@web.middleware
async def cors_middleware(request: web.Request, handler: Any) -> web.StreamResponse:
    resp = await handler(request)
    try:
        origin = str(request.headers.get("Origin") or "").strip()
        if origin and origin in _CORS_ALLOWED_ORIGINS:
            resp.headers["Access-Control-Allow-Origin"] = origin
            # 多 Origin 场景必须加 Vary，避免缓存串号
            resp.headers["Vary"] = "Origin"
        # 确保预检/实际请求都带上允许头/方法
        for k, v in _DEFAULT_CORS_HEADERS.items():
            resp.headers.setdefault(k, v)
        resp.headers.setdefault("Access-Control-Max-Age", "86400")
    except Exception:
        # 不中断主流程
        pass
    return resp


def json_response(payload: Any, status: int = 200) -> web.Response:
    body = json_dumps(payload)
    return web.Response(
        status=status,
        text=body,
        content_type="application/json",
        charset="utf-8",
        headers=dict(_DEFAULT_CORS_HEADERS),
    )


def _is_mock_response_path(path: str) -> bool:
    p = str(path or "")
    return p.startswith("/api/mock-") or p.startswith("/sn13/")


def _response_body_text(resp: web.StreamResponse) -> str:
    text = getattr(resp, "text", None)
    if isinstance(text, str):
        return text
    raw = getattr(resp, "body", b"") or b""
    if isinstance(raw, (bytes, bytearray)):
        return raw.decode("utf-8", errors="replace")
    return str(raw)


def write_mock_response_log(method: str, path: str, query: dict[str, Any], status: int, body_text: str) -> None:
    log_path = str(MOCK_RESPONSE_LOG_FILE or "").strip()
    if not log_path:
        return
    try:
        parent = os.path.dirname(log_path)
        if parent:
            os.makedirs(parent, exist_ok=True)
        ts = now_utc8().strftime("%Y-%m-%d %H:%M:%S")
        q = "&".join(f"{k}={v}" for k, v in (query or {}).items())
        req_line = f"{method} {path}" + (f"?{q}" if q else "")
        payload = body_text if str(body_text).endswith("\n") else f"{body_text}\n"
        with open(log_path, "a", encoding="utf-8") as f:
            f.write(f"\n[{ts}] {req_line} status={status}\n")
            f.write(payload)
            f.flush()
    except Exception as e:
        print(f"[mock-api] write mock response log failed: {e}")


@web.middleware
async def mock_response_log_middleware(request: web.Request, handler: Any) -> web.StreamResponse:
    resp = await handler(request)
    if request.method == "OPTIONS" or not _is_mock_response_path(request.path):
        return resp
    try:
        write_mock_response_log(
            request.method,
            request.path,
            dict(request.query),
            int(getattr(resp, "status", 0) or 0),
            _response_body_text(resp),
        )
    except Exception as e:
        print(f"[mock-api] mock response log middleware error: {e}")
    return resp


async def parse_json_body(request: web.Request) -> Any:
    raw = await request.read()
    if not raw:
        return {}
    text = raw.decode("utf-8").strip()
    if not text:
        return {}
    try:
        return json.loads(text)
    except json.JSONDecodeError as e:
        raise ValueError(f"invalid json body: {e}") from e


def parse_tx_edit_body(body_raw: Any) -> dict[str, Any]:
    """
    编辑流水时兼容三种输入：
    1) 直接 JSON 对象；
    2) 整个请求体是 JSON 字符串（字符串内容再解一次）；
    3) 对象内携带 json/body/raw 字段，字段值为 JSON 字符串。
    """

    def _parse_text_json_obj(text: str) -> dict[str, Any]:
        s = str(text or "").strip()
        if not s:
            raise ValueError("json 内容不能为空")
        if s.startswith("```"):
            s = re.sub(r"^```(?:json)?\s*", "", s, flags=re.IGNORECASE)
            s = re.sub(r"\s*```$", "", s)
        try:
            parsed = json.loads(s)
        except json.JSONDecodeError as e:
            raise ValueError(f"json 文本解析失败: {e}") from e
        if not isinstance(parsed, dict):
            raise ValueError("json 文本必须是对象（例如 {\"tx_datetime\": \"...\"}）")
        return parsed

    if isinstance(body_raw, dict):
        for key in ("json", "body", "raw"):
            v = body_raw.get(key)
            if isinstance(v, str) and v.strip():
                return _parse_text_json_obj(v)
        return body_raw
    if isinstance(body_raw, str):
        return _parse_text_json_obj(body_raw)
    raise ValueError("请求体必须是 JSON 对象或 JSON 字符串")


def apply_admin_tx_raw_overrides(body: dict[str, Any]) -> None:
    """管理端 JSON 编辑时，用户常直接改 raw 字段（txOpsName/txRemark/txAmt/txTime...）。
    由于 merge_raw_api_fields_from_normalized_import 仅在字段缺失时 fill，这里需要把 raw 的显式改动同步覆盖到规范字段，
    否则看起来“保存了但列表还是原值”。
    """

    def _text(v: Any) -> str:
        return str(v or "").strip()

    def _has(v: Any) -> bool:
        s = _text(v)
        return bool(s) and s != "-"

    # 对手方信息
    if _has(body.get("txOpsName")):
        body["counterparty_name"] = _text(body.get("txOpsName"))
    if _has(body.get("txOpsAccno")):
        body["counterparty_account"] = _text(body.get("txOpsAccno"))

    # 备注/附言
    if _has(body.get("txRemark")):
        body["remark"] = _text(body.get("txRemark"))

    # 交易类型（raw: summ）
    if _has(body.get("summ")):
        body["tx_type"] = _text(body.get("summ"))

    # 渠道（raw: chnlKindCode）→ 主列中文；数字仍保留在 chnlKindCode 字段
    if _has(body.get("chnlKindCode")):
        cd = _text(body.get("chnlKindCode"))
        body["channel"] = resolve_transaction_channel_cn(
            summ=_text(body.get("summ")),
            tx_type=_text(body.get("tx_type")),
            remark=_text(body.get("remark")),
            channel=cd,
            chnl_kind_code=cd,
        )

    # 币种（raw: currCode；项目内规范字段常用“人民币”）
    if _has(body.get("currCode")):
        cc = _text(body.get("currCode")).upper()
        if cc in ("CNY", "156", "RMB"):
            body["currency"] = "人民币"
        else:
            body["currency"] = _text(body.get("currCode"))

    # 账户余额（raw: accBal）
    if body.get("accBal") is not None and _text(body.get("accBal")) != "":
        body["account_balance"] = body.get("accBal")

    # 交易时间（raw: txTime/txDate）
    if _has(body.get("txTime")) or _has(body.get("txDate")):
        dt = fmt_ymd_hms(first_non_empty_value(body.get("tx_datetime"), body.get("txTime"), body.get("txDate")))
        ymdhms = dt.get("ymdhms") or ""
        if ymdhms and re.fullmatch(r"\d{14}", ymdhms):
            body["tx_datetime"] = f"{ymdhms[:4]}-{ymdhms[4:6]}-{ymdhms[6:8]} {ymdhms[8:10]}:{ymdhms[10:12]}:{ymdhms[12:14]}"

    # 交易金额（raw: txAmt + 借贷标志）
    if body.get("txAmt") is not None and _text(body.get("txAmt")) != "":
        amt = to_num_or_null(body.get("txAmt"))
        if amt is not None:
            dw = _text(first_non_empty_value(body.get("dwFlagCode"), body.get("incmEpnTpCd")))
            signed = -abs(amt) if dw == "2" else abs(amt)
            body["tx_amount"] = signed


async def ensure_profile_exists(profile_id: int) -> bool:
    row = await query_one(f"SELECT id FROM `{PROFILE_TABLE}` WHERE id=%s LIMIT 1", (profile_id,))
    return bool(row)


def match_path(pathname: str, pattern: str) -> Optional[dict[str, str]]:
    actual = [p for p in pathname.split("/") if p]
    expected = [p for p in pattern.split("/") if p]
    if len(actual) != len(expected):
        return None
    params: dict[str, str] = {}
    for a, p in zip(actual, expected):
        if p.startswith(":"):
            params[p[1:]] = a
        elif p != a:
            return None
    return params


def hash_admin_password(pw: str) -> str:
    iterations = 200_000
    salt = secrets.token_hex(16)
    dk = hashlib.pbkdf2_hmac(
        "sha256", pw.encode("utf-8"), salt.encode("ascii"), iterations
    )
    enc = base64.b64encode(dk).decode("ascii")
    return f"pbkdf2_sha256${iterations}${salt}${enc}"


def verify_admin_password(pw: str, stored: str) -> bool:
    try:
        parts = str(stored or "").split("$")
        if len(parts) != 4 or parts[0] != "pbkdf2_sha256":
            return False
        iterations = int(parts[1])
        salt_s = parts[2]
        want = parts[3]
        dk = hashlib.pbkdf2_hmac("sha256", pw.encode("utf-8"), salt_s.encode("ascii"), iterations)
        got = base64.b64encode(dk).decode("ascii")
        return secrets.compare_digest(got, want)
    except (ValueError, TypeError):
        return False


def _jwt_encode_to_str(token: Any) -> str:
    """PyJWT 在部分版本返回 bytes，str(bytes) 会变成 "b'eyJ...'" 导致客户端误传。"""
    if isinstance(token, (bytes, bytearray)):
        return token.decode("ascii")
    return str(token)


def jwt_issue_admin_token(admin_id: int, role: int, username: str) -> str:
    now = datetime.utcnow()
    exp = now + timedelta(days=ADMIN_JWT_EXP_DAYS)
    payload: dict[str, Any] = {
        "sub": str(admin_id),
        "role": int(role),
        "usr": str(username),
        "iat": int(now.timestamp()),
        "exp": int(exp.timestamp()),
    }
    return _jwt_encode_to_str(jwt.encode(payload, ADMIN_JWT_SECRET, algorithm="HS256"))


def _strip_malformed_b_prefix_jwt(s: str) -> str:
    """兼容已发出的错误 str(bytes) 形态：字面量 b'....'。"""
    t = str(s).strip()
    if len(t) >= 4 and t[:2] == "b'" and t.endswith("'"):
        return t[2:-1].strip()
    if len(t) >= 4 and t[:2] == 'b"' and t.endswith('"'):
        return t[2:-1].strip()
    return t


def bearer_token_from_request(request: web.Request) -> str:
    h = request.headers.get("Authorization") or request.headers.get("authorization") or ""
    h = h.strip()
    if h.startswith("Bearer "):
        return _strip_malformed_b_prefix_jwt(h[7:].strip())
    # 部分 Nginx 默认不把 Authorization 传给上游，可在网关增加 proxy_set_header Authorization $http_authorization;
    # 此处兼容前端同时发送的备用头；亦可 GET ?token=（仅排查用）
    alt = request.headers.get("X-Admin-Token") or request.headers.get("x-admin-token") or ""
    alt = _strip_malformed_b_prefix_jwt(str(alt).strip())
    if alt:
        return alt
    q = request.query.get("token") or ""
    return _strip_malformed_b_prefix_jwt(str(q).strip())


async def jwt_load_admin_record(request: web.Request) -> Optional[dict[str, Any]]:
    raw = bearer_token_from_request(request)
    if not raw:
        return None
    try:
        data = jwt.decode(raw, ADMIN_JWT_SECRET, algorithms=["HS256"])
        aid = to_id_or_null(data.get("sub"))
        if aid is None:
            return None
        row = await query_one(
            f"""SELECT id, username, role, parent_admin_id, {_admin_select_points_and_perms()}
            FROM `{ADMIN_TABLE}`
            WHERE id=%s LIMIT 1""",
            (aid,),
        )
        return dict(row) if row else None
    except jwt.PyJWTError:
        return None


async def _admin_profiles_base_sql(admin: dict[str, Any]) -> tuple[str, tuple[Any, ...]]:
    role = int(admin.get("role") or 0)
    aid = int(admin["id"])
    if role == ADMIN_ROLE_SUPER:
        sql = f"""
        SELECT p.*, (
          SELECT COUNT(1) FROM `{TX_TABLE}` t WHERE t.person_id = p.id
        ) AS tx_count
        FROM `{PROFILE_TABLE}` p
        """
        return sql, ()
    sql = f"""
    SELECT p.*, (
      SELECT COUNT(1) FROM `{TX_TABLE}` t WHERE t.person_id = p.id
    ) AS tx_count
    FROM `{PROFILE_TABLE}` p
    WHERE p.`{PROFILE_COL_CREATOR_ADMIN}` = %s
    """
    return sql, (aid,)


async def list_profiles_for_admin(admin: dict[str, Any]) -> list[dict[str, Any]]:
    base_sql, base_params = await _admin_profiles_base_sql(admin)
    sql = f"{base_sql} ORDER BY p.id DESC"
    return await query_all(sql, base_params)


async def list_profiles_for_admin_paginated(
    admin: dict[str, Any], page: int, page_size: int
) -> tuple[list[dict[str, Any]], int, int, int]:
    page = _parse_positive_int(page, 1, minimum=1)
    page_size = _parse_positive_int(page_size, 20, minimum=1, maximum=200)
    base_sql, base_params = await _admin_profiles_base_sql(admin)
    count_sql = f"SELECT COUNT(1) AS c FROM ({base_sql}) AS _p"
    crow = await query_one(count_sql, base_params)
    total = int(crow.get("c") or 0) if crow else 0

    offset = (page - 1) * page_size
    if total == 0:
        page = 1
        offset = 0
    elif offset >= total:
        page = max(1, (total + page_size - 1) // page_size)
        offset = (page - 1) * page_size

    data_sql = f"{base_sql} ORDER BY p.id DESC LIMIT %s OFFSET %s"
    rows = await query_all(data_sql, tuple(base_params) + (page_size, offset))
    return rows, total, page, page_size


async def _admin_devices_base_sql(admin: dict[str, Any]) -> tuple[str, tuple[Any, ...]]:
    role = int(admin.get("role") or 0)
    aid = int(admin["id"])
    if role == ADMIN_ROLE_SUPER:
        sql = f"""
        SELECT d.*, p.person_name, p.card_no
        FROM `{DEVICE_TABLE}` d
        LEFT JOIN `{PROFILE_TABLE}` p ON p.id = d.profile_id
        """
        return sql, ()
    sql = f"""
    SELECT d.*, p.person_name, p.card_no
    FROM `{DEVICE_TABLE}` d
    LEFT JOIN `{PROFILE_TABLE}` p ON p.id = d.profile_id
    WHERE d.`{DEVICE_COL_CREATOR_ADMIN}` = %s
    """
    return sql, (aid,)


async def list_devices_for_admin(admin: dict[str, Any]) -> list[dict[str, Any]]:
    base_sql, base_params = await _admin_devices_base_sql(admin)
    sql = f"{base_sql} ORDER BY d.id DESC"
    return await query_all(sql, base_params)


async def list_devices_for_admin_paginated(
    admin: dict[str, Any], page: int, page_size: int
) -> tuple[list[dict[str, Any]], int, int, int]:
    page = _parse_positive_int(page, 1, minimum=1)
    page_size = _parse_positive_int(page_size, 20, minimum=1, maximum=200)
    base_sql, base_params = await _admin_devices_base_sql(admin)
    count_sql = f"SELECT COUNT(1) AS c FROM ({base_sql}) AS _d"
    crow = await query_one(count_sql, base_params)
    total = int(crow.get("c") or 0) if crow else 0

    offset = (page - 1) * page_size
    if total == 0:
        page = 1
        offset = 0
    elif offset >= total:
        page = max(1, (total + page_size - 1) // page_size)
        offset = (page - 1) * page_size

    data_sql = f"{base_sql} ORDER BY d.id DESC LIMIT %s OFFSET %s"
    rows = await query_all(data_sql, tuple(base_params) + (page_size, offset))
    return rows, total, page, page_size


async def admin_profile_accessible(admin: dict[str, Any], profile_id: int) -> bool:
    if int(admin.get("role") or 0) == ADMIN_ROLE_SUPER:
        row = await query_one(f"SELECT id FROM `{PROFILE_TABLE}` WHERE id=%s LIMIT 1", (profile_id,))
        return bool(row)
    row = await query_one(
        f"SELECT `{PROFILE_COL_CREATOR_ADMIN}` AS ca FROM `{PROFILE_TABLE}` WHERE id=%s LIMIT 1",
        (profile_id,),
    )
    if not row:
        return False
    return int(row.get("ca") or 0) == int(admin["id"])


async def admin_device_accessible(admin: dict[str, Any], device_id: int) -> bool:
    if int(admin.get("role") or 0) == ADMIN_ROLE_SUPER:
        row = await query_one(f"SELECT id FROM `{DEVICE_TABLE}` WHERE id=%s LIMIT 1", (device_id,))
        return bool(row)
    row = await query_one(
        f"SELECT `{DEVICE_COL_CREATOR_ADMIN}` AS ca FROM `{DEVICE_TABLE}` WHERE id=%s LIMIT 1",
        (device_id,),
    )
    if not row:
        return False
    return int(row.get("ca") or 0) == int(admin["id"])


async def admin_validate_profile_binding_for_device(
    admin: dict[str, Any], profile_id: Optional[int]
) -> bool:
    """二级管理员绑定的资料必须是本人创建的；一级无限制。"""
    if profile_id is None:
        return True
    if int(admin.get("role") or 0) == ADMIN_ROLE_SUPER:
        return await ensure_profile_exists(profile_id)
    return await admin_profile_accessible(admin, int(profile_id))


async def admin_create_profile_with_points(admin: dict[str, Any], body: dict[str, Any]) -> dict[str, Any]:
    """创建资料：二级按发件模式扣对应积分；写入 creator_admin_id、mail_send_mode。失败返回 { ok: False, error }."""
    person_name = must_non_empty_text(body.get("person_name"))
    card_no = must_non_empty_text(body.get("card_no"))
    mail_send_mode = normalize_mail_send_mode(body.get("mail_send_mode"))
    if mail_send_mode == MAIL_SEND_MODE_SIMULATE and not admin_allows_simulate_mail(admin):
        return {"ok": False, "error": "未开通模拟真实发件权限，请联系一级管理员"}
    creator_id = int(admin["id"])
    role = int(admin.get("role") or 0)
    points_col = (
        ADMIN_COL_POINTS_MODE2
        if mail_send_mode == MAIL_SEND_MODE_SIMULATE
        else "points_balance"
    )
    mode_label = "模拟真实地址" if mail_send_mode == MAIL_SEND_MODE_SIMULATE else "普通"
    cust_lvl = normalize_cust_lvl(body.get("cust_lvl"))
    vals = (
        person_name,
        card_no,
        body.get("id_card"),
        body.get("phone_no"),
        body.get("stamp_no"),
        body.get("title"),
        to_num_or_null(body.get("balance_amount")),
        creator_id,
        mail_send_mode,
        cust_lvl,
    )
    insert_sql = f"""INSERT INTO `{PROFILE_TABLE}` (
      person_name, card_no, id_card, phone_no, stamp_no, title, `{PROFILE_COL_BALANCE}`,
      `{PROFILE_COL_CREATOR_ADMIN}`, `{PROFILE_COL_MAIL_SEND_MODE}`, `{PROFILE_COL_CUST_LVL}`
    ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"""
    if role == ADMIN_ROLE_SUPER:
        new_id = await execute_insert(insert_sql, vals)
        return {"ok": True, "id": new_id, "mail_send_mode": mail_send_mode}
    pool = await get_pool()
    async with pool.acquire() as conn:
        await conn.autocommit(False)
        try:
            async with conn.cursor() as cur:
                await cur.execute(
                    f"""UPDATE `{ADMIN_TABLE}` SET `{points_col}` = `{points_col}` - %s
                    WHERE id=%s AND role=%s AND `{points_col}` >= %s""",
                    (PROFILE_CREATE_COST_POINTS, creator_id, ADMIN_ROLE_SUB, PROFILE_CREATE_COST_POINTS),
                )
                if int(cur.rowcount or 0) != 1:
                    await conn.rollback()
                    return {
                        "ok": False,
                        "error": (
                            f"{mode_label}积分不足：创建一套「{mode_label}」资料需要 "
                            f"{PROFILE_CREATE_COST_POINTS} 积分"
                        ),
                    }
                await cur.execute(insert_sql, vals)
                new_id = int(cur.lastrowid)
            await conn.commit()
        except Exception:
            await conn.rollback()
            raise
        finally:
            await conn.autocommit(True)
    return {"ok": True, "id": new_id, "mail_send_mode": mail_send_mode}


def _hydrate_admin_tx_balance_display(row: dict[str, Any]) -> None:
    """管理端：account_balance 为空时用 accBal 回填展示；DECIMAL → float。"""
    bal = row.get("account_balance")
    if isinstance(bal, Decimal):
        row["account_balance"] = float(bal)
        return
    missing = bal is None
    if not missing and isinstance(bal, str):
        missing = not str(bal).strip()
    if missing:
        fb = first_non_empty_value(row.get("accBal"))
        if fb:
            n = to_num_or_null(fb)
            if n is not None:
                row["account_balance"] = n


def _parse_positive_int(v: Any, default: int, *, minimum: int = 1, maximum: Optional[int] = None) -> int:
    try:
        n = int(str(v).strip())
        if maximum is not None:
            n = min(maximum, max(minimum, n))
        else:
            n = max(minimum, n)
        return n
    except Exception:
        return default


def _normalize_admin_datetime_bound(raw: str, *, end_of_day: bool) -> str:
    s = str(raw or "").strip()
    if not s:
        return ""
    s = s.replace("T", " ").replace("Z", "")
    if "." in s:
        s = s.split(".", 1)[0].strip()
    if len(s) == 10 and s[4] == "-" and s[7] == "-":
        return f"{s} {'23:59:59' if end_of_day else '00:00:00'}"
    return s[:19] if len(s) >= 19 else s


def _parse_optional_decimal(v: Any) -> Optional[Decimal]:
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    try:
        return Decimal(s.replace(",", ""))
    except Exception:
        return None


def _admin_tx_filter_clause(profile_id: int, q: dict[str, Any]) -> tuple[str, list[Any]]:
    parts: list[str] = ["person_id=%s"]
    params: list[Any] = [profile_id]

    keyword = str(q.get("keyword") or q.get("q") or "").strip()
    if keyword:
        like = f"%{keyword}%"
        glo = TX_COL_GLOBAL
        cols_like = [
            "`remark`",
            "`counterparty_name`",
            "`counterparty_account`",
            "`counterparty_bank`",
            "`tx_type`",
            "`channel`",
            f"`{glo}`",
            "`summ`",
            "`txRemark`",
            "`merDesc`",
        ]
        ors = " OR ".join(f"{c} LIKE %s" for c in cols_like)
        parts.append(f"({ors})")
        params.extend([like] * len(cols_like))

    begin = _normalize_admin_datetime_bound(str(q.get("begin") or ""), end_of_day=False)
    if begin:
        parts.append("tx_datetime >= %s")
        params.append(begin)

    end = _normalize_admin_datetime_bound(str(q.get("end") or ""), end_of_day=True)
    if end:
        parts.append("tx_datetime <= %s")
        params.append(end)

    lo = _parse_optional_decimal(q.get("min_amount"))
    hi = _parse_optional_decimal(q.get("max_amount"))
    if lo is not None:
        parts.append("tx_amount >= %s")
        params.append(lo)
    if hi is not None:
        parts.append("tx_amount <= %s")
        params.append(hi)

    tx_type = str(q.get("tx_type") or "").strip()
    if tx_type:
        parts.append("tx_type LIKE %s")
        params.append(f"%{tx_type}%")

    channel = str(q.get("channel") or "").strip()
    if channel:
        parts.append("channel LIKE %s")
        params.append(f"%{channel}%")

    dc = str(q.get("device_collected") or "").strip()
    if dc == "1":
        parts.append(f"COALESCE(`{TX_COL_DEVICE_COLLECTED}`, 0) = 1")
    elif dc == "0":
        parts.append(f"COALESCE(`{TX_COL_DEVICE_COLLECTED}`, 0) = 0")

    return " AND ".join(parts), params


async def list_transactions_by_profile_id_paginated(
    profile_id: int,
    page: int,
    page_size: int,
    q: dict[str, Any],
) -> tuple[list[dict[str, Any]], int, int, int]:
    page = _parse_positive_int(page, 1, minimum=1)
    page_size = _parse_positive_int(page_size, 20, minimum=1, maximum=200)
    where_sql, base_params = _admin_tx_filter_clause(profile_id, q)
    count_sql = f"SELECT COUNT(1) AS c FROM `{TX_TABLE}` WHERE {where_sql}"
    crow = await query_one(count_sql, tuple(base_params))
    total = int(crow.get("c") or 0) if crow else 0

    offset = (page - 1) * page_size
    if total == 0:
        page = 1
        offset = 0
    elif offset >= total:
        page = max(1, (total + page_size - 1) // page_size)
        offset = (page - 1) * page_size

    data_sql = (
        f"SELECT * FROM `{TX_TABLE}` WHERE {where_sql} "
        f"ORDER BY tx_datetime DESC, id DESC LIMIT %s OFFSET %s"
    )
    data_params = tuple(base_params) + (page_size, offset)
    rows = await query_all(data_sql, data_params)
    for row in rows:
        _hydrate_admin_tx_balance_display(row)
    return rows, total, page, page_size


async def get_admin_transaction_by_id(tx_id: int) -> Optional[dict[str, Any]]:
    row = await query_one(f"SELECT * FROM `{TX_TABLE}` WHERE id=%s LIMIT 1", (int(tx_id),))
    if not row:
        return None
    d = dict(row)
    _hydrate_admin_tx_balance_display(d)
    return d


async def _bulk_update_tx_balances_case(
    person_id: int, updates: list[tuple[int, float, str]], profile_balance: float
) -> None:
    """用 CASE id 分块更新多行，并在同一连接内写回资料余额，避免逐笔 UPDATE。"""
    chunk_size = 400
    pool = await get_pool()
    async with pool.acquire() as conn:
        async with conn.cursor() as cur:
            for off in range(0, len(updates), chunk_size):
                chunk = updates[off : off + chunk_size]
                case_bal_parts: list[str] = []
                case_acc_parts: list[str] = []
                params: list[Any] = []
                ids: list[int] = []
                for tx_id, bal_f, acc_s in chunk:
                    tid = int(tx_id)
                    ids.append(tid)
                    case_bal_parts.append("WHEN %s THEN %s")
                    case_acc_parts.append("WHEN %s THEN %s")
                    params.extend([tid, float(bal_f)])
                    params.extend([tid, acc_s])
                in_ph = ", ".join(["%s"] * len(ids))
                sql = (
                    f"UPDATE `{TX_TABLE}` SET "
                    f"account_balance = CASE id {' '.join(case_bal_parts)} END, "
                    f"`accBal` = CASE id {' '.join(case_acc_parts)} END "
                    f"WHERE person_id = %s AND id IN ({in_ph})"
                )
                await cur.execute(sql, tuple(params) + (person_id,) + tuple(ids))
            await cur.execute(
                f"UPDATE `{PROFILE_TABLE}` SET `{PROFILE_COL_BALANCE}`=%s WHERE id=%s",
                (float(profile_balance), person_id),
            )


async def _bulk_update_tx_balances_rows_only(person_id: int, updates: list[tuple[int, float, str]]) -> None:
    """仅更新流水的 account_balance / accBal（不写资料表）。"""
    if not updates:
        return
    chunk_size = 400
    pool = await get_pool()
    async with pool.acquire() as conn:
        async with conn.cursor() as cur:
            for off in range(0, len(updates), chunk_size):
                chunk = updates[off : off + chunk_size]
                case_bal_parts: list[str] = []
                case_acc_parts: list[str] = []
                params: list[Any] = []
                ids: list[int] = []
                for tx_id, bal_f, acc_s in chunk:
                    tid = int(tx_id)
                    ids.append(tid)
                    case_bal_parts.append("WHEN %s THEN %s")
                    case_acc_parts.append("WHEN %s THEN %s")
                    params.extend([tid, float(bal_f)])
                    params.extend([tid, acc_s])
                in_ph = ", ".join(["%s"] * len(ids))
                sql = (
                    f"UPDATE `{TX_TABLE}` SET "
                    f"account_balance = CASE id {' '.join(case_bal_parts)} END, "
                    f"`accBal` = CASE id {' '.join(case_acc_parts)} END "
                    f"WHERE person_id = %s AND id IN ({in_ph})"
                )
                await cur.execute(sql, tuple(params) + (person_id,) + tuple(ids))


async def sync_profile_balance_from_latest_tx(person_id: int) -> tuple[str, str]:
    """按时间上最后一笔流水的赛后余额同步资料 balance_amount。"""
    row = await query_one(
        f"SELECT account_balance, accBal FROM `{TX_TABLE}` WHERE person_id=%s ORDER BY tx_datetime DESC, id DESC LIMIT 1",
        (person_id,),
    )
    eff = _effective_post_tx_balance_from_row(dict(row) if row else None)
    bal_f = float(_decimal_money_or_zero(eff))
    bs = to_amt_str(bal_f)
    await execute(
        f"UPDATE `{PROFILE_TABLE}` SET `{PROFILE_COL_BALANCE}`=%s WHERE id=%s",
        (bal_f, person_id),
    )
    return (bs, str(bal_f))


def _decimal_money_maybe(val: Any) -> Optional[Decimal]:
    """将 account_balance 等列转成两位 Decimal；空/无效则返回 None。"""
    if val is None:
        return None
    if isinstance(val, Decimal):
        return val.quantize(Decimal("0.01"))
    s = str(val).strip()
    if not s or s.lower() in ("null", "none", "-"):
        return None
    n = to_num_or_null(val)
    if n is None:
        return None
    return Decimal(str(n)).quantize(Decimal("0.01"))


def _decimal_money_or_zero(val: Any) -> Decimal:
    d = _decimal_money_maybe(val)
    return d if d is not None else Decimal("0").quantize(Decimal("0.01"))


def _effective_post_tx_balance_from_row(row: Optional[dict[str, Any]]) -> Optional[Decimal]:
    """最后一笔的「赛后余额」：优先 DECIMAL account_balance，否则 VARCHAR accBal（与列表展示逻辑一致）。"""
    if not row:
        return None
    for key in ("account_balance", "accBal"):
        raw = row.get(key)
        if raw is None:
            continue
        if isinstance(raw, str) and not raw.strip():
            continue
        d = _decimal_money_maybe(raw)
        if d is not None:
            return d
    merged = first_non_empty_value(row.get("account_balance"), row.get("accBal"))
    return _decimal_money_maybe(merged) if merged else None


async def fetch_latest_tx_balance_snapshot(person_id: int) -> Optional[dict[str, Any]]:
    """入库前调用：时间上最后一笔流水之后的余额（account_balance 与 accBal 择优）。"""
    row = await query_one(
        f"SELECT id, tx_datetime, account_balance, accBal FROM `{TX_TABLE}` WHERE person_id=%s ORDER BY tx_datetime DESC, id DESC LIMIT 1",
        (person_id,),
    )
    if not row:
        return None
    rd = dict(row)
    eff = _effective_post_tx_balance_from_row(rd)
    return {
        "id": int(rd["id"]),
        "tx_datetime": rd.get("tx_datetime"),
        "balance_after": eff,
    }


async def adjust_collected_transaction_balances_only(
    person_id: int,
    inserted_ids: list[int],
    balance_snapshot: Optional[dict[str, Any]],
) -> dict[str, Any]:
    """只改写本批入库流水的余额。

    在「入库前最后一笔」流水上的赛后余额基础上，按采集记录的时间顺序递进；不重写表中其它原有流水。
    inserted_ids：本批 INSERT 的主键。无快照时上期初视为 0。最后同步资料表的 balance_amount。
    """
    if not inserted_ids:
        pb_fmt, pb_raw = await sync_profile_balance_from_latest_tx(person_id)
        return {
            "mode": "skipped",
            "updated": 0,
            "profile_balance_set": pb_raw,
            "profile_balance_formatted": pb_fmt,
            "snapshot_balance_base": None,
        }

    def _txn_amt(r: dict[str, Any]) -> Decimal:
        raw_amt = r.get("tx_amount")
        if raw_amt is None or (isinstance(raw_amt, str) and not str(raw_amt).strip()):
            raise ValueError(f"流水 id={r.get('id')} 缺少有效交易金额 tx_amount")
        n = to_num_or_null(raw_amt)
        if n is None:
            raise ValueError(f"流水 id={r.get('id')} 交易金额非法")
        return Decimal(str(n)).quantize(Decimal("0.01"))

    in_ph = ", ".join(["%s"] * len(inserted_ids))
    collected = await query_all(
        f"SELECT id, tx_datetime, tx_amount FROM `{TX_TABLE}` WHERE person_id=%s AND id IN ({in_ph}) ORDER BY tx_datetime ASC, id ASC",
        (person_id,) + tuple(inserted_ids),
    )
    if not collected:
        raise RuntimeError(f"入库后未能按 person_id=%s id IN %s 查回采集行（用于写余额）" % (person_id, inserted_ids))

    snap_id = int(balance_snapshot["id"]) if balance_snapshot is not None else None
    snap_dt = balance_snapshot.get("tx_datetime") if balance_snapshot is not None else None

    opening_dec = (
        Decimal("0").quantize(Decimal("0.01"))
        if balance_snapshot is None
        else _decimal_money_or_zero(balance_snapshot.get("balance_after"))
    )
    running = opening_dec
    updates: list[tuple[int, float, str]] = []
    for r in collected:
        rid = int(r["id"])
        rdt = r.get("tx_datetime")
        # 仅滚动「时间上严格晚于入库前最后一笔」的插入行；更早或网关混进来的旧时点行保持插入时的余额不动
        if balance_snapshot is not None and _tx_row_leq(rdt, rid, snap_dt, snap_id):
            continue
        amt = _txn_amt(r)
        running = (running + amt).quantize(Decimal("0.01"))
        bal_str = to_amt_str(float(running))
        updates.append((rid, float(running), bal_str))

    await _bulk_update_tx_balances_rows_only(person_id, updates)
    pb_fmt, pb_raw = await sync_profile_balance_from_latest_tx(person_id)
    snap_base = (
        None if balance_snapshot is None else str(_decimal_money_or_zero(balance_snapshot.get("balance_after")))
    )
    return {
        "mode": "collected_rows_only",
        "updated": len(updates),
        "inserted_ids_adjusted": list(inserted_ids),
        "snapshot_balance_base": snap_base if balance_snapshot else None,
        "opening_carry": str(opening_dec),
        "last_collected_balance": str(running),
        "profile_balance_set": pb_raw,
        "profile_balance_formatted": pb_fmt,
    }


async def admin_delete_all_interest_rows(person_id: int) -> int:
    """删除该资料下所有判定为「利息类」的流水（管理端重算前清空利息，便于重新生成）。"""
    pid = int(person_id)
    rows = await query_all(
        f"""SELECT id, tx_type, remark, `summ`,
            `{TX_COL_INCM_TX_TP_CD}` AS incm_cd
            FROM `{TX_TABLE}` WHERE person_id=%s""",
        (pid,),
    )
    delete_ids: list[int] = []
    for r in rows or []:
        rr = dict(r)
        if not is_interest_like_transaction(
            summ=rr.get("summ"),
            tx_type=rr.get("tx_type"),
            remark=rr.get("remark"),
            tp_cd=rr.get("incm_cd"),
        ):
            continue
        delete_ids.append(int(rr["id"]))
    deleted = 0
    if delete_ids:
        chunk_size = 400
        for off in range(0, len(delete_ids), chunk_size):
            part = delete_ids[off : off + chunk_size]
            in_ph = ", ".join(["%s"] * len(part))
            await execute(
                f"DELETE FROM `{TX_TABLE}` WHERE person_id=%s AND id IN ({in_ph})",
                (pid,) + tuple(part),
            )
        deleted = len(delete_ids)
    return deleted


async def admin_deduplicate_same_calendar_day_interest_rows(person_id: int) -> dict[str, Any]:
    """
    管理端「重算」前置：同一资料同一自然日内若有多条判定为利息的流水，只保留一条（id 最大的，常为最近采集）。
    """
    pid = int(person_id)
    rows = await query_all(
        f"""SELECT id, tx_datetime, tx_type, remark, `summ`,
            `{TX_COL_INCM_TX_TP_CD}` AS incm_cd
            FROM `{TX_TABLE}` WHERE person_id=%s ORDER BY id ASC""",
        (pid,),
    )
    by_day: dict[str, list[int]] = {}
    for r in rows or []:
        rr = dict(r)
        if not is_interest_like_transaction(
            summ=rr.get("summ"),
            tx_type=rr.get("tx_type"),
            remark=rr.get("remark"),
            tp_cd=rr.get("incm_cd"),
        ):
            continue
        dkey = _calendar_date_yyyy_mm_dd_from_tx_datetime(rr.get("tx_datetime"))
        if not dkey:
            continue
        by_day.setdefault(dkey, []).append(int(rr["id"]))
    delete_ids: list[int] = []
    dup_days = 0
    for _cal, ids in by_day.items():
        if len(ids) <= 1:
            continue
        dup_days += 1
        keep_id = max(ids)
        delete_ids.extend(rid for rid in ids if rid != keep_id)
    deleted = 0
    if delete_ids:
        chunk_size = 400
        for off in range(0, len(delete_ids), chunk_size):
            part = delete_ids[off : off + chunk_size]
            in_ph = ", ".join(["%s"] * len(part))
            await execute(
                f"DELETE FROM `{TX_TABLE}` WHERE person_id=%s AND id IN ({in_ph})",
                (pid,) + tuple(part),
            )
        deleted = len(delete_ids)
    return {"interest_duplicate_days": dup_days, "interest_rows_deleted": deleted}


async def admin_recalculate_transaction_balances(
    person_id: int,
    *,
    opening_mode: str = "zero",
) -> dict[str, Any]:
    """
    管理端：先删除全部利息流水，按积数规则补回已到结息日的季度利息；再滚算余额。

    opening_mode:
      - ``zero``：期初 0（默认）
      - ``earliest_balance``：以删除利息后、最早一条流水的账户余额反推期初
        （期初 = 该条 account_balance − tx_amount），再对全部流水（含补回利息）滚算
    """
    mode = str(opening_mode or "zero").strip().lower()
    if mode in ("earliest", "earliest_balance", "from_earliest", "anchor"):
        mode = "earliest_balance"
    else:
        mode = "zero"

    interest_deleted = await admin_delete_all_interest_rows(person_id)
    baseline_rows = await query_all(
        f"""SELECT id, tx_datetime, tx_amount, account_balance
        FROM `{TX_TABLE}` WHERE person_id=%s
        ORDER BY tx_datetime ASC, id ASC""",
        (person_id,),
    )

    opening_dec = Decimal("0").quantize(Decimal("0.01"))
    opening_anchor: dict[str, Any] | None = None
    if mode == "earliest_balance":
        if not baseline_rows:
            raise ValueError("当前没有可用于锚定的流水，无法按「最早余额」重算（可改用从 0 开始）")
        first = dict(baseline_rows[0])
        bal_n = to_num_or_null(first.get("account_balance"))
        if bal_n is None:
            raise ValueError(
                f"最早一条流水 id={first.get('id')} 没有账户余额，无法作为初始余额（可改用从 0 开始）"
            )
        amt_n = to_num_or_null(first.get("tx_amount"))
        if amt_n is None:
            raise ValueError(f"最早一条流水 id={first.get('id')} 交易金额非法")
        bal_dec = Decimal(str(bal_n)).quantize(Decimal("0.01"))
        amt_dec = Decimal(str(amt_n)).quantize(Decimal("0.01"))
        # 该条余额视为入账后余额 → 期初 = 余额 − 本笔金额
        opening_dec = (bal_dec - amt_dec).quantize(Decimal("0.01"))
        opening_anchor = {
            "id": int(first["id"]),
            "account_balance": str(bal_dec),
            "tx_amount": str(amt_dec),
            "tx_datetime": str(first.get("tx_datetime") or ""),
        }

    schedule = interest_regeneration_schedule([dict(r) for r in baseline_rows], as_of=date.today())
    interest_regenerated_detail: list[dict[str, Any]] = []
    for slot in schedule:
        cd = slot["credit_date"]
        q = slot["quarter"]
        amt = slot["estimated_interest_yuan"]
        gb_base = f"ADM{person_id:06d}{str(q).replace('-', '')}"
        # 结息日记账为每季次日 21 日；入账时刻在当日凌晨 00:00:00～00:59:59 内随机
        day_total_sec = secrets.randbelow(3600)
        hh, rem = divmod(day_total_sec, 3600)
        mm, ss = divmod(rem, 60)
        body_gen: dict[str, Any] = {
            "tx_datetime": f"{cd.year:04d}-{cd.month:02d}-{cd.day:02d} {hh:02d}:{mm:02d}:{ss:02d}",
            "tx_amount": str(amt),
            "summ": "利息",
            "tx_type": "结息",
            "currency": "人民币",
            "channel": "20",
            "remark": f"管理端按积数规则重算生成季度结息（{q}）",
            "global_busi_track_no": gb_base[:31],
            "incmEpnTxTpCd": "1004",
            "dwFlagCode": "1",
        }
        ins = await admin_insert_transaction(person_id, body_gen)
        interest_regenerated_detail.append({"quarter": q, "id": ins.get("id"), "estimated_interest_yuan": str(amt)})
    rows = await query_all(
        f"SELECT id, tx_datetime, tx_amount FROM `{TX_TABLE}` WHERE person_id=%s ORDER BY tx_datetime ASC, id ASC",
        (person_id,),
    )
    interest_preview = demand_interest_preview_dict_list([dict(row) for row in rows])
    running = opening_dec
    updates: list[tuple[int, float, str]] = []

    for r in rows:
        raw_amt = r.get("tx_amount")
        if raw_amt is None or (isinstance(raw_amt, str) and not str(raw_amt).strip()):
            raise ValueError(f"流水 id={r.get('id')} 缺少有效交易金额 tx_amount")
        n = to_num_or_null(raw_amt)
        if n is None:
            raise ValueError(f"流水 id={r.get('id')} 交易金额非法")
        amt = Decimal(str(n)).quantize(Decimal("0.01"))
        running = (running + amt).quantize(Decimal("0.01"))
        bal_str = to_amt_str(float(running))
        updates.append((int(r["id"]), float(running), bal_str))

    await _bulk_update_tx_balances_case(person_id, updates, float(running))

    return {
        "updated": len(updates),
        "opening_mode": mode,
        "opening_balance_used": str(opening_dec),
        "opening_anchor": opening_anchor,
        "profile_balance_set": str(running),
        "demand_interest_preview": interest_preview,
        "interest_deleted": interest_deleted,
        "interest_regenerated": len(interest_regenerated_detail),
        "interest_regenerated_detail": interest_regenerated_detail,
    }


def normalize_xlsx_header_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date) and not isinstance(value, datetime):
        return f"{value.year:04d}-{value.month:02d}-{value.day:02d}"
    s = str(value).strip()
    s = re.sub(r"\s+", "", s)
    s = s.replace("(元)", "").replace("（元）", "")
    return s


def cell_to_tx_datetime_str(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date) and not isinstance(value, datetime):
        return f"{value.year:04d}-{value.month:02d}-{value.day:02d} 00:00:00"
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel

            dt = from_excel(value)
            return dt.strftime("%Y-%m-%d %H:%M:%S")
        except Exception:
            pass
    return str(value).strip()


def cell_to_time_hms(value: Any) -> str:
    if value is None or value == "":
        return "00:00:00"
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel

            dt = from_excel(float(value))
            if isinstance(dt, datetime):
                return dt.strftime("%H:%M:%S")
        except Exception:
            pass
    s = str(value).strip()
    digits = re.sub(r"[^\d]", "", s)
    if len(digits) == 6 and not re.search(r"[-/]", s):
        return f"{digits[:2]}:{digits[2:4]}:{digits[4:6]}"
    if len(digits) >= 14:
        return f"{digits[8:10]}:{digits[10:12]}:{digits[12:14]}"
    m = re.match(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", s)
    if m:
        hh = int(m.group(1))
        mm = m.group(2)
        ss = m.group(3) or "00"
        return f"{hh:02d}:{mm}:{ss}"
    return "00:00:00"


def merge_date_and_time_cells(date_cell: Any, time_cell: Any) -> str:
    date_part = cell_to_tx_datetime_str(date_cell)
    if not date_part:
        return ""
    m = re.match(r"(\d{4}-\d{2}-\d{2})", date_part.strip())
    ymd = m.group(1) if m else ""
    if not ymd:
        return date_part.strip()
    hms = cell_to_time_hms(time_cell)
    return f"{ymd} {hms}"


def xlsx_header_map_from_row(row: tuple[Any, ...]) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for j, c in enumerate(row):
        name = normalize_xlsx_header_cell(c)
        if name:
            header_map[name] = j
    return header_map


def xlsx_row_looks_like_tx_table_header(header_map: dict[str, int]) -> bool:
    if not header_map:
        return False
    i_amt = xlsx_header_col_index(header_map, "交易金额", "金额")
    i_debit = xlsx_header_col_index(header_map, "借方金额", "借方发生额", "支出")
    i_credit = xlsx_header_col_index(header_map, "贷方金额", "贷方发生额", "收入")
    has_amount = i_amt is not None or i_debit is not None or i_credit is not None
    i_date = xlsx_header_col_index(
        header_map, "交易日期", "记账日期", "入账日期", "交易日期时间", "记账日期时间"
    )
    i_time = xlsx_header_col_index(header_map, "交易时间", "记账时间", "入账时间")
    has_date = i_date is not None or i_time is not None
    return bool(has_amount and has_date)


def resolve_tx_datetime_from_row(
    pick: Callable[[Optional[int]], Any],
    header_map: dict[str, int],
) -> str:
    i_combo = xlsx_header_col_index(header_map, "交易日期时间", "记账日期时间")
    if i_combo is not None:
        s = cell_to_tx_datetime_str(pick(i_combo))
        if s:
            return s
    i_date = xlsx_header_col_index(header_map, "交易日期", "记账日期", "入账日期")
    i_time = xlsx_header_col_index(header_map, "交易时间", "记账时间", "入账时间")
    if i_date is not None and i_time is not None:
        merged = merge_date_and_time_cells(pick(i_date), pick(i_time))
        if merged:
            return merged
    if i_date is not None:
        s = cell_to_tx_datetime_str(pick(i_date))
        if s:
            return s
    if i_time is not None:
        s = cell_to_tx_datetime_str(pick(i_time))
        if s:
            return s
    return ""


def resolve_tx_amount_from_row(
    pick: Callable[[Optional[int]], Any],
    header_map: dict[str, int],
) -> Any:
    i_amt = xlsx_header_col_index(header_map, "交易金额", "金额")
    if i_amt is not None:
        return pick(i_amt)
    i_debit = xlsx_header_col_index(header_map, "借方金额", "借方发生额", "支出")
    i_credit = xlsx_header_col_index(header_map, "贷方金额", "贷方发生额", "收入")
    debit = to_num_or_null(pick(i_debit)) if i_debit is not None else None
    credit = to_num_or_null(pick(i_credit)) if i_credit is not None else None
    if debit is not None or credit is not None:
        d = debit or 0.0
        c = credit or 0.0
        return c - d
    return None


def parse_line_decimal_places(line: dict[str, Any]) -> int:
    """模板行小数位：-2 百位、-1 十位、0 整数、1～3 小数；兼容键名「小数位」。"""
    raw = line.get("decimal_places")
    if raw is None:
        raw = line.get("小数位")
    if raw is None:
        return 0
    if isinstance(raw, str) and not str(raw).strip():
        return 0
    try:
        n = int(float(str(raw).strip()))
    except (ValueError, TypeError):
        return 0
    if n < -2:
        return -2
    if n > 3:
        return 3
    return n


def quantize_auto_flow_abs_amount(raw: Any, dp: int) -> Decimal:
    """按自动生成流水规则，将金额绝对值量化到指定粒度（正数）。"""
    try:
        d = Decimal(str(raw).strip()).copy_abs()
    except Exception:
        d = Decimal(0)
    ip = int(dp)
    if ip == -2:
        return (d / Decimal(100)).quantize(Decimal("1"), rounding=ROUND_HALF_UP) * Decimal(100)
    if ip == -1:
        return (d / Decimal(10)).quantize(Decimal("1"), rounding=ROUND_HALF_UP) * Decimal(10)
    if ip <= 0:
        return d.quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    quant = Decimal(10) ** (-ip)
    return d.quantize(quant, rounding=ROUND_HALF_UP)


def decimal_amount_to_plain_str(d: Decimal) -> str:
    """写入 tx_amount / txAmt 的简洁字符串（去掉末尾多余 0）。"""
    s = format(d, "f")
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    if s in ("-0", "-0."):
        return "0"
    return s


def merge_raw_api_fields_from_normalized_import(
    body: dict[str, Any], tx_tp_cd: str, global_track_no: str, card_no: str, profile_person_name: str = ""
) -> None:
    """Excel 等录入往往只有中英文明细列；补全与设备采集一致的 raw 字段写入库表。"""
    auto_flow_dp = body.pop("_auto_flow_dp", None)
    raw_amt_src = first_non_empty_value(body.get("tx_amount"), body.get("txAmt"))
    amt_f = float(to_amt_str(raw_amt_src))
    dw_raw = str(first_non_empty_value(body.get("dwFlagCode"), body.get("incmEpnTpCd")) or "").strip()
    if dw_raw in ("1", "2"):
        dw = dw_raw
        if dw == "2":
            amt_f = -abs(amt_f)
        else:
            amt_f = abs(amt_f)
    else:
        dw = "2" if amt_f < 0 else "1"
    dt = fmt_ymd_hms(first_non_empty_value(body.get("tx_datetime"), body.get("txTime"), body.get("txDate")))
    ymd = dt.get("ymd") or ""
    ymdhms = dt.get("ymdhms") or ""
    abs_amt = abs(amt_f)
    chnl_raw = str(first_non_empty_value(body.get("channel"), body.get("chnlKindCode")) or "").strip()
    chnl_digit = chnl_raw if chnl_raw.isdigit() else "20"
    acc_bal = first_non_empty_value(body.get("account_balance"), body.get("accBal"))
    acc_bal_s = to_amt_str(acc_bal) if acc_bal is not None and str(acc_bal).strip() != "" else None
    card = str(card_no or "").strip()

    def fill(key: str, val: Any) -> None:
        if val is None:
            return
        cur = body.get(key)
        if cur is None or (isinstance(cur, str) and not cur.strip()):
            body[key] = val

    def force(key: str, val: Any) -> None:
        if val is None:
            return
        body[key] = val

    cp_name = first_non_empty_value(
        body.get("counterparty_name"),
        body.get("txOpsName"),
        body.get("merDesc"),
        body.get("对手方户名"),
    )
    cp_name_s = str(cp_name or "").strip()
    tx_type_s = str(first_non_empty_value(body.get("tx_type"), body.get("summ")) or "").strip()
    wallet_template_summ_set = {"快捷支付", "微信转账", "银联入账", "网联入账"}
    is_wallet_platform_tx = ("财付通" in cp_name_s) or ("支付宝" in cp_name_s) or (tx_type_s in wallet_template_summ_set)
    cp_name_tail = cp_name_s
    if is_wallet_platform_tx:
        # 例如：财付通支付科技有限公司-余其用 -> 余其用
        parts = [p.strip() for p in re.split(r"[-－]", cp_name_s) if str(p or "").strip()]
        if parts:
            cp_name_tail = parts[-1]

    fill("tx_type", body.get("summ"))
    fill("currency", body.get("currCode"))
    fill("tx_amount", f"{amt_f:.2f}")
    fill("account_balance", acc_bal_s)
    fill("counterparty_name", body.get("txOpsName"))
    fill("counterparty_account", body.get("txOpsAccno"))
    fill("remark", body.get("txRemark"))
    if ymdhms:
        fill("tx_datetime", f"{ymdhms[:4]}-{ymdhms[4:6]}-{ymdhms[6:8]} {ymdhms[8:10]}:{ymdhms[10:12]}:{ymdhms[12:14]}")
    fill("summ", body.get("tx_type"))
    if auto_flow_dp is not None:
        try:
            dpi = int(auto_flow_dp)
            abs_dec = Decimal(str(to_amt_str(raw_amt_src))).copy_abs()
            fill("txAmt", decimal_amount_to_plain_str(quantize_auto_flow_abs_amount(abs_dec, dpi)))
        except Exception:
            fill("txAmt", f"{abs_amt:.2f}")
    else:
        fill("txAmt", f"{abs_amt:.2f}")
    fill("txDate", ymd)
    fill("txTime", ymdhms)
    fill("accBal", acc_bal_s)
    fill("chnlKindCode", chnl_digit)
    fill("currCode", to_curr_code(body.get("currency")))
    fill("txOpsName", cp_name_tail if is_wallet_platform_tx else body.get("counterparty_name"))
    fill("txOpsAccno", body.get("counterparty_account"))
    if is_wallet_platform_tx:
        # 钱包平台流水按模板风格纠偏（避免被前置默认值 20/1008 固化）
        # 但管理端 JSON 编辑时可能显式修改了 merDesc/channel/类型等 raw 字段，此时不应强制覆盖用户输入。
        preserve_wallet_overrides = bool(body.get("_preserve_wallet_platform_overrides"))
        profile_name_s = str(profile_person_name or "").strip()
        if "财付通" in cp_name_s:
            platform_name = "财付通支付科技有限公司"
        elif "支付宝" in cp_name_s:
            platform_name = "支付宝（中国）网络技术有限公司"
        else:
            platform_name = cp_name_s
        mer_desc = platform_name
        if preserve_wallet_overrides:
            fill("merDesc", mer_desc)
            fill("chnlKindCode", "30")
            fill("incmEpnTxTpCd", "1010")
        else:
            force("merDesc", mer_desc)
            force("chnlKindCode", "30")
            force("incmEpnTxTpCd", "1010")
        fill("counterparty_name", cp_name_tail)
        fill("incm_epn_tx_tp_cd", "1010")
        if (not body.get("_preserve_summ_raw")) and "入账" in str(body.get("tx_type") or body.get("summ") or ""):
            force("summ", "网联入账")
            fill("tx_type", "网联入账")
    fill("txRemark", body.get("remark"))
    fill("globalBusiTrackNo", global_track_no)
    fill("global_busi_track_no", global_track_no)
    id_seed = "|".join(
        [
            str(global_track_no or "").strip(),
            str(ymdhms or "").strip(),
            str(card or "").strip(),
            str(abs_amt),
            str(cp_name_s),
            str(tx_tp_cd or "").strip(),
        ]
    )
    auto_ids = build_import_raw_ids(id_seed)
    fill("persInnerAccno", auto_ids["persInnerAccno"])
    fill("randomAssignNo", auto_ids["randomAssignNo"])
    fill("servNo", auto_ids["servNo"])
    fill("subtxNo", auto_ids["subtxNo"])
    fill("incmEpnTxTpCd", tx_tp_cd)
    fill("incmEpnTpCd", dw)
    fill("dwFlagCode", dw)
    if card:
        fill("mediumNo", card)
        _mask = format_bkcd_mask(card)
        if _mask:
            fill("bkcdMask", _mask)
    fill("ibankFlag", "1")
    fill("reckinIncmEpnFlagCd", "1")
    fill("saccnoSeqNo", "1")
    fill("dtlSeqNo", "1")
    fill("cashExgVatgCd", "2")

    # 采集/导入合并后统一写入中文「交易渠道」；网关编码在 chnlKindCode
    ch_disp = str(first_non_empty_value(body.get("channel"), "") or "").strip()
    ch_cd = str(first_non_empty_value(body.get("chnlKindCode"), "") or "").strip()
    body["channel"] = resolve_transaction_channel_cn(
        summ=str(body.get("summ") or ""),
        tx_type=str(body.get("tx_type") or ""),
        remark=str(body.get("remark") or ""),
        channel=ch_disp or ch_cd,
        chnl_kind_code=ch_cd or ch_disp,
    )
    summ_final = str(first_non_empty_value(body.get("summ"), body.get("tx_type")) or "").strip()
    apply_out_tx_sri_no_for_import_or_collect(
        body,
        summ_for_check=summ_final,
        ymd8=ymd,
        global_track_no=global_track_no,
    )


def xlsx_row_card_no_from_postal_meta(rows: list[tuple[Any, ...]]) -> tuple[Optional[str], list[str]]:
    warnings: list[str] = []
    if not rows:
        return None, warnings
    header = [normalize_xlsx_header_cell(c) for c in rows[0]]
    if not header or header[0] != "户名" or "卡号" not in header:
        return None, warnings
    try:
        card_idx = header.index("卡号")
    except ValueError:
        return None, warnings
    if len(rows) < 2:
        return None, warnings
    r1 = rows[1]
    if card_idx >= len(r1):
        return None, warnings
    raw = r1[card_idx]
    if raw is None:
        return None, warnings
    card = str(raw).strip()
    return (card if card else None), warnings


def xlsx_header_col_index(header_map: dict[str, int], *candidates: str) -> Optional[int]:
    for name in candidates:
        if name in header_map:
            return header_map[name]
    return None


def parse_naive_datetime_from_import_str(s: str) -> Optional[datetime]:
    """解析 Excel 导入产生的 tx_datetime 字符串（本地 naive），无法解析则返回 None。"""
    t = str(s or "").strip()
    if not t:
        return None
    m = re.match(
        r"(\d{4})[-/]?(\d{2})[-/]?(\d{2})(?:[T\s]?(\d{2}))?:?(\d{2})?:?(\d{2})?",
        t,
    )
    if not m:
        return None
    hh = int(m.group(4) or 0)
    mm = int(m.group(5) or 0)
    ss = int(m.group(6) or 0)
    try:
        return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)), hh, mm, ss)
    except ValueError:
        return None


def is_tx_datetime_after_now_for_import(tx_dt_str: str) -> bool:
    dt = parse_naive_datetime_from_import_str(tx_dt_str)
    if dt is None:
        return False
    # Excel 导入时间为北京时间 naive；须与 now_utc8 比较，勿用 datetime.now()（服务器常为 UTC）。
    return dt > now_utc8().replace(tzinfo=None)


def xlsx_parse_transaction_import(content: bytes) -> tuple[list[str], list[dict[str, Any]], Optional[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ValueError("服务器未安装 openpyxl，请执行: pip install openpyxl") from e
    warnings: list[str] = []
    bio = io.BytesIO(content)
    wb = load_workbook(bio, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        raise ValueError("Excel 工作表为空")
    xlsx_card, meta_warn = xlsx_row_card_no_from_postal_meta(rows)
    warnings.extend(meta_warn)
    header_row_index = -1
    header_map: dict[str, int] = {}
    scan_limit = min(len(rows), 80)
    for i in range(scan_limit):
        hm = xlsx_header_map_from_row(rows[i])
        if xlsx_row_looks_like_tx_table_header(hm):
            header_row_index = i
            header_map = hm
            break
    if header_row_index < 0 or not header_map:
        raise ValueError("未找到交易表头行（需同时包含日期/时间与金额类列，如「交易日期」+「交易金额」或「记账日期」+「借方/贷方金额」）")
    i_type = xlsx_header_col_index(header_map, "交易类型", "业务名称", "业务摘要", "交易摘要")
    if i_type is None:
        i_type = xlsx_header_col_index(header_map, "摘要")
    i_curr = xlsx_header_col_index(header_map, "交易币种", "币种", "货币")
    i_bal = xlsx_header_col_index(header_map, "账户余额", "余额", "账户余额元")
    i_cp = xlsx_header_col_index(header_map, "对手方户名", "对方户名", "对方名称")
    i_acc = xlsx_header_col_index(header_map, "对手方账户", "对方账号", "对方账户", "对方卡号")
    i_bank = xlsx_header_col_index(header_map, "对手银行", "对方开户行", "对方行名")
    i_remark = xlsx_header_col_index(header_map, "附言", "备注", "客户附言", "用途")
    i_channel = xlsx_header_col_index(header_map, "交易方式", "交易渠道", "渠道")
    i_track = xlsx_header_col_index(header_map, "外部系统流水", "全局业务跟踪号", "流水号", "主机流水号")
    i_incm = xlsx_header_col_index(header_map, "收入支出标志", "借贷标志", "收付标志")
    bodies: list[dict[str, Any]] = []
    for ri in range(header_row_index + 1, len(rows)):
        r = rows[ri]
        if not r:
            continue

        def pick(col: Optional[int]) -> Any:
            if col is None or col >= len(r):
                return None
            return r[col]

        tx_dt = resolve_tx_datetime_from_row(pick, header_map)
        amt_raw = resolve_tx_amount_from_row(pick, header_map)
        if (not tx_dt) and (amt_raw is None or str(amt_raw).strip() == ""):
            continue
        if not tx_dt:
            warnings.append(f"第 {ri + 1} 行已跳过：缺少交易日期/时间")
            continue
        if amt_raw is None or str(amt_raw).strip() == "":
            warnings.append(f"第 {ri + 1} 行已跳过：缺少交易金额")
            continue
        raw_type_cell = pick(i_type)
        remark_val = normalize_raw_tx_value(pick(i_remark))
        cp_name_val = normalize_raw_tx_value(pick(i_cp))
        body: dict[str, Any] = {
            "tx_datetime": tx_dt,
            # xlsx 导入：交易类型原样写入 raw 字段 summ（不做 strip/映射/纠偏）
            "summ": raw_type_cell,
            # tx_type 由后续 merge 从 summ 补全（用于兼容推断/展示）
            "tx_type": None,
            # 标记：后续 merge_raw_api_fields_from_normalized_import 不得改写 summ 原值
            "_preserve_summ_raw": True,
            # xlsx 导入：保留导入内容（例如 merDesc=附言），不要被“钱包平台模板纠偏”强制覆盖
            "_preserve_wallet_platform_overrides": True,
            "currency": normalize_raw_tx_value(pick(i_curr)) or "人民币",
            "tx_amount": amt_raw,
            "account_balance": pick(i_bal),
            "counterparty_name": cp_name_val,
            "counterparty_account": normalize_raw_tx_value(pick(i_acc)),
            "counterparty_bank": normalize_raw_tx_value(pick(i_bank)),
            "remark": remark_val,
            "channel": normalize_raw_tx_value(pick(i_channel)),
            # 邮政流水导入：merDesc 取“附言/备注”（而不是对手户名）
            "merDesc": remark_val or cp_name_val,
        }
        incm_raw = normalize_raw_tx_value(pick(i_incm))
        if incm_raw:
            low = incm_raw.lower()
            if any(x in low for x in ("支", "出", "借", "付")) and not any(x in low for x in ("收", "入", "还", "贷")):
                try:
                    n = float(to_amt_str(amt_raw))
                    if n > 0:
                        body["tx_amount"] = -abs(n)
                except Exception:
                    pass
            elif any(x in low for x in ("收", "入", "贷", "进")) and not any(x in low for x in ("支", "出", "借")):
                try:
                    n = float(to_amt_str(amt_raw))
                    body["tx_amount"] = abs(n)
                except Exception:
                    pass
        track = normalize_raw_tx_value(pick(i_track))
        if track:
            body["globalBusiTrackNo"] = track
            body["global_busi_track_no"] = track
        if is_tx_datetime_after_now_for_import(tx_dt):
            warnings.append(f"第 {ri + 1} 行已跳过：交易时间晚于当前时间（不允许导入未来数据）")
            continue
        ymd = fmt_ymd_hms(body["tx_datetime"])
        if ymd.get("ymd"):
            body["txDate"] = ymd["ymd"]
        bodies.append(body)
    if not bodies:
        raise ValueError("未解析到任何交易数据行")
    return warnings, bodies, xlsx_card


def split_tx_datetime_for_xlsx_template(tx_dt: Any) -> tuple[str, str]:
    if not tx_dt:
        return "", ""
    parsed = cell_to_tx_datetime_str(tx_dt)
    base = parsed or str(tx_dt).strip()
    if not base:
        return "", ""
    m = re.match(r"(\d{4}-\d{2}-\d{2})\s+(\d{2}:\d{2}:\d{2})", base)
    if m:
        return m.group(1), m.group(2)
    dtp = fmt_ymd_hms(base)
    ymd = dtp.get("ymd") or ""
    if len(ymd) == 8:
        return f"{ymd[:4]}-{ymd[4:6]}-{ymd[6:8]}", "00:00:00"
    return "", ""


def build_postal_style_import_template_xlsx_bytes(
    person_name: str,
    card_no: str,
    sample_rows: list[dict[str, Any]],
    max_samples: int = 30,
) -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise ValueError("服务器未安装 openpyxl，请执行: pip install openpyxl") from e
    wb = Workbook()
    ws = wb.active
    ws.title = "交易明细"
    meta_head = ["户名", "", "", "卡号"]
    ws.append(meta_head)
    ws.append([person_name or "", "", "", card_no or ""])
    ws.append([])
    headers = [
        "交易日期",
        "交易时间",
        "交易类型",
        "交易币种",
        "交易金额",
        "账户余额",
        "对手方户名",
        "对手方账户",
        "对手银行",
        "附言",
        "交易方式",
        "全局业务跟踪号",
    ]
    ws.append(headers)
    for row in (sample_rows or [])[:max_samples]:
        d_s, t_s = split_tx_datetime_for_xlsx_template(row.get("tx_datetime"))
        gb = str(
            row.get(TX_COL_GLOBAL) or row.get("global_busi_track_no") or row.get("globalBusiTrackNo") or ""
        ).strip()
        ws.append(
            [
                d_s,
                t_s,
                row.get("tx_type"),
                row.get("currency") or "人民币",
                row.get("tx_amount"),
                row.get("account_balance"),
                row.get("counterparty_name"),
                row.get("counterparty_account"),
                row.get("counterparty_bank"),
                row.get("remark"),
                row.get("channel"),
                gb or None,
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _json_decode_db(val: Any) -> Any:
    if val is None:
        return None
    if isinstance(val, (dict, list)):
        return val
    if isinstance(val, (bytes, bytearray)):
        val = val.decode("utf-8", errors="replace")
    if isinstance(val, str):
        if not val.strip():
            return {}
        return json.loads(val)
    return val


def parse_iso_date_only(s: Any) -> Optional[date]:
    t = str(s or "").strip()
    if not t:
        return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", t)
    if not m:
        return None
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


async def admin_auto_flow_template_accessible(admin: dict[str, Any], template_id: int) -> bool:
    row = await query_one(
        f"SELECT creator_admin_id FROM `{AUTO_FLOW_TEMPLATE_TABLE}` WHERE id=%s LIMIT 1",
        (int(template_id),),
    )
    if not row:
        return False
    if int(admin.get("role") or 0) == ADMIN_ROLE_SUPER:
        return True
    return int(row.get("creator_admin_id") or 0) == int(admin["id"])


async def list_auto_flow_templates_for_admin(admin: dict[str, Any]) -> list[dict[str, Any]]:
    if int(admin.get("role") or 0) == ADMIN_ROLE_SUPER:
        rows = await query_all(
            f"""SELECT id, ref_name, gen_settings, `lines`, creator_admin_id, created_at, updated_at
            FROM `{AUTO_FLOW_TEMPLATE_TABLE}` ORDER BY id DESC"""
        )
    else:
        rows = await query_all(
            f"""SELECT id, ref_name, gen_settings, `lines`, creator_admin_id, created_at, updated_at
            FROM `{AUTO_FLOW_TEMPLATE_TABLE}` WHERE creator_admin_id=%s ORDER BY id DESC""",
            (int(admin["id"]),),
        )
    out: list[dict[str, Any]] = []
    for r in rows or []:
        gs = _json_decode_db(r.get("gen_settings"))
        ln = _json_decode_db(r.get("lines"))
        out.append(
            {
                "id": int(r["id"]),
                "ref_name": str(r.get("ref_name") or ""),
                "gen_settings": gs if isinstance(gs, dict) else {},
                "lines": ln if isinstance(ln, list) else [],
                "creator_admin_id": int(r.get("creator_admin_id") or 0),
                "created_at": r.get("created_at"),
                "updated_at": r.get("updated_at"),
            }
        )
    return out


async def get_auto_flow_template_by_id(template_id: int) -> Optional[dict[str, Any]]:
    row = await query_one(
        f"""SELECT id, ref_name, gen_settings, `lines`, creator_admin_id
        FROM `{AUTO_FLOW_TEMPLATE_TABLE}` WHERE id=%s LIMIT 1""",
        (int(template_id),),
    )
    if not row:
        return None
    gs = _json_decode_db(row.get("gen_settings"))
    ln = _json_decode_db(row.get("lines"))
    return {
        "id": int(row["id"]),
        "ref_name": str(row.get("ref_name") or ""),
        "gen_settings": gs if isinstance(gs, dict) else {},
        "lines": ln if isinstance(ln, list) else [],
        "creator_admin_id": int(row.get("creator_admin_id") or 0),
    }


def months_spanned_inclusive(start: date, end: date) -> int:
    if end < start:
        return 0
    return (end.year - start.year) * 12 + (end.month - start.month) + 1


def auto_flow_calculate_ranges(gen: dict[str, Any]) -> dict[str, float]:
    start = parse_iso_date_only(gen.get("start_date"))
    end = parse_iso_date_only(gen.get("end_date"))
    if not start or not end or end < start:
        raise ValueError("开始日期或结束日期无效")
    months = months_spanned_inclusive(start, end)
    if months < 1:
        raise ValueError("日期范围内没有可用月份")
    opening = float(to_num_or_null(gen.get("opening_balance")) or 0)
    total_in = float(to_num_or_null(gen.get("total_income")) or 0)
    final_bal = float(to_num_or_null(gen.get("final_balance")) or 0)
    total_out = opening + total_in - final_bal
    if total_out < 0:
        total_out = 0.0
    avg_in = total_in / months
    avg_out = total_out / months
    jitter = 0.18

    def band(avg: float) -> tuple[float, float]:
        if avg <= 0:
            return 0.01, max(0.01, abs(total_in) / max(months, 1))
        lo = max(0.01, avg * (1 - jitter))
        hi = max(lo + 0.01, avg * (1 + jitter))
        return round(lo, 2), round(hi, 2)

    mi_lo, mi_hi = band(avg_in)
    mo_lo, mo_hi = band(avg_out)
    return {
        "months": float(months),
        "derived_total_expense": round(total_out, 2),
        "month_income_min": mi_lo,
        "month_income_max": mi_hi,
        "month_expense_min": mo_lo,
        "month_expense_max": mo_hi,
    }


def _auto_flow_line_ie(line: dict[str, Any]) -> str:
    t = str(line.get("ie_attr") or line.get("收支属性") or "").strip()
    if "支" in t:
        return "支出"
    return "收入"


def _auto_flow_line_selected(line: dict[str, Any]) -> bool:
    if "selected" in line:
        return bool(line.get("selected"))
    if "选择" in line:
        return bool(line.get("选择"))
    return True


def day_allowed_auto_flow(dd: date, *, weekend_ok: bool) -> bool:
    wd = dd.weekday()
    if wd >= 5 and not weekend_ok:
        return False
    return True


def valid_days_in_month(y: int, m: int, start_d: date, end_d: date, weekend_ok: bool) -> list[int]:
    _, last = calendar.monthrange(y, m)
    days: list[int] = []
    for d in range(1, last + 1):
        dd = date(y, m, d)
        if dd < start_d or dd > end_d:
            continue
        if day_allowed_auto_flow(dd, weekend_ok=weekend_ok):
            days.append(d)
    return days


def _auto_flow_tx_amount_to_decimal(raw: Any) -> Optional[Decimal]:
    if raw is None:
        return None
    s = str(raw).strip().replace(",", "")
    if not s:
        return None
    try:
        return Decimal(s).quantize(Decimal("0.01"))
    except Exception:
        return None


def _auto_flow_gen_money_decimal(val: Any) -> Decimal:
    if val is None or val == "":
        return Decimal("0").quantize(Decimal("0.01"))
    try:
        return Decimal(str(val).strip().replace(",", "")).quantize(Decimal("0.01"))
    except Exception:
        return Decimal("0").quantize(Decimal("0.01"))


def _auto_flow_min_step_for_dp(dp: int) -> Decimal:
    ip = int(dp)
    if ip == -2:
        return Decimal(100)
    if ip == -1:
        return Decimal(10)
    if ip <= 0:
        return Decimal(1)
    return (Decimal(10) ** (-ip)).quantize(Decimal(10) ** (-min(ip, 6)), rounding=ROUND_HALF_UP)


def _auto_flow_min_positive_mag_for_dp(dp: int) -> Decimal:
    return _auto_flow_min_step_for_dp(dp)


def _auto_flow_build_proportional_mags(entries: list[tuple[int, Decimal, int]], tgt: Decimal) -> list[Decimal]:
    sum_w = sum((w for _, w, _ in entries), Decimal("0"))
    if sum_w <= 0:
        raise ValueError("模板随机金额为 0，无法分摊目标总额")
    mags: list[Decimal] = []
    for _, w, dp in entries:
        ideal = tgt * (w / sum_w)
        mag = quantize_auto_flow_abs_amount(ideal, dp)
        if mag <= 0:
            mag = quantize_auto_flow_abs_amount(_auto_flow_min_positive_mag_for_dp(dp), dp)
        mags.append(mag)
    return mags


def _auto_flow_step_adjust_mags(entries: list[tuple[int, Decimal, int]], mags: list[Decimal], tgt: Decimal) -> Decimal:
    """用最小步长微调各笔金额，使合计逼近 tgt。"""
    diff = tgt - sum(mags)
    eps = Decimal("0.0001")
    guard = 0
    while abs(diff) > eps and guard < 80000:
        guard += 1
        progressed = False
        order = sorted(range(len(entries)), key=lambda i: mags[i], reverse=True)
        for i in order:
            _, _, dp = entries[i]
            step = _auto_flow_min_step_for_dp(dp)
            if diff > 0:
                new_mag = quantize_auto_flow_abs_amount(mags[i] + step, dp)
                if new_mag > mags[i]:
                    diff -= new_mag - mags[i]
                    mags[i] = new_mag
                    progressed = True
                    break
            else:
                if mags[i] >= step:
                    new_mag = quantize_auto_flow_abs_amount(mags[i] - step, dp)
                    if new_mag < mags[i]:
                        diff += mags[i] - new_mag
                        mags[i] = new_mag
                        progressed = True
                        break
        if not progressed:
            break
    return tgt - sum(mags)


def _auto_flow_residual_best_fit(entries: list[tuple[int, Decimal, int]], mags: list[Decimal], tgt: Decimal) -> Decimal:
    """
    对无法用「整步长」消除的小数差额，在单笔上直接加上剩余 diff 再按该行粒度量化，
    迭代选择使 |tgt−合计| 最小的那一笔（解决「全是整数元但目标含分」等情形）。
    """
    eps = Decimal("0.005")
    for _ in range(2500):
        diff = tgt - sum(mags)
        if abs(diff) <= eps:
            return diff
        best_i = -1
        best_cand: Optional[Decimal] = None
        best_score: Optional[Decimal] = None
        for i in range(len(entries)):
            _, _, dp = entries[i]
            cand = quantize_auto_flow_abs_amount(mags[i] + diff, dp)
            if cand < 0:
                continue
            new_sum = sum(mags[j] if j != i else cand for j in range(len(mags)))
            score = (tgt - new_sum).copy_abs()
            if best_score is None or score < best_score - Decimal("0.000001"):
                best_score = score
                best_i = i
                best_cand = cand
        if best_i < 0 or best_cand is None or best_score is None:
            break
        if best_score >= abs(diff) - Decimal("0.000001"):
            break
        mags[best_i] = best_cand
    return tgt - sum(mags)


def _auto_flow_run_allocate_pipeline(
    entries: list[tuple[int, Decimal, int]], tgt: Decimal
) -> tuple[list[Decimal], Decimal]:
    mags = _auto_flow_build_proportional_mags(entries, tgt)
    _auto_flow_step_adjust_mags(entries, mags, tgt)
    diff = _auto_flow_residual_best_fit(entries, mags, tgt)
    return mags, diff


def _auto_flow_allocate_mags_to_target(entries: list[tuple[int, Decimal, int]], tgt: Decimal) -> list[Decimal]:
    """
    返回已分摊的绝对值金额列表。若无法在小数/粒度约束下贴合 tgt，再尝试将 tgt 四舍五入到整数元（仅当最粗粒度≥1 元时）。
    """
    if not entries:
        if tgt > Decimal("0.005"):
            raise ValueError("模板中缺少收入或支出明细，无法按「总进账 / 支出合计」分摊金额")
        return []
    tgt = tgt.copy_abs().quantize(Decimal("0.01"))
    if tgt <= 0:
        return [Decimal(0) for _ in entries]
    mags, diff = _auto_flow_run_allocate_pipeline(entries, tgt)
    if abs(diff) > Decimal("0.05"):
        min_step = min(_auto_flow_min_step_for_dp(e[2]) for e in entries)
        if min_step >= Decimal("1"):
            tgt_i = quantize_auto_flow_abs_amount(tgt, 0)
            if abs(tgt_i - tgt) <= Decimal("1.5"):
                m2, d2 = _auto_flow_run_allocate_pipeline(entries, tgt_i)
                if abs(d2) < abs(diff):
                    mags, diff = m2, d2
    if abs(diff) > Decimal("0.05"):
        finest = min(_auto_flow_min_step_for_dp(e[2]) for e in entries)
        hint = (
            "若明细均为整数元（小数位 0），请把「总进账/目标余额」调到使「期初+进账−余额」为整数元，"
            "或将至少一行金额小数位设为 1～2 以容纳分位差额。"
        )
        if finest >= Decimal("1"):
            hint = "当前模板最细粒度为整数元，目标总额含分位时难以完全贴合；" + hint
        raise ValueError(
            f"无法在各行小数位规则下精确匹配目标金额（剩余差额约 {diff.quantize(Decimal('0.01'))}）。{hint}"
        )
    return mags


def _auto_flow_allocate_to_target_quantized(
    bodies: list[dict[str, Any]],
    entries: list[tuple[int, Decimal, int]],
    target: Decimal,
    *,
    sign: int,
) -> None:
    """
    按权重比例把 target 分摊到多笔流水，金额遵守各行小数位/粒度，并修正合计误差。
    sign=1 为收入（正），sign=-1 为支出（负）。
    """
    if not entries:
        if target > Decimal("0.005"):
            raise ValueError("模板中缺少收入或支出明细，无法按「总进账 / 支出合计」分摊金额")
        return
    tgt = target.copy_abs().quantize(Decimal("0.01"))
    if tgt <= 0:
        for idx, _, _ in entries:
            bodies[idx]["tx_amount"] = decimal_amount_to_plain_str(Decimal(0))
        return
    mags = _auto_flow_allocate_mags_to_target(entries, tgt)
    for k, (idx, _, _) in enumerate(entries):
        signed = mags[k] if sign > 0 else -mags[k]
        bodies[idx]["tx_amount"] = decimal_amount_to_plain_str(signed)


def _auto_flow_sum_positive_income_from_bodies(bodies: list[dict[str, Any]]) -> Decimal:
    s = Decimal("0")
    for b in bodies:
        a = _auto_flow_tx_amount_to_decimal(b.get("tx_amount"))
        if a is not None and a > 0:
            s += a
    return s.quantize(Decimal("0.01"))


def _auto_flow_apply_total_targets(
    bodies: list[dict[str, Any]], target_inc: Decimal, opening_dec: Decimal, target_fin: Decimal
) -> None:
    """
    先按「总进账」分摊收入；再用 期初+实际收入合计−目标余额 计算应分摊的支出总额，
    保证生成后滚算期末与「最后卡上余额」一致（收入侧允许在粒度约束下与填写值相差至多约 1 元）。
    """
    inc_entries: list[tuple[int, Decimal, int]] = []
    exp_entries: list[tuple[int, Decimal, int]] = []
    for i, b in enumerate(bodies):
        amt = _auto_flow_tx_amount_to_decimal(b.get("tx_amount"))
        if amt is None:
            continue
        dp_raw = b.get("_auto_flow_dp")
        try:
            dp = int(dp_raw) if dp_raw is not None else 0
        except (TypeError, ValueError):
            dp = 0
        if amt > 0:
            w = quantize_auto_flow_abs_amount(amt, dp)
            if w <= 0:
                w = _auto_flow_min_positive_mag_for_dp(dp)
            inc_entries.append((i, w, dp))
        elif amt < 0:
            w = quantize_auto_flow_abs_amount(abs(amt), dp)
            if w <= 0:
                w = _auto_flow_min_positive_mag_for_dp(dp)
            exp_entries.append((i, w, dp))
    _auto_flow_allocate_to_target_quantized(bodies, inc_entries, target_inc, sign=1)
    actual_inc = _auto_flow_sum_positive_income_from_bodies(bodies)
    target_exp_mag = (opening_dec + actual_inc - target_fin).quantize(Decimal("0.01"))
    if target_exp_mag < -Decimal("0.005"):
        raise ValueError(
            "按实际生成的收入合计计算，支出目标为负：请提高总进账或降低「最后卡上余额」，或放宽收入模板金额上限。"
        )
    if target_exp_mag < 0:
        target_exp_mag = Decimal("0").quantize(Decimal("0.01"))
    _auto_flow_allocate_to_target_quantized(bodies, exp_entries, target_exp_mag, sign=-1)


def _auto_flow_greedy_order_for_non_negative(bodies: list[dict[str, Any]], opening_dec: Decimal) -> list[dict[str, Any]]:
    eps = Decimal("0.005")
    remaining = list(bodies)
    for i, b in enumerate(remaining):
        b["_auto_flow_ord"] = i
    ordered: list[dict[str, Any]] = []
    bal = opening_dec.quantize(Decimal("0.01"))
    while remaining:
        feasible: list[tuple[dict[str, Any], Decimal]] = []
        for b in remaining:
            a = _auto_flow_tx_amount_to_decimal(b.get("tx_amount"))
            if a is None:
                continue
            if bal + a >= -eps:
                feasible.append((b, a))
        if not feasible:
            raise ValueError(
                "自动生成：按目标金额生成的收支即使重排顺序仍可能在滚算中出现负余额。"
                "请降低支出模板频率或单笔上限，或提高总进账/期初余额。"
            )
        feasible.sort(key=lambda x: (x[1], -int(x[0].get("_auto_flow_ord") or 0)), reverse=True)
        pick, a = feasible[0]
        ordered.append(pick)
        bal = (bal + a).quantize(Decimal("0.01"))
        remaining.remove(pick)
    for b in ordered:
        b.pop("_auto_flow_ord", None)
    rb = opening_dec.quantize(Decimal("0.01"))
    for b in ordered:
        a = _auto_flow_tx_amount_to_decimal(b.get("tx_amount"))
        if a is None:
            continue
        rb = (rb + a).quantize(Decimal("0.01"))
        if rb < -eps:
            raise ValueError("自动生成顺序校验失败，请调整模板后重试")
    return ordered


def _auto_flow_assign_random_monotonic_datetimes(
    rng: random.Random, ordered: list[dict[str, Any]], start: date, end: date
) -> None:
    """
    在 [start,end] 内生成严格递增的记账时间：先取大致均匀的骨架，再叠加随机抖动，
    同日多笔、间隔长短不一，避免「等间隔铺满」的不真实观感。
    """
    n = len(ordered)
    if n <= 0:
        return
    t_lo = datetime(start.year, start.month, start.day, 5, 40, rng.randint(0, 59))
    t_hi = datetime(end.year, end.month, end.day, 23, 25, rng.randint(0, 59))
    if t_hi <= t_lo:
        t_hi = t_lo + timedelta(hours=8)
    raw = _admin_spread_datetimes_evenly(n, t_lo, t_hi)
    noisy: list[datetime] = []
    for dt in raw:
        noisy.append(
            dt
            + timedelta(
                days=rng.randint(-8, 8),
                hours=rng.randint(-14, 14),
                minutes=rng.randint(0, 59),
                seconds=rng.randint(0, 59),
            )
        )
    noisy.sort()
    last = t_lo - timedelta(seconds=1)
    out: list[datetime] = []
    for dt in noisy:
        if dt < t_lo:
            dt = t_lo + timedelta(seconds=rng.randint(0, 9000))
        if dt > t_hi:
            dt = t_hi - timedelta(seconds=rng.randint(0, 7200))
        if dt <= last:
            gap_hi = max(60, int((t_hi - last).total_seconds()) - 30)
            gap_hi = min(gap_hi, 86400 * 3)
            dt = last + timedelta(seconds=rng.randint(40, max(120, gap_hi)))
        if dt > t_hi:
            dt = last + timedelta(seconds=rng.randint(120, 3600))
            if dt > t_hi:
                dt = t_hi
        if dt <= last:
            dt = min(t_hi, last + timedelta(seconds=45))
        out.append(dt)
        last = dt
    for b, dt_new in zip(ordered, out):
        b["tx_datetime"] = dt_new.strftime("%Y-%m-%d %H:%M:%S")


def _auto_flow_random_abs_amount(rng: random.Random, min_amt: float, max_amt: float) -> float:
    """在 [min,max] 内生成偏真实分布：时而偏小、时而偏大、偶有中段集中，避免金额过于接近。"""
    if max_amt <= min_amt:
        return float(min_amt)
    span = max_amt - min_amt
    roll = rng.random()
    if roll < 0.2:
        t = rng.betavariate(1.15, 2.9)
    elif roll < 0.38:
        t = rng.betavariate(2.9, 1.15)
    elif roll < 0.62:
        t = rng.random()
    elif roll < 0.82:
        t = 0.38 + 0.32 * rng.random()
    else:
        t = min(1.0, max(0.0, 0.5 + rng.gauss(0, 0.22)))
    return float(min_amt + span * t)


def plan_auto_flow_transactions(gen: dict[str, Any], lines: list[Any]) -> list[dict[str, Any]]:
    start = parse_iso_date_only(gen.get("start_date"))
    end = parse_iso_date_only(gen.get("end_date"))
    if not start or not end or end < start:
        raise ValueError("开始日期或结束日期无效")
    rng = random.SystemRandom()
    bodies: list[dict[str, Any]] = []
    gen_ord = 0
    raw_lines = lines if isinstance(lines, list) else []
    active_lines = [ln for ln in raw_lines if isinstance(ln, dict) and _auto_flow_line_selected(ln)]
    if not active_lines:
        raise ValueError("请至少勾选一条模板明细（「选择」）")

    y, m = start.year, start.month
    while True:
        first = date(y, m, 1)
        if first > end:
            break
        _, last_day = calendar.monthrange(y, m)
        month_start = date(y, m, 1)
        month_end = date(y, m, last_day)
        clip_start = max(start, month_start)
        clip_end = min(end, month_end)

        for line in active_lines:
            ie = _auto_flow_line_ie(line)
            min_amt = float(to_num_or_null(line.get("min_amount")) or 0)
            max_amt = float(to_num_or_null(line.get("max_amount")) or 0)
            if max_amt < min_amt:
                min_amt, max_amt = max_amt, min_amt
            if max_amt <= 0:
                max_amt = max(min_amt, 1.0)
            dp = parse_line_decimal_places(line)
            sh = int(to_num_or_null(line.get("start_hour")) or 9)
            eh = int(to_num_or_null(line.get("end_hour")) or 18)
            if eh < sh:
                sh, eh = eh, sh
            mf = int(to_num_or_null(line.get("monthly_frequency")) or 0)
            occ = mf if mf > 0 else rng.randint(2, 9)
            weekend_ok = bool(line.get("weekend_ok")) if "weekend_ok" in line else bool(line.get("周六日交易"))

            vd = valid_days_in_month(y, m, clip_start, clip_end, weekend_ok=weekend_ok)
            if not vd:
                vd = list(range(clip_start.day, clip_end.day + 1))
                if not vd:
                    continue
            # 少数日期集中出账/入账，易出现同日多笔与空白日，减少「每天一笔」的规律感
            pool_n = max(1, min(len(vd), rng.randint(max(1, occ // 3), max(2, occ + 2))))
            pool_n = min(pool_n, len(vd))
            day_pool = rng.sample(vd, k=pool_n) if len(vd) >= pool_n else list(vd)

            tx_type = str(line.get("tx_type") or line.get("交易类型") or "").strip() or "转账"
            currency = str(line.get("currency") or line.get("交易币种") or "人民币").strip() or "人民币"
            cp_name = str(line.get("counterparty_name") or line.get("对手方户名") or "").strip()
            cp_acc = str(line.get("counterparty_account") or line.get("对手方账户") or "").strip()
            cp_bank = str(line.get("counterparty_bank") or line.get("对手银行") or "").strip()
            remark = str(line.get("remark") or line.get("附言") or "").strip()
            channel = str(line.get("channel") or line.get("交易方式") or "").strip()

            for _ in range(occ):
                day_pick = rng.choice(day_pool)
                hh = rng.randint(int(sh), int(eh))
                mm = rng.randint(0, 59)
                ss = min(59, rng.randint(0, 59))
                raw_amt = _auto_flow_random_abs_amount(rng, min_amt, max_amt)
                mag_dec = quantize_auto_flow_abs_amount(raw_amt, dp)
                if mag_dec <= 0:
                    mag_dec = quantize_auto_flow_abs_amount(min_amt if min_amt > 0 else 1.0, dp)
                signed_dec = -mag_dec if ie == "支出" else mag_dec
                amt_plain = decimal_amount_to_plain_str(signed_dec)
                dt_str = f"{y:04d}-{m:02d}-{day_pick:02d} {hh:02d}:{mm:02d}:{ss:02d}"
                summ = tx_type
                dw_cd = "2" if ie == "支出" else "1"
                bodies.append(
                    {
                        "tx_datetime": dt_str,
                        "tx_amount": amt_plain,
                        "_auto_flow_dp": dp,
                        "tx_type": tx_type,
                        "summ": summ,
                        "currency": currency,
                        "counterparty_name": cp_name or None,
                        "counterparty_account": cp_acc or None,
                        "counterparty_bank": cp_bank or None,
                        "remark": remark or None,
                        "channel": channel or None,
                        "dwFlagCode": dw_cd,
                        "incmEpnTpCd": dw_cd,
                        "_preserve_wallet_platform_overrides": True,
                        "_gen_ord": gen_ord,
                    }
                )
                gen_ord += 1

        if m == 12:
            y += 1
            m = 1
        else:
            m += 1

    bodies.sort(key=lambda b: (str(b.get("tx_datetime") or ""), int(b.get("_gen_ord") or 0)))

    opening_dec = _auto_flow_gen_money_decimal(gen.get("opening_balance"))
    ti_raw = gen.get("total_income")
    fb_raw = gen.get("final_balance")
    constrain_totals = ti_raw not in (None, "") and fb_raw not in (None, "")

    if constrain_totals:
        target_inc = _auto_flow_gen_money_decimal(ti_raw)
        target_fin = _auto_flow_gen_money_decimal(fb_raw)
        target_exp_mag = (opening_dec + target_inc - target_fin).quantize(Decimal("0.01"))
        if target_exp_mag < -Decimal("0.005"):
            raise ValueError(
                "目标「最后卡上余额」过高：期初余额加总进账仍不足以达到该期末余额，请降低目标余额或提高总进账/期初"
            )
        if target_exp_mag < 0:
            target_exp_mag = Decimal("0").quantize(Decimal("0.01"))
        _auto_flow_apply_total_targets(bodies, target_inc, opening_dec, target_fin)
        sum_mov = Decimal("0")
        for b in bodies:
            a = _auto_flow_tx_amount_to_decimal(b.get("tx_amount"))
            if a is not None:
                sum_mov += a
        closing_chk = (opening_dec + sum_mov).quantize(Decimal("0.01"))
        if abs(closing_chk - target_fin) > Decimal("0.05"):
            raise ValueError(
                f"收支合计与目标期末余额不一致（推算期末 {closing_chk}，目标 {target_fin}），请检查模板或联系管理员"
            )
        ordered = _auto_flow_greedy_order_for_non_negative(bodies, opening_dec)
        _auto_flow_assign_random_monotonic_datetimes(rng, ordered, start, end)
        for b in ordered:
            b.pop("_gen_ord", None)
        return ordered

    for b in bodies:
        b.pop("_gen_ord", None)
    return bodies


async def roll_balances_for_person_from_opening(person_id: int, opening: float) -> float:
    rows = await query_all(
        f"SELECT id, tx_amount FROM `{TX_TABLE}` WHERE person_id=%s ORDER BY tx_datetime ASC, id ASC",
        (int(person_id),),
    )
    opening_dec = Decimal(str(float(opening))).quantize(Decimal("0.01"))
    running = opening_dec
    updates: list[tuple[int, float, str]] = []
    for r in rows:
        raw_amt = r.get("tx_amount")
        n = to_num_or_null(raw_amt)
        if n is None:
            continue
        amt = Decimal(str(n)).quantize(Decimal("0.01"))
        running = (running + amt).quantize(Decimal("0.01"))
        bal_str = to_amt_str(float(running))
        updates.append((int(r["id"]), float(running), bal_str))
    await _bulk_update_tx_balances_case(int(person_id), updates, float(running))
    return float(running)


def _admin_tx_amount_dec_from_row(r: dict[str, Any]) -> Decimal:
    n = to_num_or_null(r.get("tx_amount"))
    if n is None:
        raise ValueError(f"流水 id={r.get('id')} 缺少有效金额 tx_amount")
    return Decimal(str(n)).quantize(Decimal("0.01"))


def _admin_parse_row_tx_datetime(val: Any) -> datetime:
    s = cell_to_tx_datetime_str(val)
    if not s:
        return datetime(2000, 1, 1, 8, 0, 0)
    try:
        return datetime.strptime(s[:19], "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return datetime(2000, 1, 1, 8, 0, 0)


def _admin_spread_datetimes_evenly(n: int, t_min: datetime, t_max: datetime) -> list[datetime]:
    """在 [t_min, t_max] 内生成严格递增的时间序列（必要时拉长区间）。"""
    if n <= 0:
        return []
    if n == 1:
        return [t_min]
    if t_max <= t_min:
        t_max = t_min + timedelta(seconds=max(120, n * 2))
    span_sec = (t_max - t_min).total_seconds()
    span_sec = max(span_sec, float((n - 1) * 2))
    out: list[datetime] = []
    last: Optional[datetime] = None
    for i in range(n):
        frac = i / (n - 1)
        t = t_min + timedelta(seconds=frac * span_sec)
        if last is not None and t <= last:
            t = last + timedelta(seconds=1)
        out.append(t)
        last = t
    return out


async def admin_reorder_profile_transactions_non_negative_balance(person_id: int) -> dict[str, Any]:
    """
    在「仅调整各笔流水的记账时间顺序」的前提下，尝试消除滚算过程中的负余额。
    使用贪心：每一步在可行集合中选金额最大的一笔先进账（收入优先），失败则返回错误提示。
    """
    pid = int(person_id)
    rows = await query_all(
        f"""SELECT id, tx_datetime, tx_amount, account_balance, accBal
            FROM `{TX_TABLE}` WHERE person_id=%s""",
        (pid,),
    )
    if not rows:
        return {"ok": True, "changed": False, "message": "该资料暂无流水", "updated": 0}

    sorted_rows = sorted(
        rows,
        key=lambda r: (_admin_parse_row_tx_datetime(r.get("tx_datetime")), int(r["id"])),
    )
    total_delta = sum((_admin_tx_amount_dec_from_row(r) for r in sorted_rows), Decimal("0"))

    opening_dec: Optional[Decimal] = None
    eb0 = _effective_post_tx_balance_from_row(sorted_rows[0])
    if eb0 is not None:
        opening_dec = (eb0 - _admin_tx_amount_dec_from_row(sorted_rows[0])).quantize(Decimal("0.01"))
    if opening_dec is None:
        prow = await query_one(
            f"SELECT `{PROFILE_COL_BALANCE}` AS ba FROM `{PROFILE_TABLE}` WHERE id=%s LIMIT 1",
            (pid,),
        )
        if prow is not None and to_num_or_null(prow.get("ba")) is not None:
            closing_dec = Decimal(str(to_num_or_null(prow.get("ba")))).quantize(Decimal("0.01"))
            opening_dec = (closing_dec - total_delta).quantize(Decimal("0.01"))
    if opening_dec is None:
        opening_dec = Decimal("0").quantize(Decimal("0.01"))

    eps = Decimal("0.005")
    running_chk = opening_dec
    min_run = running_chk
    for r in sorted_rows:
        running_chk = (running_chk + _admin_tx_amount_dec_from_row(r)).quantize(Decimal("0.01"))
        min_run = min(min_run, running_chk)
    has_negative = min_run < -eps

    if not has_negative:
        return {
            "ok": True,
            "changed": False,
            "message": "当前按时间滚算不会出现负余额，无需调整顺序。",
            "updated": 0,
        }

    remaining = list(sorted_rows)
    ordered: list[dict[str, Any]] = []
    bal = opening_dec
    while remaining:
        feasible: list[tuple[dict[str, Any], Decimal]] = []
        for r in remaining:
            a = _admin_tx_amount_dec_from_row(r)
            if bal + a >= -eps:
                feasible.append((r, a))
        if not feasible:
            return {
                "ok": False,
                "error": "无法通过调整时间顺序消除负余额：按当前期初滚算，支出相对于可用收入过多，请调低支出模板金额或增加收入后再试。",
                "changed": False,
                "updated": 0,
            }
        feasible.sort(key=lambda x: (x[1], -int(x[0]["id"])), reverse=True)
        pick, a = feasible[0]
        ordered.append(pick)
        bal = (bal + a).quantize(Decimal("0.01"))
        remaining.remove(pick)

    orig_ids = [int(r["id"]) for r in sorted_rows]
    new_ids = [int(r["id"]) for r in ordered]
    order_changed = orig_ids != new_ids

    rb = opening_dec
    for r in ordered:
        rb = (rb + _admin_tx_amount_dec_from_row(r)).quantize(Decimal("0.01"))
        if rb < -eps:
            return {
                "ok": False,
                "error": "重排后校验仍出现负余额，请检查数据异常或稍后重试。",
                "changed": False,
                "updated": 0,
            }

    dts = [_admin_parse_row_tx_datetime(r.get("tx_datetime")) for r in sorted_rows]
    t_min, t_max = min(dts), max(dts)
    new_dts = _admin_spread_datetimes_evenly(len(ordered), t_min, t_max)

    for r, dt_new in zip(ordered, new_dts):
        dt_str = dt_new.strftime("%Y-%m-%d %H:%M:%S")
        fm = fmt_ymd_hms(dt_str)
        ymd = fm.get("ymd") or ""
        ymdhms = fm.get("ymdhms") or ""
        await execute(
            f"""UPDATE `{TX_TABLE}` SET tx_datetime=%s, `txDate`=%s, `txTime`=%s WHERE id=%s AND person_id=%s""",
            (dt_str, ymd, ymdhms, int(r["id"]), pid),
        )

    final_bal = await roll_balances_for_person_from_opening(pid, float(opening_dec))
    return {
        "ok": True,
        "changed": True,
        "message": f"已按避免负余额的顺序调整 {len(ordered)} 笔流水的记账时间并重新滚算余额。",
        "updated": len(ordered),
        "opening_balance_used": str(opening_dec),
        "final_balance_after": final_bal,
        "order_changed": order_changed,
    }


async def admin_auto_flow_generate(
    admin: dict[str, Any],
    profile_id: int,
    body: dict[str, Any],
) -> dict[str, Any]:
    if not await admin_profile_accessible(admin, int(profile_id)):
        raise ValueError("无权操作该资料流水")
    tpl_id = to_id_or_null(body.get("template_id"))
    gen: dict[str, Any]
    raw_lines: list[Any]
    if tpl_id:
        tpl = await get_auto_flow_template_by_id(int(tpl_id))
        if not tpl:
            raise ValueError("模板不存在")
        if not await admin_auto_flow_template_accessible(admin, int(tpl_id)):
            raise ValueError("无权使用该模板")
        gen = dict(tpl.get("gen_settings") or {})
        raw_lines = list(tpl.get("lines") or [])
        og = body.get("gen_settings")
        ol = body.get("lines")
        if isinstance(og, dict):
            gen.update(og)
        if isinstance(ol, list):
            raw_lines = ol
    else:
        gen = dict(body.get("gen_settings") or {})
        raw_lines = list(body.get("lines") or [])
    if not gen.get("start_date") or not gen.get("end_date"):
        raise ValueError("gen_settings 需提供 start_date / end_date")
    if not raw_lines:
        raise ValueError("模板明细 lines 不能为空")

    clear_existing = bool(body.get("clear_existing"))
    preserve_dc = to_bool_int(body.get("preserve_device_collected")) == 1
    if clear_existing:
        await admin_delete_transactions_for_profile(int(profile_id), preserve_device_collected=preserve_dc)

    opening = float(to_num_or_null(gen.get("opening_balance")) or 0)
    planned = plan_auto_flow_transactions(gen, raw_lines)
    if len(planned) > AUTO_FLOW_MAX_GENERATE:
        raise ValueError(f"生成笔数过多（{len(planned)}），上限 {AUTO_FLOW_MAX_GENERATE}")

    inserted = 0
    for pb in planned:
        await admin_insert_transaction(int(profile_id), pb)
        inserted += 1

    final_bal = await roll_balances_for_person_from_opening(int(profile_id), opening)
    return {"inserted": inserted, "opening_balance_used": opening, "final_balance_after": final_bal}


def _auto_flow_tx_row_text(row: dict[str, Any], *keys: str) -> str:
    if not row:
        return ""
    lk = {str(k).lower(): k for k in row.keys()}
    for want in keys:
        ck = lk.get(str(want).lower())
        if ck is None:
            continue
        v = row.get(ck)
        if v is None:
            continue
        s = str(v).strip()
        if s:
            return s
    return ""


def _auto_flow_tx_datetime_hour(tx_dt: Any) -> Optional[int]:
    if tx_dt is None:
        return None
    if isinstance(tx_dt, datetime):
        return int(tx_dt.hour)
    s = str(tx_dt).strip()
    m = re.search(r"[ T](\d{1,2}):", s)
    if m:
        return int(m.group(1)) % 24
    return None


def _infer_auto_flow_decimal_places(abs_amounts: list[float]) -> int:
    """根据样本金额推断模板小数位（0 或 2；细则仍可由用户在表格里改成 -2～3）。"""
    if not abs_amounts:
        return 0
    non_int = 0
    for a in abs_amounts:
        try:
            d = Decimal(str(a))
            if d != d.quantize(Decimal("1"), rounding=ROUND_HALF_UP):
                non_int += 1
        except Exception:
            non_int += 1
    if non_int == 0:
        return 0
    if non_int / len(abs_amounts) < 0.22:
        return 0
    return 2


async def build_auto_flow_line_presets_for_admin(admin: dict[str, Any], scan_limit: int) -> dict[str, Any]:
    """
    从当前管理员可访问的流水表中抽样，按「交易类型 + 解析后渠道 + 收支」聚类，
    生成可一键插入自动生成流水明细表的 line 模板。
    """
    role = int(admin.get("role") or 0)
    aid = int(admin["id"])
    lim = max(80, min(int(scan_limit or 2500), 8000))
    sel_cols = (
        "t.id, t.tx_datetime, t.tx_type, t.currency, t.tx_amount, "
        "t.counterparty_name, t.counterparty_account, t.counterparty_bank, "
        "t.remark, t.channel, t.summ, t.merDesc, t.txOpsName, t.txOpsAccno, t.txRemark, t.`chnlKindCode`"
    )
    if role == ADMIN_ROLE_SUPER:
        rows = await query_all(
            f"SELECT {sel_cols} FROM `{TX_TABLE}` t ORDER BY t.id DESC LIMIT %s",
            (lim,),
        )
    else:
        rows = await query_all(
            f"""SELECT {sel_cols} FROM `{TX_TABLE}` t
            INNER JOIN `{PROFILE_TABLE}` p ON p.id = t.person_id
            WHERE p.`{PROFILE_COL_CREATOR_ADMIN}` = %s
            ORDER BY t.id DESC LIMIT %s""",
            (aid, lim),
        )
    rows = rows or []
    groups: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    group_order: list[tuple[str, str, str]] = []
    for r in rows:
        tx_t = _auto_flow_tx_row_text(r, "tx_type", "summ") or "其它"
        summ_s = _auto_flow_tx_row_text(r, "summ")
        remark_s = _auto_flow_tx_row_text(r, "remark", "txRemark")
        ch_raw = _auto_flow_tx_row_text(r, "channel")
        ch_cd = _auto_flow_tx_row_text(r, "chnlKindCode")
        channel_cn = resolve_transaction_channel_cn(
            summ=summ_s or tx_t,
            tx_type=tx_t,
            remark=remark_s,
            channel=ch_raw,
            chnl_kind_code=ch_cd,
        ).strip() or (ch_raw or "其它")
        amt_n = to_num_or_null(r.get("tx_amount"))
        if amt_n is None:
            continue
        ie = "支出" if float(amt_n) < 0 else "收入"
        tx_key = tx_t[:120]
        ch_key = channel_cn[:48]
        gk = (tx_key, ch_key, ie)
        if gk not in groups:
            group_order.append(gk)
            groups[gk] = []
        groups[gk].append(dict(r))

    presets: list[dict[str, Any]] = []
    for gk in group_order:
        grp = groups.get(gk) or []
        if not grp:
            continue
        tx_key, channel_cn, ie = gk
        sample = grp[0]
        abs_amts: list[float] = []
        for x in grp:
            n = to_num_or_null(x.get("tx_amount"))
            if n is None:
                continue
            abs_amts.append(abs(float(n)))
        abs_amts = [a for a in abs_amts if a > 1e-9]
        if not abs_amts:
            mn_f, mx_f = 1.0, 5000.0
        else:
            lo, hi = min(abs_amts), max(abs_amts)
            mn_f = max(0.01, round(lo * 0.82, 2))
            mx_f = max(mn_f * 1.2, round(hi * 1.18, 2))
        dp = _infer_auto_flow_decimal_places(abs_amts)
        hrs = [_auto_flow_tx_datetime_hour(x.get("tx_datetime")) for x in grp]
        hrs = [h for h in hrs if h is not None]
        if hrs:
            sh, eh = min(hrs), max(hrs)
            sh = max(6, min(22, sh))
            eh = max(6, min(22, eh))
            if eh < sh:
                sh, eh = eh, sh
            if eh - sh < 2:
                sh = max(6, sh - 1)
                eh = min(22, eh + 1)
        else:
            sh, eh = 9, 18
        cp_name = _auto_flow_tx_row_text(sample, "counterparty_name", "txOpsName")
        cp_acc = _auto_flow_tx_row_text(sample, "counterparty_account", "txOpsAccno")
        cp_bank = _auto_flow_tx_row_text(sample, "counterparty_bank")
        remark_s = _auto_flow_tx_row_text(sample, "remark", "txRemark", "merDesc")
        curr = _auto_flow_tx_row_text(sample, "currency") or "人民币"
        line: dict[str, Any] = {
            "selected": True,
            "ie_attr": ie,
            "min_amount": str(mn_f),
            "max_amount": str(mx_f),
            "decimal_places": dp,
            "start_hour": int(sh),
            "end_hour": int(eh),
            "monthly_frequency": 0,
            "holiday_ok": True,
            "weekend_ok": ie == "支出",
            "tx_type": tx_key,
            "currency": curr,
            "counterparty_name": cp_name,
            "counterparty_account": cp_acc,
            "counterparty_bank": cp_bank,
            "remark": remark_s,
            "channel": channel_cn,
        }
        label = tx_key if channel_cn in ("", tx_key, "其它") else f"{tx_key} · {channel_cn}"
        presets.append(
            {
                "shortcut_label": label,
                "line": line,
                "sample_count": len(grp),
                "tx_type": tx_key,
                "channel_resolved": channel_cn,
                "ie_attr": ie,
            }
        )

    presets.sort(key=lambda x: int(x.get("sample_count") or 0), reverse=True)
    total_groups = len(presets)
    max_out = 48
    return {"presets": presets[:max_out], "scanned": len(rows), "groups": total_groups}


async def admin_delete_transactions_for_profile(
    profile_id: int, *, preserve_device_collected: bool
) -> tuple[int, int]:
    """
    清空某资料下流水。preserve_device_collected=True 时保留 device_collected=1（设备采集入库）的记录。
    返回 (删除条数, 保留的采集条数)。
    """
    pid = int(profile_id)
    if preserve_device_collected:
        deleted = await execute_rowcount(
            f"DELETE FROM `{TX_TABLE}` WHERE person_id=%s AND COALESCE(`{TX_COL_DEVICE_COLLECTED}`, 0)=0",
            (pid,),
        )
        crow = await query_one(
            f"SELECT COUNT(1) AS c FROM `{TX_TABLE}` WHERE person_id=%s AND COALESCE(`{TX_COL_DEVICE_COLLECTED}`, 0)=1",
            (pid,),
        )
        kept = int(crow.get("c") or 0) if crow else 0
        return deleted, kept
    deleted = await execute_rowcount(f"DELETE FROM `{TX_TABLE}` WHERE person_id=%s", (pid,))
    return deleted, 0


async def admin_insert_transaction(person_id: int, body: dict[str, Any]) -> dict[str, Any]:
    if not str(body.get("tx_datetime") or "").strip():
        raise ValueError("tx_datetime 不能为空")
    if body.get("tx_amount") == "" or body.get("tx_amount") is None or to_num_or_null(body.get("tx_amount")) is None:
        raise ValueError("tx_amount 非法")
    tx_row = build_tx_row_from_body(body, 0, person_id)
    identity = build_stable_tx_identity(tx_row, fmt_ymd_hms(tx_row["tx_datetime"]))
    tx_type_text = "" if tx_row.get("tx_type") is None else str(tx_row.get("tx_type"))
    amt_num_raw = float(to_amt_str(tx_row.get("tx_amount")))
    is_out = amt_num_raw < 0
    is_in_text = "汇入" in tx_type_text or "转入" in tx_type_text
    is_out_text = "汇出" in tx_type_text or "转出" in tx_type_text or "支出" in tx_type_text
    tp_cd = "2" if is_out else "1"
    if is_in_text and not is_out_text:
        tp_cd = "1"
    if is_out_text and not is_in_text:
        tp_cd = "2"
    tx_tp_cd = resolve_incm_epn_tx_tp_cd_from_row(tx_row, tx_type_text, tp_cd)
    global_track_no = normalize_raw_tx_value(body.get("global_busi_track_no") or body.get("globalBusiTrackNo")) or identity[
        "globalBusiTrackNo"
    ]
    prow = await query_one(f"SELECT card_no, person_name FROM `{PROFILE_TABLE}` WHERE id=%s LIMIT 1", (person_id,))
    profile_card = must_non_empty_text((prow or {}).get("card_no"))
    profile_person_name = str((prow or {}).get("person_name") or "").strip()
    if profile_card:
        body["mediumNo"] = profile_card
        body["bkcdMask"] = format_bkcd_mask(profile_card)
    merge_raw_api_fields_from_normalized_import(body, tx_tp_cd, global_track_no, profile_card, profile_person_name)
    raw_model = build_raw_tx_model_from_body(body, global_track_no, tx_tp_cd)
    raw_cols = [d[0] for d in RAW_TX_COLUMN_DEFS]
    insert_cols = [
        "person_id",
        "tx_datetime",
        "tx_type",
        "currency",
        "tx_amount",
        "account_balance",
        "counterparty_name",
        "counterparty_account",
        "counterparty_bank",
        "remark",
        "channel",
        TX_COL_INCM_TX_TP_CD,
        TX_COL_GLOBAL,
        TX_COL_DEVICE_COLLECTED,
    ] + raw_cols
    insert_vals = [
        person_id,
        body.get("tx_datetime") or "1970-01-01 00:00:00",
        body.get("tx_type"),
        body.get("currency") or "人民币",
        to_num_or_null(body.get("tx_amount")) or 0,
        to_num_or_null(body.get("account_balance")),
        body.get("counterparty_name"),
        body.get("counterparty_account"),
        body.get("counterparty_bank"),
        body.get("remark"),
        body.get("channel"),
        tx_tp_cd,
        global_track_no,
        0,
    ] + [raw_model.get(col) for col in raw_cols]
    quoted = ", ".join(f"`{c}`" for c in insert_cols)
    placeholders = ", ".join(["%s"] * len(insert_cols))
    new_tx_id = await execute_insert(
        f"INSERT INTO `{TX_TABLE}` ({quoted}) VALUES ({placeholders})",
        tuple(insert_vals),
    )
    out: dict[str, Any] = {"id": new_tx_id, "incmEpnTxTpCd": tx_tp_cd, "globalBusiTrackNo": global_track_no}
    row = await get_admin_transaction_by_id(int(new_tx_id))
    if row:
        out["data"] = row
    return out


async def run_startup_migrations() -> None:
    try:
        await ensure_device_profile_nullable()
        print("[mock-api] device profile binding nullable ready")
        await ensure_device_collect_tx_column()
        print("[mock-api] device collect flag column ready")
        await ensure_profile_balance_column()
        print("[mock-api] profile balance column ready")
        await ensure_device_auto_sync_tx_column()
        print("[mock-api] device auto-sync transactions flag ready")
        await ensure_admin_accounts_table()
        print("[mock-api] admin accounts table ready")
        await bootstrap_default_super_admin_if_empty()
        await ensure_profile_creator_admin_column()
        print("[mock-api] profile creator_admin_id ready")
        await ensure_profile_mail_send_mode_column()
        print("[mock-api] profile mail_send_mode ready")
        await ensure_profile_cust_lvl_column()
        print("[mock-api] profile cust_lvl ready")
        await ensure_admin_points_mode2_column()
        print("[mock-api] admin points_mode2 ready")
        await ensure_admin_allow_simulate_mail_column()
        print("[mock-api] admin allow_simulate_mail ready")
        await ensure_device_creator_admin_column()
        print("[mock-api] device creator_admin_id ready")
        await ensure_device_remark_column()
        print("[mock-api] device remark column ready")
        await ensure_raw_tx_model_columns()
        print("[mock-api] raw tx model columns ready")
        await ensure_tx_device_collected_column()
        print("[mock-api] tx device_collected column ready")
        await ensure_history_apply_mail_table()
        print("[mock-api] history apply mail table ready")
        await ensure_tx_identity_columns_and_backfill()
        print("[mock-api] tx identity columns ready")
        await drop_legacy_tx_identity_columns()
        print("[mock-api] dropped legacy tx identity columns")
        await ensure_auto_flow_templates_table()
        print("[mock-api] auto flow templates table ready")
    except Exception as e:
        print(f"[mock-api] tx identity init failed: {e}")


async def handle(request: web.Request) -> web.StreamResponse:
    path = request.path
    method = request.method
    print("===============>", method, path, dict(request.query))
    if method == "OPTIONS":
        return json_response({"ok": True})

    if method == "GET" and path == "/healthz":
        try:
            await execute("SELECT 1")
            return json_response({"ok": True, "db": "up", "ts": datetime.now().isoformat()})
        except Exception as e:
            return json_response({"ok": False, "db": "down", "error": str(e)}, 500)

    if method == "GET" and path == "/api/device-info":
        try:
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            device = await load_device_context(did)
            if not device:
                return json_response({"ok": False, "error": f"设备不存在或未启用: did={did}"}, 404)
            return json_response(
                {
                    "ok": True,
                    "data": {
                        "did": device.get("did"),
                        "profile_id": device.get("profile_id"),
                        "profile_bound": is_device_profile_bound(device),
                        "collect_transactions_enabled": to_bool_int(device.get("collect_transactions_enabled")) == 1,
                        "is_active": True,
                    },
                }
            )
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "POST" and path == "/api/device-collect-transactions":
        try:
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            body = await parse_json_body(request)
            payload = body.get("payload")
            if isinstance(payload, str):
                payload = json.loads(payload)
            result = await collect_transactions_by_did(did, payload)
            if not result.get("ok"):
                return json_response(
                    {"ok": False, "error": result.get("error") or "collect failed"},
                    int(result.get("status") or 400),
                )
            return json_response({"ok": True, "data": result})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "POST" and path == "/api/mock-xhx":
        try:
            body = await parse_json_body(request)
            req_msg_id = str(body.get("reqMsgId") or request.headers.get("x-req-msg-id") or "").strip()
            did = pick_did(request)
            begin_date = str(body.get("beginDate") or "").strip()
            dline_date = str(body.get("dlineDate") or "").strip()
            ebank_qry_type_flag_cd = str(body.get("ebankQryTypeFlagCd") or "").strip()
            host_rsp_plain = str(body.get("hostRspPlain") or "")
            tx_tp_raw = body.get("txTpCdList")
            if isinstance(tx_tp_raw, list):
                tx_tp_cd_list = tx_tp_raw
            else:
                tx_tp_cd_list = "" if tx_tp_raw is None else str(tx_tp_raw).strip()
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            device_any = await load_device_context_any(did)
            if device_any and to_bool_int(device_any.get("is_active")) != 1:
                empty = build_empty_xhx_response(str(req_msg_id))
                print("服务端响应===============>", empty)
                return json_response(empty)
            collect_live = await fetch_device_collect_transactions_enabled_by_did(did)
            if device_any is not None and collect_live == 1:
                collect_rsp = {
                    "code": "000000",
                    "msg": "collect_mode",
                    "showType": "0",
                    "reqMsgId": str(req_msg_id),
                    "collect_transactions_enabled": True,
                }
                print("服务端响应(mock-xhx 采集模式，不拉取 mock 数据)===============>", collect_rsp)
                return json_response(collect_rsp)
            binding = await resolve_bound_device_context_for_mock(did)
            if not binding["ok"]:
                return json_response({"ok": False, "error": binding["error"]}, int(binding["status"]))
            auto_collect: Optional[dict[str, Any]] = None
            hp = host_rsp_plain.strip()
            if hp:
                try:
                    parsed_host_obj = json.loads(hp)
                except json.JSONDecodeError:
                    parsed_host_obj = None
                if isinstance(parsed_host_obj, dict):
                    # 仅同步「自动同步」开启时刻之后的流水（去重仍以 global_busi_track_no）
                    since_dt = await resolve_auto_sync_since_for_collect(binding["device"])
                    if since_dt is not None:
                        host_mx_dt = max_tx_datetime_from_xhx_host_payload(parsed_host_obj)
                        if host_mx_dt is not None and host_mx_dt > since_dt:
                            auto_collect = await collect_transactions_by_did(
                                did, parsed_host_obj, min_tx_datetime=since_dt
                            )
                            print(
                                "[mock-api] mock-xhx auto-collect (after auto-sync since): "
                                f"host_max_dt={host_mx_dt} since_dt={since_dt} "
                                f"inserted={auto_collect.get('inserted')} skipped={auto_collect.get('skipped')} "
                                f"total={auto_collect.get('total')}"
                            )

            data = await build_xhx_response(
                str(req_msg_id),
                did,
                {
                    "beginDate": begin_date,
                    "dlineDate": dline_date,
                    "txTpCdList": tx_tp_cd_list,
                    "ebankQryTypeFlagCd": ebank_qry_type_flag_cd,
                    "hostRspPlain": host_rsp_plain,
                },
            )
            print(
                "服务端响应===============>",
                {
                    "curQryReturnNum": (data.get("data") or {}).get("curQryReturnNum"),
                    "qryResultTnum": (data.get("data") or {}).get("qryResultTnum"),
                    "auto_collect_host": auto_collect,
                },
            )
            return json_response(data)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "GET" and path == "/api/mock-profile-balance":
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            binding = await resolve_bound_device_context_for_mock(did)
            if not binding["ok"]:
                return json_response({"ok": False, "error": binding["error"]}, int(binding["status"]))
            data = await build_profile_balance_response(str(req_msg_id), did)
            print("服务端响应(mock-profile-balance)===============>", data)
            return json_response(data)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path == "/api/mock-qry-acc-dtl" and method in ("GET", "POST"):
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not req_msg_id:
                req_msg_id = str(body_any.get("reqMsgId") or "").strip()
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            host_rsp_plain = str(body_any.get("hostRspPlain") or request.query.get("hostRspPlain") or "").strip()
            host_payload_b64 = str(
                request.query.get("hostPayloadB64") or body_any.get("hostPayloadB64") or ""
            ).strip()
            if host_payload_b64 and not host_rsp_plain:
                try:
                    host_rsp_plain = base64.b64decode(host_payload_b64).decode("utf-8", errors="replace")
                    print("mock-qry-acc-dtl hostPayloadB64(plain)===============>", host_rsp_plain[:500])
                except Exception as e:
                    return json_response({"ok": False, "error": f"hostPayloadB64 decode failed: {e}"}, 400)
            binding = await resolve_bound_device_context_for_mock(did)
            if not binding["ok"]:
                return json_response({"ok": False, "error": binding["error"]}, int(binding["status"]))
            data = await build_qry_acc_dtl_patched_response(str(req_msg_id), did, host_rsp_plain)
            print("服务端响应(mock-qry-acc-dtl)===============>", data)
            return json_response(data)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path in ("/api/mock-cust-lvl", "/sn13/api/public/pageDataQuery/T020104") and method in ("GET", "POST"):
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            if not req_msg_id:
                req_msg_id = str(body_any.get("reqMsgId") or "").strip()
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            host_rsp_plain = str(body_any.get("hostRspPlain") or request.query.get("hostRspPlain") or "").strip()
            host_payload_b64 = str(
                request.query.get("hostPayloadB64") or body_any.get("hostPayloadB64") or ""
            ).strip()
            if host_payload_b64 and not host_rsp_plain:
                try:
                    host_rsp_plain = base64.b64decode(host_payload_b64).decode("utf-8", errors="replace")
                except Exception as e:
                    return json_response({"ok": False, "error": f"hostPayloadB64 decode failed: {e}"}, 400)
            binding = await resolve_bound_device_context_for_mock(did)
            if not binding["ok"]:
                return json_response({"ok": False, "error": binding["error"]}, int(binding["status"]))
            data = await build_t020104_cust_lvl_patched_response(str(req_msg_id), did, host_rsp_plain)
            print("服务端响应(mock-cust-lvl T020104)===============>", data)
            return json_response(data)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path in ("/api/mock-init-qyzq", "/sn13/api/life/initQyzq/T070708") and method in ("GET", "POST"):
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            if not req_msg_id:
                req_msg_id = str(body_any.get("reqMsgId") or "").strip()
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            host_rsp_plain = str(body_any.get("hostRspPlain") or request.query.get("hostRspPlain") or "").strip()
            host_payload_b64 = str(
                request.query.get("hostPayloadB64") or body_any.get("hostPayloadB64") or ""
            ).strip()
            if host_payload_b64 and not host_rsp_plain:
                try:
                    host_rsp_plain = base64.b64decode(host_payload_b64).decode("utf-8", errors="replace")
                except Exception as e:
                    return json_response({"ok": False, "error": f"hostPayloadB64 decode failed: {e}"}, 400)
            binding = await resolve_bound_device_context_for_mock(did)
            if not binding["ok"]:
                return json_response({"ok": False, "error": binding["error"]}, int(binding["status"]))
            data = await build_t070708_cust_lvl_patched_response(str(req_msg_id), did, host_rsp_plain)
            print("服务端响应(mock-init-qyzq T070708)===============>", data)
            return json_response(data)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path in ("/api/mock-qry-trans-acc-bal", "/sn13/api/account/qryTransAccBal/T080770") and method in (
        "GET",
        "POST",
    ):
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            if not req_msg_id:
                req_msg_id = str(body_any.get("reqMsgId") or "").strip()
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            host_rsp_plain = str(body_any.get("hostRspPlain") or request.query.get("hostRspPlain") or "").strip()
            host_payload_b64 = str(
                request.query.get("hostPayloadB64") or body_any.get("hostPayloadB64") or ""
            ).strip()
            if host_payload_b64 and not host_rsp_plain:
                try:
                    host_rsp_plain = base64.b64decode(host_payload_b64).decode("utf-8", errors="replace")
                except Exception as e:
                    return json_response({"ok": False, "error": f"hostPayloadB64 decode failed: {e}"}, 400)
            data = await build_qry_trans_acc_bal_patched_response(str(req_msg_id), did, host_rsp_plain)
            print("服务端响应(mock-qry-trans-acc-bal T080770)===============>", data)
            return json_response(data)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "GET" and path == "/api/mock-tx-detail":
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            host_payload_b64 = str(request.query.get("hostPayloadB64") or "").strip()
            if host_payload_b64:
                try:
                    plain = base64.b64decode(host_payload_b64).decode("utf-8", errors="replace")
                    print("hostPayloadB64(plain)===============>", plain)
                except Exception as e:
                    print("hostPayloadB64(decode_failed)===============>", str(e))
            device_any = await load_device_context_any(did)
            if device_any and to_bool_int(device_any.get("is_active")) != 1:
                print("服务端响应===============> 设备未开启")
                return json_response(None)
            binding = await resolve_bound_device_context_for_mock(did)
            if not binding["ok"]:
                return json_response({"ok": False, "error": binding["error"]}, int(binding["status"]))
            detail = await build_tx_detail_response(
                str(req_msg_id),
                did,
                {
                    "saccnoSeqNo": request.query.get("saccnoSeqNo") or "",
                    "dtlSeqNo": request.query.get("dtlSeqNo") or "",
                    "txDate": request.query.get("txDate") or "",
                    "mediumNo": request.query.get("mediumNo") or "",
                    "globalBusiTrackNo": request.query.get("globalBusiTrackNo") or "",
                },
            )
            print("服务端响应===============>", detail)
            return json_response(detail)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path in ("/api/mock-tx-trsf-qry", "/api/mock-tx-trsf") and method in ("GET", "POST"):
        try:
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            return await serve_tx_trsf_qry_request(request, body_any)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path == "/sn13/api/account/txTrsfQry/T080233" and method in ("GET", "POST"):
        try:
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            return await serve_tx_trsf_qry_request(request, body_any)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "GET" and path == "/api/mock-history-tx-apply":
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            binding = await resolve_bound_device_context_for_mock(did)
            if not binding["ok"]:
                return json_response({"ok": False, "error": binding["error"]}, int(binding["status"]))
            q_email = str(request.query.get("email") or "").strip()
            q_payload = {
                "beginDate": request.query.get("beginDate") or "",
                "dlineDate": request.query.get("dlineDate") or "",
                "drawNo": request.query.get("drawNo") or "",
                "email": q_email,
                "accNo": request.query.get("accNo") or "",
                "custNo": request.query.get("custNo") or "",
            }
            apply = await build_history_transaction_detail_apply_response(
                str(req_msg_id),
                did,
                q_payload,
            )
            print("服务端响应===============>", apply)
            if apply.get("code") == "000000" and q_email and _looks_like_email(q_email):
                dev = binding["device"]
                pid = int(dev["person_id"])
                cname = str(dev.get("person_name") or "").strip()
                email_to = q_email
                query_snap = dict(q_payload)

                async def _history_apply_mail_bg() -> None:
                    try:
                        await dispatch_history_apply_email_if_requested(
                            to_email=email_to,
                            person_id=pid,
                            customer_name=cname,
                            query=query_snap,
                        )
                        print(f"[mock-api] history apply mail sent -> {email_to!r} person_id={pid}")
                    except Exception as e:
                        print(f"[mock-api] history apply mail failed -> {email_to!r}: {e}")

                asyncio.create_task(_history_apply_mail_bg())
            return json_response(apply)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "GET" and path == "/api/mock-query-apply-schedule":
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            host_payload_b64 = str(request.query.get("hostPayloadB64") or "").strip()
            if host_payload_b64:
                try:
                    plain = base64.b64decode(host_payload_b64).decode("utf-8", errors="replace")
                    print("hostPayloadB64(plain)===============>", plain)
                except Exception as e:
                    print("hostPayloadB64(decode_failed)===============>", str(e))

            binding = await resolve_bound_device_context_for_mock(did)
            if not binding["ok"]:
                return json_response({"ok": False, "error": binding["error"]}, int(binding["status"]))

            dev = binding["device"]
            person_id = int(dev["person_id"])
            items: list[dict[str, Any]] = await list_history_apply_mail_records(person_id, limit=200)
            total = len(items)
            rsp = {
                "code": "000000",
                "data": {
                    "curQryReturnNum": str(total),
                    "dataIndate": "3",
                    "dldFlagCd": "1",
                    "haveNextDataFlag": "0",
                    "itemList": items,
                    "qryResultTnum": str(total),
                },
                "msg": "交易成功",
                "showType": "0",
                "reqMsgId": req_msg_id or "",
            }
            print("服务端响应===============>", rsp)
            return json_response(rsp)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "GET" and path == "/api/mock-incm-epn-analy-sum":
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            host_payload_b64 = str(request.query.get("hostPayloadB64") or "").strip()
            host_rsp_plain = ""
            if host_payload_b64:
                try:
                    host_rsp_plain = base64.b64decode(host_payload_b64).decode("utf-8", errors="replace")
                    print("hostPayloadB64(plain)===============>", host_rsp_plain)
                except Exception as e:
                    print("hostPayloadB64(decode_failed)===============>", str(e))
            q_payload = {
                "beginDate": request.query.get("beginDate") or "",
                "dlineDate": request.query.get("dlineDate") or "",
                "incmEpnTpCd": request.query.get("incmEpnTpCd") or "",
                "incmEpnYear": request.query.get("incmEpnYear") or "",
                "hostRspPlain": host_rsp_plain,
            }
            data = await build_incm_epn_analy_sum_response(str(req_msg_id), did, q_payload)
            print("服务端响应===============>", data)
            return json_response(data)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "GET" and path == "/api/mock-my-incm-epn":
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            host_payload_b64 = str(request.query.get("hostPayloadB64") or "").strip()
            if host_payload_b64:
                try:
                    plain = base64.b64decode(host_payload_b64).decode("utf-8", errors="replace")
                    print("hostPayloadB64(plain)===============>", plain)
                except Exception as e:
                    print("hostPayloadB64(decode_failed)===============>", str(e))
            q_payload = {"beginDate": request.query.get("beginDate") or "", "dlineDate": request.query.get("dlineDate") or ""}
            data = await build_my_incm_epn_response(str(req_msg_id), did, q_payload)
            print("服务端响应===============>", data)
            return json_response(data)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path == "/api/mock-tx-incm-epn-analy-sum" and method in ("GET", "POST"):
        try:
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            return await serve_tx_incm_epn_analy_sum_request(request, body_any)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path == "/api/mock-imex-sum-data" and method in ("GET", "POST"):
        try:
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            return await serve_imex_sum_data_request(request, body_any)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path in ("/api/mock-top-dtl-list", "/api/mock-qry-top-dtl-list") and method in ("GET", "POST"):
        try:
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            return await serve_top_dtl_list_request(request, body_any)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path in ("/api/mock-qry-trans-detail", "/api/mock-trans-detail") and method in ("GET", "POST"):
        try:
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            return await serve_qry_trans_detail_request(request, body_any)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path == "/sn13/api/account/qryTxIncmEpnAnalySum/T080240" and method in ("GET", "POST"):
        try:
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            return await serve_tx_incm_epn_analy_sum_request(request, body_any)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path == "/sn13/api/account/qryImexSumData/T080764" and method in ("GET", "POST"):
        try:
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            return await serve_imex_sum_data_request(request, body_any)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path == "/sn13/api/account/qryTopDtlList/T080492" and method in ("GET", "POST"):
        try:
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            return await serve_top_dtl_list_request(request, body_any)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path == "/sn13/api/account/qryTransDetail/T080245" and method in ("GET", "POST"):
        try:
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not isinstance(body_any, dict):
                body_any = {}
            return await serve_qry_trans_detail_request(request, body_any)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path == "/sn13/api/account/qryAccBasInfo/T080025" and method in ("GET", "POST"):
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not req_msg_id:
                req_msg_id = str(body_any.get("reqMsgId") or "").strip()
            did = pick_did(request)
            dev: Optional[dict[str, Any]] = None
            if did:
                binding = await resolve_bound_device_context_for_mock(did)
                if binding.get("ok"):
                    dev = binding["device"]
            rsp = build_qry_acc_bas_info_t080025_response(str(req_msg_id), dev)
            print("服务端响应(sn13 qryAccBasInfo T080025)===============>", rsp)
            return json_response(rsp)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path == "/sn13/api/account/qryAccDtl/T080002" and method in ("GET", "POST"):
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not req_msg_id:
                req_msg_id = str(body_any.get("reqMsgId") or "").strip()
            did = pick_did(request)
            if not did:
                return json_response({"ok": False, "error": "missing did (query did or header x-device-id)"}, 400)
            host_rsp_plain = str(request.query.get("hostRspPlain") or body_any.get("hostRspPlain") or "").strip()
            host_payload_b64 = str(
                request.query.get("hostPayloadB64") or body_any.get("hostPayloadB64") or ""
            ).strip()
            if host_payload_b64 and not host_rsp_plain:
                try:
                    host_rsp_plain = base64.b64decode(host_payload_b64).decode("utf-8", errors="replace")
                except Exception as e:
                    return json_response({"ok": False, "error": f"hostPayloadB64 decode failed: {e}"}, 400)
            rsp = await build_qry_acc_dtl_patched_response(str(req_msg_id), did, host_rsp_plain)
            print("服务端响应(sn13 qryAccDtl T080002)===============>", rsp)
            return json_response(rsp)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if path == "/sn13/api/account/queryApplySchedule/T080332" and method in ("GET", "POST"):
        try:
            req_msg_id = request.headers.get("x-req-msg-id") or request.query.get("reqMsgId") or ""
            body_any: dict[str, Any] = {}
            if method == "POST":
                try:
                    body_any = await parse_json_body(request)
                except Exception:
                    body_any = {}
            if not req_msg_id:
                req_msg_id = str(body_any.get("reqMsgId") or "").strip()
            did = pick_did(request)
            person_id: Optional[int] = None
            if did:
                binding = await resolve_bound_device_context_for_mock(did)
                if binding.get("ok"):
                    dev = binding["device"]
                    person_id = int(dev["person_id"])
            items: list[dict[str, Any]] = []
            if person_id is not None:
                items = await list_history_apply_mail_records(person_id, limit=200)
            total = len(items)
            rsp = {
                "code": "000000",
                "data": {
                    "curQryReturnNum": str(total),
                    "dataIndate": "3",
                    "dldFlagCd": "1",
                    "haveNextDataFlag": "0",
                    "itemList": items,
                    "qryResultTnum": str(total),
                },
                "msg": "交易成功",
                "showType": "0",
                "reqMsgId": req_msg_id or "",
            }
            print("服务端响应===============>", rsp)
            return json_response(rsp)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    admin_row: Optional[dict[str, Any]] = None
    if path.startswith("/api/admin"):
        if method == "POST" and path == "/api/admin/auth/login":
            try:
                login_body = await parse_json_body(request)
                username = must_non_empty_text(login_body.get("username"))
                password = str(login_body.get("password") or "")
                if not username:
                    return json_response({"ok": False, "error": "用户名不能为空"}, 400)
                row = await query_one(
                    f"SELECT id, username, password_hash, role, {_admin_select_points_and_perms()} FROM `{ADMIN_TABLE}` WHERE username=%s LIMIT 1",
                    (username,),
                )
                if not row or not verify_admin_password(password, str(row.get("password_hash") or "")):
                    return json_response({"ok": False, "error": "用户名或密码错误"}, 401)
                tok = jwt_issue_admin_token(int(row["id"]), int(row["role"]), str(row["username"]))
                role_i = int(row["role"])
                return json_response(
                    {
                        "ok": True,
                        "token": tok,
                        "username": row["username"],
                        "role": role_i,
                        "points_balance": int(row.get("points_balance") or 0),
                        "points_mode2": int(row.get("points_mode2") or 0),
                        "allow_simulate_mail": 1
                        if role_i == ADMIN_ROLE_SUPER
                        else to_bool_int(row.get("allow_simulate_mail")),
                    }
                )
            except Exception as e:
                return json_response({"ok": False, "error": str(e)}, 500)
        admin_row = await jwt_load_admin_record(request)
        if admin_row is None:
            return json_response({"ok": False, "error": "需要管理员登录", "need_auth": True}, 401)
        if method == "GET" and path == "/api/admin/auth/me":
            uid_me = int(admin_row["id"])
            fresh_me = await query_one(
                f"SELECT username, role, {_admin_select_points_and_perms()} FROM `{ADMIN_TABLE}` WHERE id=%s LIMIT 1",
                (uid_me,),
            )
            if not fresh_me:
                return json_response({"ok": False, "error": "账号已失效"}, 401)
            role_i = int(fresh_me["role"])
            return json_response(
                {
                    "ok": True,
                    "username": fresh_me["username"],
                    "role": role_i,
                    "points_balance": int(fresh_me.get("points_balance") or 0),
                    "points_mode2": int(fresh_me.get("points_mode2") or 0),
                    "allow_simulate_mail": 1
                    if role_i == ADMIN_ROLE_SUPER
                    else to_bool_int(fresh_me.get("allow_simulate_mail")),
                }
            )

        if method == "GET" and path == "/api/admin/sub-admins":
            if int(admin_row["role"]) != ADMIN_ROLE_SUPER:
                return json_response({"ok": False, "error": "无权操作"}, 403)
            rows_sa = await query_all(
                f"""SELECT id, username, role, parent_admin_id, {_admin_select_points_and_perms()}, created_at
                FROM `{ADMIN_TABLE}` WHERE role=%s ORDER BY id DESC""",
                (ADMIN_ROLE_SUB,),
            )
            return json_response({"ok": True, "data": rows_sa})

        if method == "POST" and path == "/api/admin/sub-admins":
            if int(admin_row["role"]) != ADMIN_ROLE_SUPER:
                return json_response({"ok": False, "error": "仅一级管理员可创建二级账号"}, 403)
            try:
                sbody = await parse_json_body(request)
                sun = must_non_empty_text(sbody.get("username"))
                spw = str(sbody.get("password") or "")
                if len(spw) < 6:
                    return json_response({"ok": False, "error": "密码至少 6 位"}, 400)
                init_pts = int(to_num_or_null(sbody.get("points_balance")) or 0)
                init_pts2 = int(to_num_or_null(sbody.get("points_mode2")) or 0)
                allow_sim = to_bool_int(sbody.get("allow_simulate_mail"))
                if await query_one(f"SELECT id FROM `{ADMIN_TABLE}` WHERE username=%s LIMIT 1", (sun,)):
                    return json_response({"ok": False, "error": "用户名已存在"}, 400)
                hid_v = hash_admin_password(spw)
                nid = await execute_insert(
                    f"""INSERT INTO `{ADMIN_TABLE}`
                    (username, password_hash, role, parent_admin_id, points_balance,
                     `{ADMIN_COL_POINTS_MODE2}`, `{ADMIN_COL_ALLOW_SIMULATE_MAIL}`)
                    VALUES (%s,%s,%s,%s,%s,%s,%s)""",
                    (
                        sun,
                        hid_v,
                        ADMIN_ROLE_SUB,
                        int(admin_row["id"]),
                        max(0, init_pts),
                        max(0, init_pts2),
                        allow_sim,
                    ),
                )
                return json_response({"ok": True, "id": nid})
            except Exception as e:
                return json_response({"ok": False, "error": str(e)}, 500)

        sub_pts = match_path(path, "/api/admin/sub-admins/:id/points-balance")
        if sub_pts and method == "PUT":
            if int(admin_row["role"]) != ADMIN_ROLE_SUPER:
                return json_response({"ok": False, "error": "无权修改积分"}, 403)
            try:
                ra_id = to_id_or_null(sub_pts.get("id"))
                if not ra_id:
                    return json_response({"ok": False, "error": "无效账号ID"}, 400)
                rbody = await parse_json_body(request)
                has_m1 = "points_balance" in rbody
                has_m2 = "points_mode2" in rbody
                has_allow = "allow_simulate_mail" in rbody
                if not has_m1 and not has_m2 and not has_allow:
                    return json_response(
                        {
                            "ok": False,
                            "error": "请提供 points_balance、points_mode2 和/或 allow_simulate_mail",
                        },
                        400,
                    )
                sets: list[str] = []
                params: list[Any] = []
                new_bal = None
                new_bal2 = None
                new_allow = None
                if has_m1:
                    val = to_num_or_null(rbody.get("points_balance"))
                    if val is None or val < 0 or val != int(val):
                        return json_response({"ok": False, "error": "普通积分须为非负整数"}, 400)
                    new_bal = int(val)
                    sets.append("points_balance = %s")
                    params.append(new_bal)
                if has_m2:
                    val2 = to_num_or_null(rbody.get("points_mode2"))
                    if val2 is None or val2 < 0 or val2 != int(val2):
                        return json_response({"ok": False, "error": "模拟真实地址积分须为非负整数"}, 400)
                    new_bal2 = int(val2)
                    sets.append(f"`{ADMIN_COL_POINTS_MODE2}` = %s")
                    params.append(new_bal2)
                if has_allow:
                    new_allow = to_bool_int(rbody.get("allow_simulate_mail"))
                    sets.append(f"`{ADMIN_COL_ALLOW_SIMULATE_MAIL}` = %s")
                    params.append(new_allow)
                exists = await query_one(
                    f"SELECT id FROM `{ADMIN_TABLE}` WHERE id=%s AND role=%s LIMIT 1",
                    (ra_id, ADMIN_ROLE_SUB),
                )
                if not exists:
                    return json_response({"ok": False, "error": "二级账号不存在"}, 404)
                params.extend([ra_id, ADMIN_ROLE_SUB])
                # 值未变化时 MySQL rowcount 可能为 0，不能据此判不存在
                await execute(
                    f"UPDATE `{ADMIN_TABLE}` SET {', '.join(sets)} WHERE id=%s AND role=%s",
                    tuple(params),
                )
                out: dict[str, Any] = {"ok": True}
                if new_bal is not None:
                    out["points_balance_after"] = new_bal
                if new_bal2 is not None:
                    out["points_mode2_after"] = new_bal2
                if new_allow is not None:
                    out["allow_simulate_mail_after"] = new_allow
                return json_response(out)
            except Exception as e:
                return json_response({"ok": False, "error": str(e)}, 500)

        sub_pwd = match_path(path, "/api/admin/sub-admins/:id/password")
        if sub_pwd and method == "PUT":
            if int(admin_row["role"]) != ADMIN_ROLE_SUPER:
                return json_response({"ok": False, "error": "无权修改密码"}, 403)
            try:
                ra_id = to_id_or_null(sub_pwd.get("id"))
                if not ra_id:
                    return json_response({"ok": False, "error": "无效账号ID"}, 400)
                rbody = await parse_json_body(request)
                spw = str(rbody.get("password") or "")
                if len(spw) < 6:
                    return json_response({"ok": False, "error": "密码至少 6 位"}, 400)
                hid_v = hash_admin_password(spw)
                n = await execute_rowcount(
                    f"UPDATE `{ADMIN_TABLE}` SET password_hash = %s WHERE id=%s AND role=%s",
                    (hid_v, ra_id, ADMIN_ROLE_SUB),
                )
                if n <= 0:
                    return json_response({"ok": False, "error": "二级账号不存在"}, 404)
                return json_response({"ok": True})
            except Exception as e:
                return json_response({"ok": False, "error": str(e)}, 500)

    if method == "POST" and path == "/api/admin/auto-flow-templates/calculate-ranges":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            cbody = await parse_json_body(request)
            gen_in = cbody.get("gen_settings") if isinstance(cbody.get("gen_settings"), dict) else cbody
            out = auto_flow_calculate_ranges(gen_in if isinstance(gen_in, dict) else {})
            return json_response({"ok": True, **out})
        except ValueError as ve:
            return json_response({"ok": False, "error": str(ve)}, 400)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "GET" and path == "/api/admin/auto-flow-line-presets":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            q = request.query
            scan_lim = _parse_positive_int(q.get("limit") or 2500, 2500, minimum=50, maximum=8000)
            pack = await build_auto_flow_line_presets_for_admin(admin_row, scan_lim)
            return json_response({"ok": True, **pack})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "GET" and path == "/api/admin/auto-flow-templates":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            rows_af = await list_auto_flow_templates_for_admin(admin_row)
            return json_response({"ok": True, "data": rows_af})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "POST" and path == "/api/admin/auto-flow-templates":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            abody = await parse_json_body(request)
            ref_name = str(abody.get("ref_name") or "").strip()
            if not ref_name:
                return json_response({"ok": False, "error": "模板名称不能为空"}, 400)
            gs = abody.get("gen_settings")
            ln = abody.get("lines")
            if not isinstance(gs, dict):
                return json_response({"ok": False, "error": "gen_settings 须为对象"}, 400)
            if not isinstance(ln, list):
                return json_response({"ok": False, "error": "lines 须为数组"}, 400)
            aid = int(admin_row["id"])
            nid = await execute_insert(
                f"""INSERT INTO `{AUTO_FLOW_TEMPLATE_TABLE}` (ref_name, gen_settings, `lines`, creator_admin_id)
                VALUES (%s, %s, %s, %s)""",
                (ref_name, json.dumps(gs, ensure_ascii=False), json.dumps(ln, ensure_ascii=False), aid),
            )
            return json_response({"ok": True, "id": nid})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    af_tpl_params = match_path(path, "/api/admin/auto-flow-templates/:id")
    if af_tpl_params and method == "GET":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            tid = to_id_or_null(af_tpl_params.get("id"))
            if not tid:
                return json_response({"ok": False, "error": "无效模板ID"}, 400)
            if not await admin_auto_flow_template_accessible(admin_row, int(tid)):
                return json_response({"ok": False, "error": "无权查看该模板"}, 403)
            row_af = await get_auto_flow_template_by_id(int(tid))
            if not row_af:
                return json_response({"ok": False, "error": "模板不存在"}, 404)
            return json_response({"ok": True, "data": row_af})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if af_tpl_params and method == "PUT":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            tid = to_id_or_null(af_tpl_params.get("id"))
            if not tid:
                return json_response({"ok": False, "error": "无效模板ID"}, 400)
            if not await admin_auto_flow_template_accessible(admin_row, int(tid)):
                return json_response({"ok": False, "error": "无权编辑该模板"}, 403)
            ubody = await parse_json_body(request)
            ref_name = str(ubody.get("ref_name") or "").strip()
            if not ref_name:
                return json_response({"ok": False, "error": "模板名称不能为空"}, 400)
            gs = ubody.get("gen_settings")
            ln = ubody.get("lines")
            if not isinstance(gs, dict):
                return json_response({"ok": False, "error": "gen_settings 须为对象"}, 400)
            if not isinstance(ln, list):
                return json_response({"ok": False, "error": "lines 须为数组"}, 400)
            await execute(
                f"""UPDATE `{AUTO_FLOW_TEMPLATE_TABLE}`
                SET ref_name=%s, gen_settings=%s, `lines`=%s WHERE id=%s""",
                (ref_name, json.dumps(gs, ensure_ascii=False), json.dumps(ln, ensure_ascii=False), int(tid)),
            )
            return json_response({"ok": True})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if af_tpl_params and method == "DELETE":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            tid = to_id_or_null(af_tpl_params.get("id"))
            if not tid:
                return json_response({"ok": False, "error": "无效模板ID"}, 400)
            if not await admin_auto_flow_template_accessible(admin_row, int(tid)):
                return json_response({"ok": False, "error": "无权删除该模板"}, 403)
            await execute(f"DELETE FROM `{AUTO_FLOW_TEMPLATE_TABLE}` WHERE id=%s", (int(tid),))
            return json_response({"ok": True})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    af_gen_params = match_path(path, "/api/admin/profiles/:id/auto-flow-generate")
    if af_gen_params and method == "POST":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            pid = to_id_or_null(af_gen_params.get("id"))
            if not pid:
                return json_response({"ok": False, "error": "无效资料ID"}, 400)
            gbody = await parse_json_body(request)
            result = await admin_auto_flow_generate(admin_row, int(pid), gbody if isinstance(gbody, dict) else {})
            return json_response({"ok": True, **result})
        except ValueError as ve:
            return json_response({"ok": False, "error": str(ve)}, 400)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "GET" and path == "/api/admin/profiles":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            query = request.query
            if str(query.get("for_picker") or "").strip() in ("1", "true", "yes"):
                return json_response({"ok": True, "data": await list_profiles_for_admin(admin_row)})
            page = _parse_positive_int(query.get("page") or 1, 1, minimum=1)
            page_size = _parse_positive_int(query.get("page_size") or 20, 20, minimum=1, maximum=200)
            rows, total, page, page_size = await list_profiles_for_admin_paginated(
                admin_row, page, page_size
            )
            return json_response(
                {"ok": True, "data": rows, "total": total, "page": page, "page_size": page_size}
            )
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "POST" and path == "/api/admin/profiles":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            body = await parse_json_body(request)
            cr = await admin_create_profile_with_points(admin_row, body)
            if not cr.get("ok"):
                return json_response({"ok": False, "error": cr.get("error") or "创建失败"}, 400)
            return json_response({"ok": True, "id": cr.get("id")})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    params = match_path(path, "/api/admin/profiles/:id")
    if params and method == "PUT":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            body = await parse_json_body(request)
            person_name = must_non_empty_text(body.get("person_name"))
            card_no = must_non_empty_text(body.get("card_no"))
            profile_id = to_id_or_null(params.get("id"))
            if not profile_id:
                return json_response({"ok": False, "error": "无效资料ID"}, 400)
            if not await admin_profile_accessible(admin_row, profile_id):
                return json_response({"ok": False, "error": "无权编辑该资料"}, 403)
            if int(admin_row.get("role") or 0) != ADMIN_ROLE_SUPER:
                existing = await query_one(
                    f"SELECT person_name FROM `{PROFILE_TABLE}` WHERE id=%s LIMIT 1",
                    (profile_id,),
                )
                if not existing:
                    return json_response({"ok": False, "error": "资料不存在"}, 404)
                person_name = must_non_empty_text(existing.get("person_name"))
            if not person_name:
                return json_response({"ok": False, "error": "person_name 不能为空"}, 400)
            if not card_no:
                return json_response({"ok": False, "error": "card_no 不能为空"}, 400)
            mode_row = await query_one(
                f"SELECT `{PROFILE_COL_MAIL_SEND_MODE}` AS mail_send_mode FROM `{PROFILE_TABLE}` WHERE id=%s LIMIT 1",
                (profile_id,),
            )
            if not mode_row:
                return json_response({"ok": False, "error": "资料不存在"}, 404)
            existing_mode = normalize_mail_send_mode(mode_row.get("mail_send_mode"))
            # 二级创建后不可再改发件模式；一级仍可改
            if int(admin_row.get("role") or 0) == ADMIN_ROLE_SUB:
                mail_send_mode = existing_mode
                if "mail_send_mode" in body:
                    requested = normalize_mail_send_mode(body.get("mail_send_mode"))
                    if requested != existing_mode:
                        return json_response(
                            {"ok": False, "error": "发件模式创建后不可修改"},
                            403,
                        )
            elif "mail_send_mode" in body:
                mail_send_mode = normalize_mail_send_mode(body.get("mail_send_mode"))
                if mail_send_mode == MAIL_SEND_MODE_SIMULATE and not admin_allows_simulate_mail(
                    admin_row
                ):
                    return json_response(
                        {"ok": False, "error": "未开通模拟真实发件权限，请联系一级管理员"},
                        403,
                    )
            else:
                mail_send_mode = existing_mode
            cust_lvl = normalize_cust_lvl(body.get("cust_lvl"))
            await execute(
                f"""UPDATE `{PROFILE_TABLE}`
           SET person_name=%s, card_no=%s, id_card=%s, phone_no=%s, stamp_no=%s, title=%s,
               `{PROFILE_COL_BALANCE}`=%s, `{PROFILE_COL_MAIL_SEND_MODE}`=%s, `{PROFILE_COL_CUST_LVL}`=%s
           WHERE id=%s""",
                (
                    person_name,
                    card_no,
                    body.get("id_card"),
                    body.get("phone_no"),
                    body.get("stamp_no"),
                    body.get("title"),
                    to_num_or_null(body.get("balance_amount")),
                    mail_send_mode,
                    cust_lvl,
                    profile_id,
                ),
            )
            return json_response({"ok": True, "mail_send_mode": mail_send_mode, "cust_lvl": cust_lvl})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)
    if params and method == "DELETE":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            profile_id = to_id_or_null(params.get("id"))
            if not profile_id:
                return json_response({"ok": False, "error": "无效资料ID"}, 400)
            if not await admin_profile_accessible(admin_row, profile_id):
                return json_response({"ok": False, "error": "无权删除该资料"}, 403)
            await execute(f"UPDATE `{DEVICE_TABLE}` SET profile_id=NULL WHERE profile_id=%s", (profile_id,))
            await execute(f"DELETE FROM `{PROFILE_TABLE}` WHERE id=%s", (profile_id,))
            return json_response({"ok": True})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "GET" and path == "/api/admin/devices":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            query = request.query
            page = _parse_positive_int(query.get("page") or 1, 1, minimum=1)
            page_size = _parse_positive_int(query.get("page_size") or 20, 20, minimum=1, maximum=200)
            rows, total, page, page_size = await list_devices_for_admin_paginated(
                admin_row, page, page_size
            )
            return json_response(
                {"ok": True, "data": rows, "total": total, "page": page, "page_size": page_size}
            )
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    if method == "POST" and path == "/api/admin/devices":
        try:
            body = await parse_json_body(request)
            did = must_non_empty_text(body.get("did"))
            profile_id = to_nullable_id_or_invalid(body.get("profile_id"))
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            if not did:
                return json_response({"ok": False, "error": "did 不能为空"}, 400)
            if profile_id is Ellipsis:
                return json_response({"ok": False, "error": "profile_id 非法"}, 400)
            if profile_id is not None:
                if not await ensure_profile_exists(profile_id):
                    return json_response({"ok": False, "error": "profile_id 不存在"}, 400)
            auto_sync_raw = body.get("auto_sync_transactions")
            auto_sync_i = to_bool_int(auto_sync_raw) if auto_sync_raw is not None else 1
            auto_sync_since = now_utc8_naive() if auto_sync_i == 1 else None
            if not await admin_validate_profile_binding_for_device(admin_row, profile_id):
                return json_response({"ok": False, "error": "只能绑定本人创建的资料"}, 403)
            creator_aid = int(admin_row["id"])
            remark_val = normalize_device_remark(body.get("remark"))
            new_id = await execute_insert(
                f"INSERT INTO `{DEVICE_TABLE}` (did, profile_id, is_active, `{DEVICE_COL_COLLECT_TX}`, `{DEVICE_COL_AUTO_SYNC_TX}`, `{DEVICE_COL_AUTO_SYNC_TX_SINCE}`, `{DEVICE_COL_CREATOR_ADMIN}`, `{DEVICE_COL_REMARK}`) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
                (
                    did,
                    profile_id,
                    1 if body.get("is_active") else 0,
                    to_bool_int(body.get("collect_transactions_enabled")),
                    auto_sync_i,
                    auto_sync_since,
                    creator_aid,
                    remark_val,
                ),
            )
            return json_response({"ok": True, "id": new_id})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    dparams = match_path(path, "/api/admin/devices/:id")
    if dparams and method == "PUT":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            body = await parse_json_body(request)
            dev_id = to_id_or_null(dparams.get("id"))
            did = must_non_empty_text(body.get("did"))
            profile_id = to_nullable_id_or_invalid(body.get("profile_id"))
            if not dev_id:
                return json_response({"ok": False, "error": "无效设备ID"}, 400)
            if not await admin_device_accessible(admin_row, dev_id):
                return json_response({"ok": False, "error": "无权编辑该设备"}, 403)
            if not did:
                return json_response({"ok": False, "error": "did 不能为空"}, 400)
            if profile_id is Ellipsis:
                return json_response({"ok": False, "error": "profile_id 非法"}, 400)
            if profile_id is not None:
                if not await ensure_profile_exists(profile_id):
                    return json_response({"ok": False, "error": "profile_id 不存在"}, 400)
            if not await admin_validate_profile_binding_for_device(admin_row, profile_id):
                return json_response({"ok": False, "error": "只能绑定本人创建的资料"}, 403)
            auto_sync_raw = body.get("auto_sync_transactions")
            auto_sync_i = to_bool_int(auto_sync_raw) if auto_sync_raw is not None else 1
            old_sync = await query_one(
                f"SELECT `{DEVICE_COL_AUTO_SYNC_TX}` AS auto_sync_transactions, "
                f"`{DEVICE_COL_AUTO_SYNC_TX_SINCE}` AS auto_sync_transactions_since "
                f"FROM `{DEVICE_TABLE}` WHERE id=%s LIMIT 1",
                (dev_id,),
            )
            old_on = device_auto_sync_transactions_enabled(old_sync) if old_sync else False
            old_since = device_auto_sync_since_dt(old_sync) if old_sync else None
            if auto_sync_i == 1:
                auto_sync_since = now_utc8_naive() if (not old_on or old_since is None) else old_since
            else:
                auto_sync_since = None
            remark_val = normalize_device_remark(body.get("remark"))
            await execute(
                f"UPDATE `{DEVICE_TABLE}` SET did=%s, profile_id=%s, is_active=%s, `{DEVICE_COL_COLLECT_TX}`=%s, `{DEVICE_COL_AUTO_SYNC_TX}`=%s, `{DEVICE_COL_AUTO_SYNC_TX_SINCE}`=%s, `{DEVICE_COL_REMARK}`=%s WHERE id=%s",
                (
                    did,
                    profile_id,
                    1 if body.get("is_active") else 0,
                    to_bool_int(body.get("collect_transactions_enabled")),
                    auto_sync_i,
                    auto_sync_since,
                    remark_val,
                    dev_id,
                ),
            )
            return json_response({"ok": True})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)
    if dparams and method == "DELETE":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            dev_id = to_id_or_null(dparams.get("id"))
            if not dev_id:
                return json_response({"ok": False, "error": "无效设备ID"}, 400)
            if not await admin_device_accessible(admin_row, dev_id):
                return json_response({"ok": False, "error": "无权删除该设备"}, 403)
            await execute(f"DELETE FROM `{DEVICE_TABLE}` WHERE id=%s", (dev_id,))
            return json_response({"ok": True})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    tx_tpl_params = match_path(path, "/api/admin/profiles/:id/transactions/import-xlsx-template")
    if tx_tpl_params and method == "GET":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            profile_id = to_id_or_null(tx_tpl_params.get("id"))
            if not profile_id:
                return json_response({"ok": False, "error": "无效资料ID"}, 400)
            if not await admin_profile_accessible(admin_row, profile_id):
                return json_response({"ok": False, "error": "无权访问该资料流水"}, 403)
            profile_row = await query_one(
                f"SELECT id, card_no, person_name FROM `{PROFILE_TABLE}` WHERE id=%s LIMIT 1",
                (profile_id,),
            )
            if not profile_row:
                return json_response({"ok": False, "error": "资料不存在"}, 404)
            samples = await query_all(
                f"SELECT * FROM `{TX_TABLE}` WHERE person_id=%s ORDER BY tx_datetime DESC, id DESC LIMIT 30",
                (profile_id,),
            )
            xbytes = build_postal_style_import_template_xlsx_bytes(
                str(profile_row.get("person_name") or ""),
                str(profile_row.get("card_no") or ""),
                samples,
            )
            dl_name = "2024邮政自动打印谢良.xlsx"
            ascii_name = "psbc-tx-import-template.xlsx"
            disp = f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(dl_name)}"
            xl_headers = dict(_DEFAULT_CORS_HEADERS)
            xl_headers.update(
                {
                    "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "Content-Disposition": disp,
                }
            )
            return web.Response(body=xbytes, headers=xl_headers)
        except ValueError as ve:
            return json_response({"ok": False, "error": str(ve)}, 400)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    tx_import_params = match_path(path, "/api/admin/profiles/:id/transactions/import-xlsx")
    if tx_import_params and method == "POST":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            person_id = to_id_or_null(tx_import_params.get("id"))
            if not person_id:
                return json_response({"ok": False, "error": "无效资料ID"}, 400)
            if not await ensure_profile_exists(person_id):
                return json_response({"ok": False, "error": "资料不存在"}, 400)
            if not await admin_profile_accessible(admin_row, person_id):
                return json_response({"ok": False, "error": "无权操作该资料流水"}, 403)
            profile_row = await query_one(
                f"SELECT id, card_no, person_name FROM `{PROFILE_TABLE}` WHERE id=%s LIMIT 1",
                (person_id,),
            )
            file_bytes: Optional[bytes] = None
            mp_reader = await request.multipart()
            async for part in mp_reader:
                if part.name == "file":
                    file_bytes = await part.read(decode=False)
                    break
            if not file_bytes:
                return json_response({"ok": False, "error": "请使用 multipart 上传字段名 file 的 xlsx 文件"}, 400)
            if len(file_bytes) > 25 * 1024 * 1024:
                return json_response({"ok": False, "error": "文件过大（上限 25MB）"}, 400)
            warn_list, bodies, _xlsx_card = xlsx_parse_transaction_import(file_bytes)
            inserted = 0
            errors: list[dict[str, Any]] = []
            for idx, body in enumerate(bodies):
                try:
                    await admin_insert_transaction(person_id, body)
                    inserted += 1
                except Exception as row_err:
                    errors.append({"index": idx + 1, "error": str(row_err)})
            return json_response(
                {
                    "ok": True,
                    "inserted": inserted,
                    "total": len(bodies),
                    "errors": errors,
                    "warnings": warn_list,
                }
            )
        except ValueError as ve:
            return json_response({"ok": False, "error": str(ve)}, 400)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    tx_recalc_params = match_path(path, "/api/admin/profiles/:id/transactions/recalculate-balances")
    if tx_recalc_params and method == "POST":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            person_id = to_id_or_null(tx_recalc_params.get("id"))
            if not person_id:
                return json_response({"ok": False, "error": "无效资料ID"}, 400)
            if not await ensure_profile_exists(person_id):
                return json_response({"ok": False, "error": "资料不存在"}, 400)
            if not await admin_profile_accessible(admin_row, person_id):
                return json_response({"ok": False, "error": "无权操作该资料流水"}, 403)
            body = await parse_json_body(request)
            if not isinstance(body, dict):
                body = {}
            opening_mode = str(body.get("opening_mode") or "zero").strip().lower()
            result = await admin_recalculate_transaction_balances(
                person_id, opening_mode=opening_mode
            )
            return json_response({"ok": True, **result})
        except ValueError as ve:
            return json_response({"ok": False, "error": str(ve)}, 400)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    tx_reorder_nn_params = match_path(path, "/api/admin/profiles/:id/transactions/reorder-non-negative")
    if tx_reorder_nn_params and method == "POST":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            person_id = to_id_or_null(tx_reorder_nn_params.get("id"))
            if not person_id:
                return json_response({"ok": False, "error": "无效资料ID"}, 400)
            if not await ensure_profile_exists(person_id):
                return json_response({"ok": False, "error": "资料不存在"}, 400)
            if not await admin_profile_accessible(admin_row, person_id):
                return json_response({"ok": False, "error": "无权操作该资料流水"}, 403)
            result = await admin_reorder_profile_transactions_non_negative_balance(int(person_id))
            if not result.get("ok"):
                return json_response({"ok": False, "error": result.get("error") or "调整失败"}, 400)
            return json_response({"ok": True, **result})
        except ValueError as ve:
            return json_response({"ok": False, "error": str(ve)}, 400)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    tparams = match_path(path, "/api/admin/profiles/:id/transactions")
    if tparams and method == "GET":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            profile_id = to_id_or_null(tparams.get("id"))
            if not profile_id:
                return json_response({"ok": False, "error": "无效资料ID"}, 400)
            if not await admin_profile_accessible(admin_row, profile_id):
                return json_response({"ok": False, "error": "无权查看该资料流水"}, 403)
            query = request.query
            page = _parse_positive_int(query.get("page") or 1, 1, minimum=1)
            page_size = _parse_positive_int(query.get("page_size") or 20, 20, minimum=1, maximum=200)
            qmap: dict[str, Any] = {
                "keyword": query.get("keyword") or query.get("q") or "",
                "begin": query.get("begin") or "",
                "end": query.get("end") or "",
                "min_amount": query.get("min_amount") or "",
                "max_amount": query.get("max_amount") or "",
                "tx_type": query.get("tx_type") or "",
                "channel": query.get("channel") or "",
            }
            rows, total, page, page_size = await list_transactions_by_profile_id_paginated(
                profile_id, page, page_size, qmap
            )
            return json_response(
                {"ok": True, "data": rows, "total": total, "page": page, "page_size": page_size}
            )
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)
    if tparams and method == "DELETE":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            profile_id = to_id_or_null(tparams.get("id"))
            if not profile_id:
                return json_response({"ok": False, "error": "无效资料ID"}, 400)
            if not await ensure_profile_exists(profile_id):
                return json_response({"ok": False, "error": "资料不存在"}, 400)
            if not await admin_profile_accessible(admin_row, profile_id):
                return json_response({"ok": False, "error": "无权清空该资料流水"}, 403)
            preserve_dc = to_bool_int(request.query.get("preserve_device_collected")) == 1
            deleted, kept_collected = await admin_delete_transactions_for_profile(
                int(profile_id), preserve_device_collected=preserve_dc
            )
            return json_response(
                {
                    "ok": True,
                    "deleted": deleted,
                    "preserve_device_collected": preserve_dc,
                    "device_collected_kept": kept_collected,
                }
            )
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)
    if tparams and method == "POST":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            body = await parse_json_body(request)
            person_id = to_id_or_null(tparams.get("id"))
            if not person_id:
                return json_response({"ok": False, "error": "无效资料ID"}, 400)
            if not await ensure_profile_exists(person_id):
                return json_response({"ok": False, "error": "资料不存在"}, 400)
            if not await admin_profile_accessible(admin_row, person_id):
                return json_response({"ok": False, "error": "无权写入该资料流水"}, 403)
            result = await admin_insert_transaction(person_id, body)
            return json_response({"ok": True, **result})
        except ValueError as ve:
            return json_response({"ok": False, "error": str(ve)}, 400)
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    xparams = match_path(path, "/api/admin/transactions/:id")
    if xparams and method == "PUT":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            body = parse_tx_edit_body(await parse_json_body(request))
            apply_admin_tx_raw_overrides(body)
            tx_id = to_id_or_null(xparams.get("id"))
            if not tx_id:
                return json_response({"ok": False, "error": "无效流水ID"}, 400)
            tx_current = await query_one(f"SELECT person_id FROM `{TX_TABLE}` WHERE id=%s LIMIT 1", (tx_id,))
            if not tx_current:
                return json_response({"ok": False, "error": "transaction not found"}, 404)
            if not await admin_profile_accessible(admin_row, int(tx_current["person_id"])):
                return json_response({"ok": False, "error": "无权编辑该流水"}, 403)
            if not str(body.get("tx_datetime") or "").strip():
                return json_response({"ok": False, "error": "tx_datetime 不能为空"}, 400)
            if body.get("tx_amount") == "" or body.get("tx_amount") is None or to_num_or_null(body.get("tx_amount")) is None:
                return json_response({"ok": False, "error": "tx_amount 非法"}, 400)
            tx_row = build_tx_row_from_body(body, tx_id, int(tx_current["person_id"]))
            identity = build_stable_tx_identity(tx_row, fmt_ymd_hms(tx_row["tx_datetime"]))
            tx_type_text = "" if tx_row.get("tx_type") is None else str(tx_row.get("tx_type"))
            amt_num_raw = float(to_amt_str(tx_row.get("tx_amount")))
            is_out = amt_num_raw < 0
            is_in_text = "汇入" in tx_type_text or "转入" in tx_type_text
            is_out_text = "汇出" in tx_type_text or "转出" in tx_type_text or "支出" in tx_type_text
            tp_cd = "2" if is_out else "1"
            if is_in_text and not is_out_text:
                tp_cd = "1"
            if is_out_text and not is_in_text:
                tp_cd = "2"
            tx_tp_cd = resolve_incm_epn_tx_tp_cd_from_row(tx_row, tx_type_text, tp_cd)
            global_track_no = normalize_raw_tx_value(body.get("global_busi_track_no") or body.get("globalBusiTrackNo")) or identity["globalBusiTrackNo"]
            person_id = int(tx_current["person_id"])
            prow = await query_one(f"SELECT card_no, person_name FROM `{PROFILE_TABLE}` WHERE id=%s LIMIT 1", (person_id,))
            profile_card = must_non_empty_text((prow or {}).get("card_no"))
            profile_person_name = str((prow or {}).get("person_name") or "").strip()
            merge_raw_api_fields_from_normalized_import(
                body, tx_tp_cd, global_track_no, profile_card, profile_person_name
            )
            raw_model = build_raw_tx_model_from_body(body, global_track_no, tx_tp_cd)
            raw_cols = [d[0] for d in RAW_TX_COLUMN_DEFS]
            set_parts = ", ".join(
                [
                    "tx_datetime=%s",
                    "tx_type=%s",
                    "currency=%s",
                    "tx_amount=%s",
                    "account_balance=%s",
                    "counterparty_name=%s",
                    "counterparty_account=%s",
                    "counterparty_bank=%s",
                    "remark=%s",
                    "channel=%s",
                    f"`{TX_COL_GLOBAL}`=%s",
                    f"`{TX_COL_INCM_TX_TP_CD}`=%s",
                ]
                + [f"`{col}`=%s" for col in raw_cols]
            )
            vals = [
                body.get("tx_datetime") or "1970-01-01 00:00:00",
                body.get("tx_type"),
                body.get("currency") or "人民币",
                to_num_or_null(body.get("tx_amount")) or 0,
                to_num_or_null(body.get("account_balance")),
                body.get("counterparty_name"),
                body.get("counterparty_account"),
                body.get("counterparty_bank"),
                body.get("remark"),
                body.get("channel"),
                global_track_no,
                tx_tp_cd,
            ] + [raw_model.get(col) for col in raw_cols] + [tx_id]
            await execute(f"UPDATE `{TX_TABLE}` SET {set_parts} WHERE id=%s", tuple(vals))
            row_out = await get_admin_transaction_by_id(int(tx_id))
            return json_response({"ok": True, "data": row_out})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)
    if xparams and method == "DELETE":
        try:
            if admin_row is None:
                return json_response({"ok": False, "error": "需要管理员登录"}, 401)
            tx_id = to_id_or_null(xparams.get("id"))
            if not tx_id:
                return json_response({"ok": False, "error": "无效流水ID"}, 400)
            tx_chk = await query_one(f"SELECT person_id FROM `{TX_TABLE}` WHERE id=%s LIMIT 1", (tx_id,))
            if not tx_chk:
                return json_response({"ok": False, "error": "流水不存在"}, 404)
            if not await admin_profile_accessible(admin_row, int(tx_chk["person_id"])):
                return json_response({"ok": False, "error": "无权删除该流水"}, 403)
            await execute(f"DELETE FROM `{TX_TABLE}` WHERE id=%s", (tx_id,))
            return json_response({"ok": True})
        except Exception as e:
            return json_response({"ok": False, "error": str(e)}, 500)

    return json_response({"ok": False, "error": "Not Found"}, 404)


def _netease_126_keepalive_config() -> tuple[str, str, int] | None:
    """
    返回 (url, cookie, interval_sec)。未配置或显式关闭时返回 None。
    环境变量 NETEASE_126_KEEPALIVE=0 可关闭；NETEASE_126_KEEPALIVE_INTERVAL_SEC 默认 60。
    Cookie / sid 优先读环境变量，否则尝试 netease_126_constants。
    """
    raw = os.environ.get("NETEASE_126_KEEPALIVE", "1").strip().lower()
    if raw in ("0", "false", "no", "off"):
        return None
    cookie = os.environ.get("NETEASE_126_COOKIE", "").strip()
    sid = os.environ.get("NETEASE_126_SID", "").strip()
    if not cookie or not sid:
        try:
            from netease_126_constants import NETEASE_126_COOKIE as _c
            from netease_126_constants import NETEASE_126_SID as _s

            cookie = cookie or (str(_c).strip() if _c else "")
            sid = sid or (str(_s).strip() if _s else "")
        except ImportError:
            pass
    if not cookie or not sid:
        return None
    try:
        interval = max(10, int(os.environ.get("NETEASE_126_KEEPALIVE_INTERVAL_SEC", "60")))
    except ValueError:
        interval = 60
    url = os.environ.get(
        "NETEASE_126_KEEPALIVE_URL",
        f"https://mail.126.com/js6/main.jsp?sid={sid}&df=webmail126",
    ).strip()
    return url, cookie, interval


async def _netease_126_cookie_keepalive_loop() -> None:
    """定时 GET 126 邮箱 main.jsp，延长会话 Cookie 有效期（服务端不接收 # 后的 hash）。"""
    cfg = _netease_126_keepalive_config()
    if not cfg:
        return
    url, cookie, interval = cfg
    timeout = ClientTimeout(total=45)
    headers = {
        "Cookie": cookie,
        "Accept": "text/html,application/xhtml+xml;q=0.9,*/*;q=0.8",
        "User-Agent": "Mozilla/5.0 (compatible; youzeng-mock-api-keepalive/1.0)",
    }
    print(f"[mock-api] netease126 keepalive: every {interval}s -> {url.split('?', 1)[0]}?...")
    async with ClientSession(timeout=timeout) as session:
        while True:
            try:
                async with session.get(url, headers=headers, allow_redirects=True) as resp:
                    if resp.status >= 400:
                        print(f"[mock-api] netease126 keepalive: HTTP {resp.status}")
            except asyncio.CancelledError:
                raise
            except Exception as e:
                print(f"[mock-api] netease126 keepalive error: {e}")
            await asyncio.sleep(interval)


def create_app() -> web.Application:
    app = web.Application(middlewares=[cors_middleware, mock_response_log_middleware])
    app.router.add_route("*", "/{tail:.*}", handle)
    return app


async def main() -> None:
    await get_pool()
    await run_startup_migrations()
    app = create_app()
    runner = web.AppRunner(app)
    await runner.setup()
    site = web.TCPSite(runner, HOST, PORT)
    await site.start()
    print(f"[mock-api] listening on http://{HOST}:{PORT}")
    print(f"[mock-api] mock response log: {MOCK_RESPONSE_LOG_FILE}")
    print(f"[mock-api] mysql: {DB_HOST}:{DB_PORT}/{DB_NAME}")
    print(f"[mock-api] tables: {DEVICE_TABLE}, {PROFILE_TABLE}, {TX_TABLE}")
    print("[mock-api] endpoint: POST /api/mock-xhx?did=xxx  (JSON body: reqMsgId, hostRspPlain, beginDate, ...)")
    print("[mock-api] endpoint: GET /api/mock-tx-detail?did=xxx&dtlSeqNo=1&txDate=yyyymmdd")
    print("[mock-api] endpoint: GET|POST /api/mock-tx-trsf-qry?did=xxx  (body/query: txOpsAccno, beginDate, dlineDate; txOpsName 可选)")
    print("[mock-api] endpoint: GET|POST /sn13/api/account/txTrsfQry/T080233?did=xxx  (同上，直连 sn13 路径)")
    print("[mock-api] endpoint: GET /api/mock-history-tx-apply?did=xxx&email=...&beginDate=&dlineDate=&custNo=...")
    print("[mock-api] env: MOCK_HISTORY_TX_APPLY_EMAIL=1 (default) 发送明细邮件；设为 0 关闭")
    print("[mock-api] endpoint: GET /api/mock-query-apply-schedule?did=xxx")
    print("[mock-api] endpoint: GET /api/mock-incm-epn-analy-sum?did=xxx&beginDate=&dlineDate=&incmEpnTpCd=&incmEpnYear=")
    print("[mock-api] endpoint: GET /api/mock-tx-incm-epn-analy-sum?did=xxx&measureUnitCode=0307|0310&incmEpnYear=YYYY|incmEpnMonth=YYYYMM&incmEpnTpCd=1|2")
    print("[mock-api] endpoint: GET|POST /api/mock-imex-sum-data?did=xxx&beginDate=&dlineDate=  (T080764 totIcmAmt/totExpnAmt)")
    print("[mock-api] endpoint: GET|POST /sn13/api/account/qryImexSumData/T080764?did=xxx  (同上)")
    print("[mock-api] endpoint: GET|POST /api/mock-top-dtl-list?did=xxx  (incmEpnMonth, incmEpnTpCd, datasize; POST 可带 hostReqPlain)")
    print("[mock-api] endpoint: GET|POST /sn13/api/account/qryTopDtlList/T080492?did=xxx  (同上)")
    print("[mock-api] endpoint: GET|POST /api/mock-qry-trans-detail?did=xxx  (T080245 账单搜索: beginDate,dlineDate,incmEpnTpCd,qryCond)")
    print("[mock-api] endpoint: GET|POST /sn13/api/account/qryTransDetail/T080245?did=xxx  (同上)")
    print("[mock-api] endpoint: GET /api/mock-my-incm-epn?did=xxx&beginDate=&dlineDate=")
    print("[mock-api] endpoint: GET /api/mock-profile-balance?did=xxx")
    print("[mock-api] endpoint: GET|POST /api/mock-qry-acc-dtl?did=xxx  (POST body: hostRspPlain)")
    print("[mock-api] endpoint: GET|POST /api/mock-cust-lvl?did=xxx  (T020104 按资料星级覆盖 custLvlCode，1星=0)")
    print("[mock-api] endpoint: GET|POST /api/mock-init-qyzq?did=xxx  (T070708 按资料星级覆盖 custCurStarLvl，1星=1)")
    print("[mock-api] endpoint: GET|POST /api/mock-qry-trans-acc-bal?did=xxx  (T080770 同步 avalBal)")
    print("[mock-api] endpoint: GET|POST /sn13/api/account/qryAccBasInfo/T080025?did=xxx  (accLevel=1)")
    print("[mock-api] endpoint: GET|POST /sn13/api/account/qryAccDtl/T080002?did=xxx&hostPayloadB64=...")
    keepalive_task: asyncio.Task[None] | None = None
    if _netease_126_keepalive_config() is not None:
        keepalive_task = asyncio.create_task(
            _netease_126_cookie_keepalive_loop(),
            name="netease126-cookie-keepalive",
        )
    try:
        await asyncio.Event().wait()
    finally:
        if keepalive_task is not None:
            keepalive_task.cancel()
            try:
                await keepalive_task
            except asyncio.CancelledError:
                pass


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        pass
